# -*- coding: utf-8 -*-
"""Agente Faturas Amil - gravação.
Grava na Base_Faturas_Vidas.xlsx o JSON gerado por faturas_amil.py e copia os PDFs para Faturas/AAAA-MM/<Operadora>/.
Uso: python3 gravar_amil.py <amil.json> <planilha entrada> <planilha saida> <pasta FATURAS> <AAAA-MM-DD> [--sem-copia]
Regras:
- Uma empresa (aba Empresas) por contrato que recebe a NF/boleto (contrato principal). Código estipulante = contrato (10 díg.),
  igual ao que aparece no extrato de comissão da Amil. Os demais contratos do mesmo demonstrativo vão para Subestipulantes.
- Competência = mês da "Mensalidade - MM/AAAA" do demonstrativo (= período de cobertura do boleto).
- Nº fatura = nº da NF (N. Fiscal), o mesmo do boleto/NFS-e.
- Operadora: AMIL -> Amil (planos DENTAL -> Amil Dental); SHAM -> Santa Helena.
- Vidas: uma linha por beneficiário (marca ótica) por fatura. Mens.* = prêmio base; Co-participação = copart;
  Cobrança/Devolução retroativa e descontos = pró-rata/movimento. Sem mensalidade no mês = Excluído.
- Copart separada por tipo quando há o Demonstrativo de Utilização de Coparticipação da mesma NF e a soma bate.
- Fatura já gravada (mesma competência + empresa + NF) não é gravada de novo.
- Contratos de empresas vinculadas como subestipulantes (GRUPOS em agrupar_subestipulantes.py) são gravados na empresa principal.
"""
import json, sys, os, re, copy, shutil, calendar
from datetime import datetime
from collections import defaultdict, OrderedDict, Counter
import openpyxl
from agrupar_subestipulantes import DE, aplicar as agrupar_subs  # subs vinculadas à empresa principal
from shaddai_comum import d, br, cs, last_row, fim_mes, mmaa, estender  # módulo comum
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.cell_range import MultiCellRange

OPS = {'AMIL': 'Amil', 'SHAM': 'Santa Helena'}
PASTA = {'Amil': 'Amil', 'Amil Dental': 'AmilDental', 'Santa Helena': 'SantaHelena'}
PROD = {'Amil': 'AMIL SAUDE', 'Amil Dental': 'AMIL DENTAL', 'Santa Helena': 'SANTA HELENA SAUDE'}
PAREN = {'Conjuge': 'Cônjuge', 'Filho/Filha': 'Filho(a)', 'Companheiro': 'Companheiro(a)'}
SUBFORM = ('=IF(C{r}="","",IF(ISNUMBER(SEARCH("Bradesco",IFERROR(INDEX(Empresas!$D:$D,MATCH(A{r},Empresas!$A:$A,0)),""))),'
           'B{r}&" - "&G{r},C{r}&" - "&G{r}))')
NEMP = 300   # empresas cobertas pela aba Faturamento


def grupo_de(rel):
    p = rel.split('/')
    return p[1] if len(p) > 2 and re.match(r'\d\d-\d\d$', p[0]) else None


def prefixo(rel):
    m = re.match(r'(\d{10})', os.path.basename(rel))
    return m.group(1) if m else None


def nome_curto_arquivo(rel):
    b = os.path.basename(rel)
    m = re.match(r'^\d{10}\s*-?\s*(.*?)\s*-?\s*(?:FATURADOS|BOLETO|COPART|RE\s+FAT|AMIL\b|SANT|SAUDE)', b, re.I)
    if m and m.group(1).strip(' -'):
        return re.sub(r'\s+', ' ', m.group(1).strip(' -')).upper()
    return None


def operadora_fatura(f):
    op = OPS.get(f.get('operadora') or '', None)
    if op is None:
        op = 'Santa Helena' if 'SANT' in f['rel'].upper() else 'Amil'
    if op == 'Amil' and f['linhas'] and sum(1 for x in f['linhas'] if x['plano'].upper().startswith('DENTAL')) > len(f['linhas']) / 2:
        op = 'Amil Dental'
    return op


def copart_tipo(g):
    g = (g or '').lower()
    if 'pronto' in g: return 23
    if 'interna' in g: return 24
    if 'consulta' in g: return 21
    if 'exame' in g or 'terap' in g or 'procedim' in g: return 22
    return 32


def preparar(docs, log):
    fats = [x for x in docs if x['tipo'] == 'fatura' and x['contratos']]
    bols = [x for x in docs if x['tipo'] == 'boleto']
    cops = [x for x in docs if x['tipo'] == 'copart']
    for x in docs:
        if x['tipo'] in ('ignorado', 'outro', 'erro') or (x['tipo'] == 'fatura' and not x['contratos']):
            motivo = x.get('motivo') or x.get('erro') or ('não é demonstrativo Amil (implantação Porto Odonto)' if x['tipo'] == 'fatura'
                                                          else 'não é fatura nem boleto (ex.: carteirinha, relatório de outra operadora)')
            log.append((None, None, 'Extração', f'Arquivo não lido: {x["rel"]} ({motivo}).', 'Ignorado pelo Agente Faturas Amil.', 'Sim'))
    conhecidos = {prefixo(f['rel']) for f in fats if prefixo(f['rel']) in f['contratos']}
    for f in fats:
        p = prefixo(f['rel']); ks = list(f['contratos'])
        if p in ks:
            f['principal'] = p
        else:
            inter = [k for k in ks if k in conhecidos]
            f['principal'] = inter[0] if inter else ks[0]
            if p and p not in ks:
                log.append((f['ref'], f['principal'], 'Conferência',
                            f'Arquivo {f["rel"]} tem no nome o contrato {p}, mas o conteúdo é do contrato {f["principal"]} ({f["contratos"][f["principal"]]["nome"]}).',
                            'Gravado pelo conteúdo. Verificar se falta a fatura do contrato do nome do arquivo.', 'Não'))
        c0 = f['contratos'][f['principal']]
        f['nf'] = c0.get('nf') or next((v['nf'] for v in f['contratos'].values() if v.get('nf')), None)
        f['venc'] = c0.get('venc'); f['emissao'] = c0.get('emissao')
        f['op'] = operadora_fatura(f)
        if not f.get('ref'):
            f['ref'] = f['venc'][:7]
    por_nf = OrderedDict()
    for f in sorted(fats, key=lambda x: (x.get('origem') == 'TXT', 'FATURAD' not in x['rel'].upper(), x['rel'])):
        k = (f['principal'], f['nf'])
        if k in por_nf:
            log.append((f['ref'], f['principal'], 'Conferência', f'NF {f["nf"]} repetida: {f["rel"]} (já lida de {por_nf[k]["rel"]}).',
                        'Arquivo repetido ignorado (gravada uma vez só).', 'Sim'))
            continue
        por_nf[k] = f
    fats = list(por_nf.values())
    bol_nf = {}
    for b in bols:
        bol_nf.setdefault(b.get('nf'), b)
    usados = set()
    for f in fats:
        f['boleto'] = bol_nf.get(f['nf'])
        if f['boleto']: usados.add(f['nf'])
    principais = {f['principal'] for f in fats}
    so_boleto = []
    for nf, b in bol_nf.items():
        if nf in usados: continue
        c9 = b.get('contrato9') or ''
        p = prefixo(b['rel'])
        cand = [k for k in principais if c9 and k.startswith(c9)]
        if p and c9 and p.startswith(c9): princ = p
        elif len(cand) == 1: princ = cand[0]
        else: princ = (c9 + '0') if c9 else p
        if p and c9 and not p.startswith(c9):
            log.append(((b['cob'][0] or b['venc'])[:7], princ, 'Conferência',
                        f'Boleto {b["rel"]} tem no nome o contrato {p}, mas é do contrato {c9}x ({b.get("tomador")}).', 'Gravado pelo conteúdo.', 'Não'))
        comp = (b['cob'][0] or b['venc'])[:7]
        so_boleto.append(dict(principal=princ, nf=nf, boleto=b, ref=comp, venc=b['venc'], emissao=None, rel=b['rel'], linhas=[], contratos={},
                              op='Santa Helena' if 'SANTA HELENA' in (b.get('prestador') or '').upper() else ('Amil Dental' if 'ODONTO' in b['rel'].upper() else 'Amil'),
                              total=b.get('valor_nota') or b.get('valor'), so_boleto=True))
    for f in fats:
        if not f['boleto']:
            log.append((f['ref'], f['principal'], 'Download', f'Sem boleto/NFS-e para a NF {f["nf"]} ({f["rel"]}).',
                        'Fatura gravada sem dados do boleto (conferência do boleto em branco).', 'Não'))
    todas = fats + so_boleto
    cop_nf = {c['nf']: c for c in cops if c.get('nf')}
    for f in so_boleto:
        cp = cop_nf.get(f['nf'])
        if not cp: continue
        f['so_copart'] = True
        for cod, gr in cp['por'].items():
            f['linhas'].append(dict(contrato=cp['contrato'], cod=cod, nome=cp.get('nomes', {}).get(cod), mat=None, cpf=None, plano=None,
                                    tp='T' if cp.get('titular', {}).get(cod) == cod else 'D', idade=None, dep=None, limite=None, inclusao=None,
                                    rubrica='Co-participação', classe='copart', valor=round(sum(gr.values()), 2)))
        if cp.get('contrato') and cp['contrato'] != f['principal']:
            f['principal_cop'] = cp['contrato']
    for f in so_boleto:
        if not f.get('so_copart'):
            log.append((f['ref'], f['principal'], 'Download', f'Só boleto/NFS-e, sem o Demonstrativo Analítico (NF {f["nf"]}, {f["rel"]}).',
                        'Fatura gravada com o valor da NF; sem vidas.', 'Não'))
    info = defaultdict(lambda: dict(comps=[], nomes=Counter(), grupo=None, subs=OrderedDict()))
    for f in sorted(todas, key=lambda x: x['ref']):
        i = info[f['principal']]
        i['comps'].append(f['ref']); i['op'] = f['op']; i['venc'] = f['venc']
        nc = nome_curto_arquivo(f['rel']) if prefixo(f['rel']) == f['principal'] else None
        if nc: i['nomes'][nc] += 1 + (f['ref'] >= '2026-01')
        if f['contratos']:
            i['nome_contrato'] = f['contratos'][f['principal']]['nome']
        b = f.get('boleto')
        if b and b.get('tomador') and (not f.get('so_boleto') or prefixo(b['rel']) in (None, f['principal'])):
            i['razao'] = b['tomador']; i['cnpj'] = b.get('cnpj')
        elif b and b.get('tomador') and not i.get('razao'):
            i['razao'] = b['tomador']; i['cnpj'] = b.get('cnpj')
        if f.get('cnpj') and not i.get('cnpj'): i['cnpj'] = f['cnpj']
        g = grupo_de(f['rel'])
        if g: i['grupo'] = re.sub(r'\s+\d\d-\d\d$', '', g)
        for k, v in f['contratos'].items():
            i['subs'][k] = v['nome']
    return todas, info, cop_nf


def cadastrar(wb, info, proc):
    E, S, LS = wb['Empresas'], wb['Subestipulantes'], wb['Listas']
    ops = [x for x in (LS.cell(r, 1).value for r in range(2, 31)) if x]
    novos = [o for o in ('Amil Dental', 'Santa Helena') if o not in ops]
    if novos:
        i0 = ops.index('Amil') + 1
        for o in reversed(novos): ops.insert(i0, o)
        for i, v in enumerate(ops):
            cs(LS.cell(3, 1), LS.cell(2 + i, 1)); LS.cell(2 + i, 1).value = v
    rE = last_row(E); tplE = rE
    nid = max(int(E.cell(r, 1).value[1:]) for r in range(2, rE + 1) if E.cell(r, 1).value)
    ids = {}
    for r in range(2, rE + 1):
        if E.cell(r, 4).value in ('Amil', 'Amil Dental', 'Santa Helena'):
            ids[str(E.cell(r, 5).value)] = E.cell(r, 1).value; info[str(E.cell(r, 5).value)]['nc'] = E.cell(r, 11).value
    novas = []
    for princ in sorted(info, key=lambda k: (info[k]['grupo'] or 'ZZZ', k)):
        if princ in ids: continue
        i = info[princ]
        nid += 1; eid = f'E{nid:03d}'; ids[princ] = eid; rE += 1
        nc = i['nomes'].most_common(1)[0][0] if i['nomes'] else re.sub(r'\s+', ' ', (i.get('nome_contrato') or i.get('razao') or princ)).upper()[:25]
        i['nc'] = nc
        ult = max(i['comps'])
        obs = f'Cadastrada pelo Agente Faturas Amil a partir dos demonstrativos/boletos ({proc:%d/%m/%Y}). Código estipulante = nº do contrato Amil.'
        if i['grupo']: obs += f' Grupo/pasta: {i["grupo"]}.'
        if len(i['subs']) > 1: obs += f' {len(i["subs"])} contratos no mesmo demonstrativo (aba Subestipulantes).'
        obs += f' Faturas na pasta: {min(i["comps"])} a {ult}.'
        row = [eid, i.get('razao') or i.get('nome_contrato'), i.get('cnpj'), i['op'], princ, princ, None,
               int(i['venc'][8:10]) if i.get('venc') else None, 'Sim' if ult >= '2026-09' else 'Não', obs, nc,
               f'=IF(A{rE}="","",COUNTIFS(Subestipulantes!$A:$A,A{rE}))']
        for c, v in enumerate(row, 1):
            cs(E.cell(tplE, c), E.cell(rE, c)); E.cell(rE, c).value = v
        novas.append(eid)
    rS = last_row(S); tplS = rS
    exist = {(S.cell(r, 1).value, str(S.cell(r, 3).value)) for r in range(2, rS + 1)}
    for princ, i in info.items():
        subs = dict(i['subs']) or {princ: i.get('nome_contrato') or i.get('razao')}
        subs.setdefault(princ, i.get('nome_contrato') or i.get('razao'))
        pe = DE.get(ids[princ], ids[princ])  # contrato de empresa que já é sub: cadastra na principal
        for k, nome in subs.items():
            if (pe, k) in exist or (ids[princ], k) in exist: continue
            exist.add((pe, k)); rS += 1
            row = [pe, princ, k, None, nome, i.get('cnpj') if k == princ else None,
                   i['nc'] if k == princ else re.sub(r'\s+', ' ', nome or k).upper()[:25], 'Sim', SUBFORM.format(r=rS),
                   'Contrato principal (recebe a NF/boleto).' if k == princ else 'Contrato Amil faturado junto, na NF do contrato principal.']
            for c, v in enumerate(row, 1):
                cs(S.cell(tplS, c), S.cell(rS, c)); S.cell(rS, c).value = v
    estender(E, NEMP + 1, 12, tplE)
    estender(S, max(rS, S.max_row), 10, tplS)
    return ids, novas


def faturamento(FT):
    if FT.cell(6 + NEMP, 2).value == 'Total da carteira':
        return
    N0 = 200
    if FT.cell(6 + N0, 2).value != 'Total da carteira':
        raise SystemExit('Layout da aba Faturamento inesperado')

    def snap(r):
        return [copy.copy(FT.cell(r, c)._style) for c in range(1, 18)], FT.row_dimensions[r].height
    sd1, st1, sgap, sti2, sh2, sd2, st2 = (snap(r) for r in (6, 6 + N0, 7 + N0, 9 + N0, 10 + N0, 11 + N0, 11 + 2 * N0))
    head2 = [FT.cell(10 + N0, c).value for c in range(1, 18)]; title2 = FT.cell(9 + N0, 1).value
    for r in range(6, 6 + 2 * NEMP + 20):
        for c in range(1, 18): FT.cell(r, c).value = None

    def apply(r, st):
        for c in range(1, 18): FT.cell(r, c)._style = copy.copy(st[0][c - 1])
        FT.row_dimensions[r].height = st[1]
    mcols = [L(c) for c in range(4, 16)]
    t1 = 6 + NEMP
    for blk, (d0, tot, cv, hdr) in enumerate([(6, t1, 'J', 5), (t1 + 5, t1 + 5 + NEMP, 'L', t1 + 4)]):
        if blk == 1:
            apply(t1 + 1, sgap); apply(t1 + 2, sgap); apply(t1 + 3, sti2); FT.cell(t1 + 3, 1).value = title2
            apply(t1 + 4, sh2)
            for c, v in enumerate(head2, 1): FT.cell(t1 + 4, c).value = v
        for i in range(NEMP):
            r = d0 + i; apply(r, sd1 if blk == 0 else sd2)
            FT.cell(r, 1).value = f'=IF(INDEX(Empresas!$A$2:$A${NEMP+1},{i+1})="","",INDEX(Empresas!$A$2:$A${NEMP+1},{i+1}))'
            FT.cell(r, 2).value = f'=IF($A{r}="","",INDEX(Empresas!$B$2:$B${NEMP+1},{i+1}))'
            FT.cell(r, 3).value = f'=IF($A{r}="","",INDEX(Empresas!$D$2:$D${NEMP+1},{i+1}))'
            for col in mcols:
                FT[f'{col}{r}'] = f'=IF($A{r}="","",SUMIFS(Faturas!${cv}:${cv},Faturas!$A:$A,{col}${hdr},Faturas!$B:$B,$A{r}))'
            FT[f'P{r}'] = f'=IF($A{r}="","",SUM(D{r}:O{r}))'
            if blk == 0: FT[f'Q{r}'] = f'=IF(OR($A{r}="",$P${t1}=0),"",P{r}/$P${t1})'
        apply(tot, st1 if blk == 0 else st2); FT.cell(tot, 2).value = 'Total da carteira'
        for col in mcols + ['P']: FT[f'{col}{tot}'] = f'=SUM({col}{d0}:{col}{tot-1})'


def gravar_faturas(wb, todas, info, ids, cop_nf, proc, log):
    F, V = wb['Faturas'], wb['Vidas']
    rF = last_row(F); tplF = rF
    rV = last_row(V); tplV = rV
    ja = {(F.cell(r, 1).value, F.cell(r, 2).value, str(F.cell(r, 5).value)) for r in range(2, rF + 1)}
    res = Counter(); copias = []
    for f in sorted(todas, key=lambda x: (x['ref'], ids[x['principal']])):
        eid = DE.get(ids[f['principal']], ids[f['principal']]); comp = f['ref']; nf = f['nf']; op = f['op']; b = f.get('boleto') or {}
        if (comp, eid, nf) in ja:
            res['ja_na_planilha'] += 1; continue
        ja.add((comp, eid, nf))
        nc = info[f['principal']]['nc']
        base = f'Faturas/{comp}/{PASTA[op]}/{f["principal"]} - {nc}/{f["principal"]} - {nc}'
        arq_f = arq_b = None
        if not f.get('so_boleto'):
            arq_f = f'{base} FATURADOS {PROD[op]} {mmaa(comp)}{".txt" if f.get("origem") == "TXT" else ".pdf"}'
            copias.append((f['arquivo'], arq_f))
        if b:
            arq_b = f'{base} BOLETO {PROD[op]} {mmaa(comp)}.pdf'; copias.append((b['arquivo'], arq_b))
        cp = cop_nf.get(nf)
        if cp: copias.append((cp['arquivo'], f'{base} COPART {PROD[op]} {mmaa(comp)}.pdf'))
        vinf = sum(v.get('vidas') or 0 for v in f['contratos'].values()) if f['contratos'] else None
        subs = [k for k in f['contratos'] if k != f['principal']]
        mov = round(sum(x['valor'] for x in f['linhas'] if x['classe'] == 'mov'), 2)
        copt = round(sum(x['valor'] for x in f['linhas'] if x['classe'] == 'copart'), 2)
        obs = ['Fatura só de coparticipação (boleto + Demonstrativo de Utilização de Coparticipação; sem mensalidade).' if f.get('so_copart') else
               'Só boleto/NFS-e: falta o Demonstrativo Analítico (sem vidas).' if f.get('so_boleto') else
               f'Origem: {"TXT (delimitador #)" if f.get("origem") == "TXT" else "PDF Demonstrativo Analítico de Faturamento"}.']
        if subs: obs.append(f'{len(f["contratos"])} contratos na mesma NF.')
        if mov: obs.append(f'Movimentos retroativos/descontos: R$ {br(mov)}.')
        if copt: obs.append(f'Coparticipação: R$ {br(copt)}' + (' (por tipo, do demonstrativo de copart).' if cp else '.'))
        if f.get('venc') and f['venc'][:7] != comp:
            obs.append(f'Vencimento {f["venc"][8:10]}/{f["venc"][5:7]}/{f["venc"][:4]}.')
        obs.append(f'Arquivo original: AMIL/{f["rel"]}.')
        cob = b.get('cob') or (None, None)
        ini = d(cob[0]) or datetime(int(comp[:4]), int(comp[5:7]), 1)
        fim = d(cob[1]) or fim_mes(comp)
        rF += 1; r = rF
        row = {1: comp, 2: eid, 4: op, 5: nf, 6: d(f.get('emissao')), 7: d(f.get('venc')), 8: ini, 9: fim,
               10: f.get('total'), 11: vinf, 17: b.get('valor'), 19: b.get('nosso'), 20: b.get('linha'),
               21: arq_f, 22: arq_f, 23: arq_b, 24: proc, 25: ' '.join(obs), 26: f['principal'], 27: ', '.join(subs) or None}
        for col in range(1, 33):
            cs(F.cell(tplF, col), F.cell(r, col)); F.cell(r, col).value = row.get(col)
        F.cell(r, 3).value = f'=IF(B{r}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{r},Empresas!$A:$A,0)),"ID não cadastrado"))'
        F.cell(r, 12).value = f'=IF(A{r}="","",COUNTIFS(Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r},Vidas!$P:$P,"Ativo"))'
        F.cell(r, 13).value = f'=IF(A{r}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r}))'
        F.cell(r, 14).value = f'=IF(A{r}="","",J{r}-M{r})'
        F.cell(r, 15).value = f'=IF(A{r}="","",IF(L{r}=0,"Sem vidas",IF(ABS(N{r})<=0.01,"OK","Divergência")))'
        F.cell(r, 16).value = f'=IF(OR(A{r}="",K{r}=""),"",IF(L{r}=K{r},"OK","Divergência"))'
        F.cell(r, 18).value = f'=IF(OR(A{r}="",Q{r}=""),"",IF(ABS(Q{r}-J{r}-N(AE{r}))<=0.01,"OK","Divergência"))'
        F.cell(r, 29).value = f'=IF(A{r}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{r}))'
        F.cell(r, 30).value = (f'=IF(A{r}="","",IF(AC{r}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{r},'
                               f'Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')
        res['faturas'] += 1; res['so_boleto'] += bool(f.get('so_boleto')) and not f.get('so_copart'); res['so_copart'] += bool(f.get('so_copart'))
        pessoas = OrderedDict()
        for x in f['linhas']:
            pessoas.setdefault(x['cod'], []).append(x)
        ativos = 0; soma = 0.0
        for cod, ls in pessoas.items():
            x0 = next((x for x in ls if x['classe'] == 'base'), ls[0])
            bv = round(sum(x['valor'] for x in ls if x['classe'] == 'base'), 2)
            mv = round(sum(x['valor'] for x in ls if x['classe'] == 'mov'), 2)
            cv = round(sum(x['valor'] for x in ls if x['classe'] == 'copart'), 2)
            tipos = defaultdict(float)
            if cv and cp and cod in cp['por'] and abs(sum(cp['por'][cod].values()) - cv) <= 0.01:
                for g, v in cp['por'][cod].items(): tipos[copart_tipo(g)] += v
                res['copart_por_tipo'] += 1
            elif cv:
                tipos[32] = cv
            ativo = any(x['classe'] == 'base' for x in ls)
            ativos += ativo; soma += bv + mv + cv
            ob = []
            if mv or any(x['classe'] == 'mov' for x in ls):
                ob.append('Movimento: ' + '; '.join(f"{x['rubrica']} {br(x['valor'])}" for x in ls if x['classe'] == 'mov'))
            if not ativo:
                ob.append('Sem mensalidade nesta fatura (só ' + ', '.join(sorted({x['rubrica'] for x in ls})) + ')')
                if f.get('so_copart'): ob.append('dados do demonstrativo de coparticipação (sem CPF/plano)')
            if x0.get('limite'):
                ob.append(f"Data limite {x0['limite'][8:10]}/{x0['limite'][5:7]}/{x0['limite'][:4]}")
            tp = x0['tp']
            paren = 'Titular' if tp == 'T' else PAREN.get(x0.get('dep') or '', x0.get('dep') or ('Agregado' if tp == 'A' else 'Dependente'))
            rV += 1
            vals = {1: comp, 2: eid, 3: op, 4: nf, 5: cod, 7: x0['nome'], 8: x0.get('cpf'), 10: x0.get('idade'), 11: paren,
                    12: 'Titular' if tp == 'T' else 'Dependente', 14: x0['plano'], 15: x0.get('mat'),
                    16: 'Ativo' if ativo else 'Excluído', 17: d(x0.get('inclusao')), 18: d(x0.get('exclusao')),
                    19: bv, 20: mv, 21: round(tipos.get(21, 0), 2), 22: round(tipos.get(22, 0), 2), 23: round(tipos.get(23, 0), 2),
                    24: round(tipos.get(24, 0), 2), 25: 0, 29: '; '.join(ob) or None, 30: x0['contrato'], 32: round(tipos.get(32, 0), 2)}
            for col in range(1, 33):
                cs(V.cell(tplV, col if col < 32 else 21), V.cell(rV, col)); V.cell(rV, col).value = vals.get(col)
            V.cell(rV, 26).value = f'=IF(A{rV}="","",N(U{rV})+N(V{rV})+N(W{rV})+N(X{rV})+N(AF{rV}))'
            V.cell(rV, 27).value = f'=IF(A{rV}="","",N(S{rV})+N(T{rV})+N(Z{rV})+N(Y{rV}))'
            V.cell(rV, 28).value = f'=IF(A{rV}="","",IF(P{rV}="Excluído","Exclusão",IF(N(T{rV})>0,"Inclusão (pró-rata)","-")))'
            res['vidas'] += 1
        if f['linhas']:
            ok_v = abs(round(soma, 2) - (f.get('total') or 0)) <= 0.01
            ok_n = vinf is None or ativos == vinf
            b_ok = (not b) or abs((b.get('valor') or 0) - (f.get('total') or 0)) <= 0.01
            res['valor_ok'] += ok_v; res['vidas_ok'] += ok_n; res['boleto_ok'] += bool(b) and b_ok
            res['ok' if (ok_v and ok_n and b_ok and b) else 'pend'] += 1
            if not ok_v:
                log.append((comp, f['principal'], 'Conferência', f'Soma das vidas R$ {br(soma)} ≠ total da fatura R$ {br(f["total"])} (NF {nf}).', 'Verificar o PDF.', 'Não'))
            if not ok_n:
                log.append((comp, f['principal'], 'Conferência', f'Vidas com mensalidade ({ativos}) ≠ "Total Contrato" da fatura ({vinf}) (NF {nf}).',
                            'Normal quando há beneficiário só com devolução/copart retroativa ou listado sem valor. Valores conferidos.', 'Sim'))
            if b and not b_ok:
                log.append((comp, f['principal'], 'Conferência', f'Boleto R$ {br(b["valor"])} ≠ fatura R$ {br(f["total"])} (NF {nf}).', 'Verificar.', 'Não'))
        else:
            res['pend'] += 1
    estender(F, max(rF, F.max_row), 32, tplF)
    estender(V, max(rV + 500, V.max_row), 32, tplV)
    for r in range(tplV + 1, V.max_row + 1):
        cs(V.cell(tplV, 21), V.cell(r, 32))
    return res, copias, rF, rV


def registrar(wb, res, log, info, ids, novas, comps, proc, backup):
    AG, EX, LG, EXH = wb['Agentes'], wb['Execuções'], wb['Log de Erros'], wb['Execução']
    ra = last_row(AG)
    if not any(AG.cell(r, 1).value == 'Faturas Amil' for r in range(2, ra + 1)):
        ra += 1
        for c in range(1, 13):
            cs(AG.cell(2, c), AG.cell(ra, c))
            v = AG.cell(2, c).value
            AG.cell(ra, c).value = re.sub(r'(?<=[A-Z$])2\b', str(ra), v) if isinstance(v, str) and v.startswith('=') else None
        for c, v in enumerate(['Faturas Amil', 'Lê os demonstrativos analíticos (PDF/TXT) e boletos da Amil, Amil Dental e Santa Helena que a equipe coloca na pasta e grava Faturas + Vidas',
                               'Amil', 'Empresas, Subestipulantes, Faturas, Vidas, Log de Erros', 'Mensal', 10,
                               'Scripts/faturas_amil.py + Scripts/gravar_amil.py', 'Ativo'], 1):
            AG.cell(ra, c).value = v
    pend = sum(1 for x in log if x[5] == 'Não')
    re_ = last_row(EX) + 1
    for c in range(1, 14): cs(EX.cell(re_ - 1, c), EX.cell(re_, c))
    for c, v in enumerate(['Faturas Amil', f'{comps[0]} a {comps[-1]}', proc, datetime.now().replace(microsecond=0),
                           '⚠️ OK com pendência' if pend else '✅ OK', res['faturas'], res['ok'], res['pend'], 0,
                           f'{res["so_boleto"]} faturas só com boleto (falta o demonstrativo); {pend} pendências no Log de Erros',
                           backup, f'Primeira carga Amil a partir da pasta "boletos a enviar/AMIL": {res["faturas"]} faturas, {res["vidas"]} linhas de vidas, '
                           f'{len(novas)} empresas novas.', 'Bruno'], 1):
        EX.cell(re_, c).value = v
    rl = last_row(LG); tplL = rl
    for comp, k, etapa, erro, acao, ok in log:
        rl += 1
        for c in range(1, 10): cs(LG.cell(tplL, c), LG.cell(rl, c))
        op = info[k]['op'] if k in info else 'Amil'
        for c, v in enumerate([proc, comp, op, ids.get(k) if k else None, etapa, erro, acao, ok, 'Agente Faturas'], 1):
            LG.cell(rl, c).value = v
    estender(LG, max(rl, LG.max_row), 9, tplL)
    rh = last_row(EXH, 2) + 1
    for c in range(2, 10): cs(EXH.cell(rh - 1, c), EXH.cell(rh, c))
    for c, v in enumerate([proc, f'{comps[0]} a {comps[-1]}', f'Amil: {len(info)} contratos principais (arquivos da equipe)', res['faturas'],
                           res['ok'], res['pend'], f'{res["so_boleto"]} só boleto', 'Primeira carga Amil / Amil Dental / Santa Helena a partir dos arquivos da equipe.'], 2):
        EXH.cell(rh, c).value = v


LEIA = ['AGENTE FATURAS AMIL (AMIL, AMIL DENTAL E SANTA HELENA)',
        'Sem portal por enquanto: a equipe salva os arquivos em "Área de Trabalho/boletos a enviar/AMIL/MM-AA/<grupo>/" e o agente lê a pasta inteira. Scripts: Scripts/faturas_amil.py (leitura) e Scripts/gravar_amil.py (gravação).',
        'Arquivos lidos: Demonstrativo Analítico de Faturamento (PDF, ou TXT com delimitador #), boleto + NFS-e (PDF) e Demonstrativo de Utilização de Coparticipação (PDF). O que não é Amil (Porto, Bradesco, carteirinhas) é ignorado e vai para o Log de Erros.',
        'Empresas: uma linha por contrato que recebe a NF/boleto (Código estipulante = contrato de 10 dígitos, o mesmo do extrato de comissão). Os outros contratos do mesmo demonstrativo (ex.: Quanta, Lipos) ficam em Subestipulantes.',
        'Competência = "Mensalidade - MM/AAAA" do demonstrativo (= período de cobertura do boleto). Nº fatura = nº da NF (N. Fiscal), igual ao do boleto. Operadora: AMIL = Amil (plano DENTAL = Amil Dental); SHAM = Santa Helena.',
        'Vidas: uma linha por beneficiário (Nº seguro = marca ótica). Mens. = prêmio base; Cobrança/Devolução retroativa e descontos = pró-rata; Co-participação separada por tipo quando há o demonstrativo de copart da mesma NF (senão vai em Copart sem tipo). Sem mensalidade no mês = Excluído.',
        'Vidas informadas = "Total Contrato" do demonstrativo, que conta também quem só tem devolução/copart; por isso a conferência de vidas pode dar Divergência com valores OK.',
        'Arquivos copiados para Faturas/AAAA-MM/Amil | AmilDental | SantaHelena/<contrato> - <NOME CURTO>/<contrato> - <NOME CURTO> FATURADOS | BOLETO | COPART <PRODUTO> <MM-AA>.pdf. O original fica na pasta da equipe (caminho nas Observações da fatura).']


def leia_me(wb):
    ws = wb['Leia-me']
    if any(ws.cell(r, 2).value == LEIA[0] for r in range(1, ws.max_row + 1)):
        return
    hdr = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or '').startswith('AGENTE FATURAS BRADESCO'))
    r = last_row(ws, 2) + 1
    for i, t in enumerate(LEIA):
        r += 1
        cs(ws.cell(hdr if i == 0 else hdr + 1, 2), ws.cell(r, 2)); ws.cell(r, 2).value = t


def main(js, xin, xout, base_dir, proc, *opts):
    proc = datetime.fromisoformat(proc)
    docs = json.load(open(js, encoding='utf-8'))
    log = []
    todas, info, cop_nf = preparar(docs, log)
    wb = openpyxl.load_workbook(xin)
    ids, novas = cadastrar(wb, info, proc)
    res, copias, rF, rV = gravar_faturas(wb, todas, info, ids, cop_nf, proc, log)
    faturamento(wb['Faturamento'])
    comps = sorted({f['ref'] for f in todas})
    backup = f'Backups/Base_Faturas_Vidas antes Faturas Amil {proc:%Y-%m-%d}.xlsx'
    agrupar_subs(wb, proc, registrar=False)  # código estipulante/subestipulante das faturas das subs
    registrar(wb, res, log, info, ids, novas, comps, proc, backup)
    leia_me(wb)
    wb.save(xout)
    n = 0
    if '--sem-copia' not in opts:
        for src, dst in copias:
            p = os.path.join(base_dir, dst)
            if not os.path.exists(p):
                os.makedirs(os.path.dirname(p), exist_ok=True); shutil.copy2(src, p); n += 1
    print(json.dumps(dict(res, novas_empresas=len(novas), ids=f'{novas[0]}..{novas[-1]}' if novas else '', copiados=n, a_copiar=len(copias),
                          log=len(log), log_pend=sum(1 for x in log if x[5] == 'Não'), ult_fatura=rF, ult_vida=rV, comps=comps), ensure_ascii=False))


if __name__ == '__main__':
    main(*sys.argv[1:])

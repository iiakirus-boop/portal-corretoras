# -*- coding: utf-8 -*-
"""Módulo comum dos agentes da Shaddai.

Funções e regras que vários agentes usam. Cada agente importa daqui, e não de outro agente: assim uma mudança
na SulAmérica não quebra a Hapvida, e a regra de agenciamento/vitalício fica escrita num lugar só.
Antes de mudar qualquer coisa aqui, rode os testes: python testes/regressao.py conferir

Conteúdo:
- planilha: d, br, dbr, cs, last_row, fim_mes, mmaa, estender, NEMP, SUBFORM
- regra da Shaddai: classe_por_taxa (acima de 50% = Agenciamento; até 10% = Vitalício)
- abas de comissão: preparar (cria "Pagamentos Comissão" e "Comissões"), _prox, _w, estilos e formatos
"""
import copy, re, calendar
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.cell_range import MultiCellRange


# ---------------------------------------------------------------- planilha (aba Faturas, Vidas, Empresas...)
SUBFORM = ('=IF(C{r}="","",IF(ISNUMBER(SEARCH("Bradesco",IFERROR(INDEX(Empresas!$D:$D,MATCH(A{r},Empresas!$A:$A,0)),""))),'
           'B{r}&" - "&G{r},C{r}&" - "&G{r}))')
NEMP = 300   # empresas cobertas pela aba Faturamento


def d(s):
    """'AAAA-MM-DD' -> datetime (None se vazio)."""
    return datetime.fromisoformat(s) if s else None


def br(x):
    """1234.5 -> '1.234,50'."""
    return f'{x:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def dbr(s):
    """'AAAA-MM-DD' -> 'DD/MM/AAAA'."""
    return f'{s[8:10]}/{s[5:7]}/{s[:4]}' if s else ''


def cs(src, dst):
    """Copia o estilo de uma célula para outra."""
    if src.has_style:
        dst._style = copy.copy(src._style)


def last_row(ws, col=1):
    """Última linha preenchida na coluna (ignora linhas só com formatação)."""
    r = ws.max_row
    while r > 1 and ws.cell(r, col).value in (None, ''):
        r -= 1
    return r


def fim_mes(comp):
    """'AAAA-MM' -> último dia do mês."""
    y, m = int(comp[:4]), int(comp[5:7])
    return datetime(y, m, calendar.monthrange(y, m)[1])


def mmaa(comp):
    """'AAAA-MM' -> 'MM-AA' (como a equipe escreve no nome dos arquivos)."""
    return comp[5:7] + '-' + comp[2:4]


def estender(ws, ate, ncols, tpl):
    """Estende formatação (da linha tpl) e validações da aba até a linha 'ate'."""
    for r in range(ws.max_row + 1, ate + 1):
        for c in range(1, ncols + 1):
            cs(ws.cell(tpl, c), ws.cell(r, c))
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(' '.join(re.sub(r'(\D)\d+$', lambda m: m.group(1) + str(ate), p) for p in str(dv.sqref).split()))


# ---------------------------------------------------------------- regra da Shaddai para comissão
AGENCIAMENTO_ACIMA_DE = 50   # taxa (%) acima disso = Agenciamento
VITALICIO_ATE = 10           # taxa (%) até isso = Vitalício
VERIFICAR = f'Verificar (taxa entre {VITALICIO_ATE}% e {AGENCIAMENTO_ACIMA_DE}%)'


def classe_por_taxa(perc):
    """Classe da comissão pela taxa: Agenciamento, Vitalício ou VERIFICAR (entre as duas faixas)."""
    if perc > AGENCIAMENTO_ACIMA_DE:
        return 'Agenciamento'
    if perc <= VITALICIO_ATE:
        return 'Vitalício'
    return VERIFICAR


# ---------------------------------------------------------------- abas de comissão (todas as operadoras)
AZUL = Font(color="FF0000FF")
PRETO = Font(color="FF000000")
CINZA = PatternFill("solid", fgColor="FFEEF1F4")
CAB_FILL = PatternFill("solid", fgColor="FF0F2A44")
CAB_FONT = Font(color="FFFFFFFF", bold=True)
MOEDA = '"R$ "#,##0.00;"-R$ "#,##0.00;\\-'
DATA = "dd/mm/yyyy"
N = 1000  # linhas preparadas com fórmula

PAG = "Pagamentos Comissão"
COM = "Comissões"


def _cab(ws, headers, widths):
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, j, h)
        c.fill, c.font = CAB_FILL, CAB_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{N + 1}"


def preparar(wb):
    if PAG not in wb.sheetnames:
        ws = wb.create_sheet(PAG, index=wb.sheetnames.index("Log de Erros"))
        _cab(ws, ["Competência do extrato (AAAA-MM)", "Empresa Porto", "Ordem de pagamento", "Data de pagamento",
                  "Comissão bruta (R$)", "INSS (R$)", "IRRF (R$)", "ISS (R$)", "Líquido (R$)", "Débito/Crédito (R$)",
                  "A pagar (R$)", "Lançamentos no detalhe", "Soma dos lançamentos (R$)", "Conferência",
                  "Descrição débitos/créditos", "Arquivo", "Processado em", "Observações"],
             [12, 34, 12, 12, 15, 10, 10, 10, 15, 15, 15, 11, 15, 13, 50, 60, 12, 40])
        for r in range(2, N + 2):
            ws[f"L{r}"] = f"=IF(C{r}=\"\",\"\",COUNTIFS('{COM}'!$B:$B,C{r}))"
            ws[f"M{r}"] = f"=IF(C{r}=\"\",\"\",SUMIFS('{COM}'!$O:$O,'{COM}'!$B:$B,C{r}))"
            ws[f"N{r}"] = (f"=IF(C{r}=\"\",\"\",IF(AND(ABS(M{r}-E{r})<=0.01,"
                           f"ABS(I{r}+J{r}-K{r})<=0.01,ABS(E{r}-F{r}-G{r}-H{r}-I{r})<=0.01),\"OK\",\"Divergência\"))")
            for col in "LMN":
                ws[f"{col}{r}"].fill = CINZA
            for col in "EFGHIJKM":
                ws[f"{col}{r}"].number_format = MOEDA
            ws[f"D{r}"].number_format = DATA
            ws[f"Q{r}"].number_format = DATA
    if COM not in wb.sheetnames:
        ws = wb.create_sheet(COM, index=wb.sheetnames.index(PAG) + 1)
        _cab(ws, ["Competência do extrato (AAAA-MM)", "Ordem de pagamento", "Data de pagamento", "Empresa Porto",
                  "Histórico (segurado)", "Sucursal", "Ramo", "Apl/Prop", "Código estipulante", "Nº fatura (Fat/Eds)",
                  "Parcela", "Data do lançamento", "Prêmio líquido (R$)", "Taxa (%)", "Comissão (R$)", "Tipo",
                  "ID Empresa", "Competência da fatura", "Valor da fatura (R$)", "Copart + despesas da fatura (R$)",
                  "Prêmio líquido esperado (R$)", "Diferença prêmio (R$)", "Comissão calculada (R$)", "Conferência",
                  "Arquivo", "Processado em", "Observações"],
             [12, 12, 12, 30, 38, 9, 7, 12, 12, 12, 8, 12, 15, 8, 14, 26, 10, 12, 15, 15, 15, 13, 14, 18, 60, 12, 45])
        for r in range(2, N + 2):
            ws[f"Q{r}"] = (f"=IF(B{r}=\"\",\"\",IFERROR(INDEX(Faturas!$B:$B,MATCH(J{r},Faturas!$E:$E,0)),"
                           f"IFERROR(INDEX(Empresas!$A:$A,MATCH(I{r},Empresas!$E:$E,0)),\"\")))")
            ws[f"R{r}"] = f"=IF(B{r}=\"\",\"\",IFERROR(INDEX(Faturas!$A:$A,MATCH(J{r},Faturas!$E:$E,0)),\"\"))"
            ws[f"S{r}"] = f"=IF(B{r}=\"\",\"\",IFERROR(INDEX(Faturas!$J:$J,MATCH(J{r},Faturas!$E:$E,0)),\"\"))"
            ws[f"T{r}"] = (f"=IF(OR(B{r}=\"\",S{r}=\"\"),\"\",SUMIFS(Vidas!$Z:$Z,Vidas!$D:$D,J{r})"
                           f"+SUMIFS(Vidas!$Y:$Y,Vidas!$D:$D,J{r}))")
            ws[f"U{r}"] = (f"=IF(OR(B{r}=\"\",S{r}=\"\"),\"\",ROUND((S{r}-T{r})/"
                           f"IF(ISNUMBER(SEARCH(\"SEGURO SAUDE\",D{r})),1.0238,1),2))")
            ws[f"V{r}"] = f"=IF(OR(B{r}=\"\",U{r}=\"\"),\"\",M{r}-U{r})"
            ws[f"W{r}"] = f"=IF(B{r}=\"\",\"\",ROUND(M{r}*N{r}/100,2))"
            ws[f"X{r}"] = (f"=IF(B{r}=\"\",\"\",IF(LEFT(P{r},3)=\"226\",\"Bônus\",IF(S{r}=\"\",\"Sem fatura na base\","
                           f"IF(AND(ABS(V{r})<=MAX(1,0.0005*U{r}),ABS(W{r}-O{r})<=0.01),\"OK\",\"Divergência\"))))")
            for col in "QRSTUVWX":
                ws[f"{col}{r}"].fill = CINZA
            for col in "MOSTUVW":
                ws[f"{col}{r}"].number_format = MOEDA
            for col in "CLZ":
                ws[f"{col}{r}"].number_format = DATA
            ws[f"N{r}"].number_format = "0.00"
        ws["U1"].comment = Comment(
            "Prêmio líquido esperado = (valor da fatura − copart − despesas não cobertas) ÷ 1,0238 na "
            "Porto Seguro Saúde S.A. (IOF 2,38%); ÷ 1 na Portomed. Regra tirada da conferência de 09/2026.", "Agente Comissão")
        ws["X1"].comment = Comment(
            "OK = prêmio bate com o esperado (tolerância R$ 1 ou 0,05%) e comissão = prêmio × taxa. "
            "Bônus = tipo 226 (Muito + Porto Saúde). Sem fatura na base = odonto ou fatura não importada.", "Agente Comissão")
    ws = wb[COM]
    if ws["AB1"].value is None:  # v2: desconto não comissionável + comissão sobre o faturamento
        for col, h, w in (("AB", "Taxa não comissionável (R$)", 13), ("AC", "Comissão sobre o faturamento (R$)", 15),
                          ("AD", "Diferença vs faturamento (R$)", 14)):
            c = ws[f"{col}1"]
            c.value, c.fill, c.font = h, CAB_FILL, CAB_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[col].width = w
        ws["AB1"].comment = Comment(
            "Preenchido pela equipe: valor da fatura que a Porto não comissiona (ex.: taxa de R$ 15,00 na 1ª fatura). "
            "É descontado do prêmio esperado.", "Agente Comissão")
        ws["AC1"].comment = Comment(
            "Valor da fatura × taxa. A Porto paga sobre o prêmio líquido (sem IOF e copart), por isso a diferença "
            "da coluna AD é esperada e não entra na conferência.", "Agente Comissão")
        for r in range(2, N + 2):
            ws[f"U{r}"] = (f"=IF(OR(B{r}=\"\",S{r}=\"\"),\"\",ROUND((S{r}-T{r}-N(AB{r}))/"
                           f"IF(ISNUMBER(SEARCH(\"SEGURO SAUDE\",D{r})),1.0238,1),2))")
            ws[f"AC{r}"] = f"=IF(OR(B{r}=\"\",S{r}=\"\"),\"\",ROUND(S{r}*N{r}/100,2))"
            ws[f"AD{r}"] = f"=IF(OR(B{r}=\"\",AC{r}=\"\"),\"\",O{r}-AC{r})"
            ws[f"AB{r}"].font = AZUL
            for col in ("AB", "AC", "AD"):
                ws[f"{col}{r}"].number_format = MOEDA
            ws[f"AC{r}"].fill = ws[f"AD{r}"].fill = CINZA
        ws.auto_filter.ref = f"A1:AD{N + 1}"
    fat = wb["Faturas"]
    if fat["AC1"].value is None:
        ref = fat["AB1"]
        for col, h in (("AC", "Comissão recebida (R$)"), ("AD", "Status comissão")):
            c = fat[f"{col}1"]
            c.value, c.fill, c.font = h, CAB_FILL, CAB_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            fat.column_dimensions[col].width = 15
        for r in range(2, 502):
            fat[f"AC{r}"] = f"=IF(A{r}=\"\",\"\",SUMIFS('{COM}'!$O:$O,'{COM}'!$J:$J,E{r}))"
            fat[f"AD{r}"] = (f"=IF(A{r}=\"\",\"\",IF(AC{r}=0,\"Pendente\",IF(COUNTIFS('{COM}'!$J:$J,E{r},"
                             f"'{COM}'!$X:$X,\"Divergência\")>0,\"Divergência\",\"Recebida\")))")
            fat[f"AC{r}"].number_format = MOEDA
            fat[f"AC{r}"].fill = fat[f"AD{r}"].fill = CINZA
        fat.auto_filter.ref = "A1:AD501"


def _prox(ws, col="B"):
    r = 2
    while ws[f"{col}{r}"].value not in (None, ""):
        r += 1
    return r


def _w(ws, ref, v, fmt=None):
    c = ws[ref]
    c.value, c.font = v, AZUL
    if fmt:
        c.number_format = fmt

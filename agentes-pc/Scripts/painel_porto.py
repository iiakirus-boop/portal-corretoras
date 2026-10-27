# -*- coding: utf-8 -*-
"""Gera os JSON da Central para a Porto Saúde a partir da base (só lê a planilha).

Uso: python painel_porto.py [<Base_Faturas_Vidas.xlsx>] [--comp AAAA-MM ...] [--publicar] [--comissao]
Saída: _painel/fat_AAAA-MM_porto.json (aba Faturas; sem --comp, só os meses sem documento e o mais recente da base, e a
NF que já estava no documento mantém os ajustes feitos à mão) e, só com --comissao, _painel/com_AAAA-MM_porto.json (abas
Pagamentos Comissão e Comissões; o documento de comissão já existente é mais completo, então só ganha as ordens novas). Porto Seguros Gerais fica de fora (não é plano de saúde). Depois, painel_subs.py acrescenta o detalhe das subestipulantes e banco.publicar() sobe para o site.
As colunas de fórmula da base (conferências, prêmio esperado...) são calculadas aqui com as mesmas regras das fórmulas:
quando os scripts gravam a planilha, o Excel ainda não recalculou e o valor salvo da fórmula vem vazio.
faturas() e comissoes() servem também ao painel_bradesco.py.
"""
import openpyxl, json, sys, os, datetime
from collections import defaultdict

AQUI = os.path.dirname(os.path.abspath(__file__))
SAUDE = ('PORTO SEGURO - SEGURO SAUDE', 'PORTOMED')
d2 = lambda v: v.strftime('%d/%m/%Y') if isinstance(v, datetime.datetime) else ''
n = lambda v: round(float(v), 2) if isinstance(v, (int, float)) else 0.0  # fórmula (texto) ou vazio = 0


def faturas(wb, ops):
    """Aba Faturas, um item por NF das operadoras ops ({nome na base: slug}).
    Devolve {(comp, slug): [itens]}, {NF: linha da fatura} e {NF: copart + despesas das vidas}."""
    hoje = datetime.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    emp = {r[0]: r[1] for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0]}
    viv, cop = defaultdict(lambda: [0, 0.0, 0]), defaultdict(float)
    for r in wb['Vidas'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[2] in ops:
            z = sum(n(r[c]) for c in (20, 21, 22, 23, 31))          # Z = copart total
            k = viv[(r[0], r[1], str(r[3]))]
            k[0] += r[15] == 'Ativo'; k[1] += n(r[18]) + n(r[19]) + z + n(r[24]); k[2] += 1   # L e M da aba Faturas
            cop[str(r[3])] += z + n(r[24])
    fat, nfs = defaultdict(list), {}
    for r in wb['Faturas'].iter_rows(min_row=2, values_only=True):
        if r[3] not in ops or not r[0] or str(r[0]) < '2026-01':
            continue
        if not isinstance(r[6], datetime.datetime):  # sem vencimento (ex.: "SEM BOLETO", crédito de cancelamento): não é fatura a pagar e quebra a Projeção
            continue
        at, soma, linhas = viv.get((r[0], r[1], str(r[4])), [0, 0.0, 0])
        tot, venc = n(r[9]), r[6]
        nfs[str(r[4])] = r
        fat[(str(r[0]), ops[r[3]])].append(dict(eid=r[1], empresa=(emp.get(r[1]) or '')[:40], titulo=str(r[4]), venc=d2(venc),
            valor=tot, vidas=r[10] if r[10] is not None else (at if linhas else None), ativos=at if linhas else None, subs=1,
            boleto='Em aberto' if isinstance(venc, datetime.datetime) and venc >= hoje else 'Pago',
            confValor='Sem vidas' if not at else ('OK' if abs(tot - soma) <= 0.01 else 'Divergência'),
            confVidas=('Sem relação de vidas' if not linhas else '') if r[10] is None else ('OK' if at == r[10] else 'Divergência')))
    return fat, nfs, cop


def comissoes(wb, seguradora, nfs, cop, rotulo):
    """Abas Comissões e Pagamentos Comissão das seguradoras em que seguradora(nome) é verdadeiro.
    Devolve {comp: {ordens, itens}}. rotulo(nome da seguradora) = nome curto da ordem no painel."""
    itens = defaultdict(list)
    for r in wb['Comissões'].iter_rows(min_row=2, values_only=True):
        if not r[0] or not seguradora(str(r[3] or '').upper()):
            continue
        f = nfs.get(str(r[9] or ''))
        premio, taxa, com, tipo = n(r[12]), n(r[13]), n(r[14]), str(r[15] or '')
        porto = 'PORTO' in str(r[3]).upper()
        esp = None
        if f:  # coluna U: Porto tira copart e taxa fixa e o IOF de 2,38% (Seguro Saúde); Bradesco só a taxa fixa
            esp = round((n(f[9]) - cop.get(str(r[9]), 0) - n(r[27])) / (1.0238 if 'SEGURO SAUDE' in str(r[3]).upper() else 1), 2) \
                if porto else round(n(f[9]) - n(r[27]), 2)
        ok = f and abs(premio - esp) <= max(1, 0.0005 * esp) and abs(round(premio * taxa / 100, 2) - com) <= (0.01 if porto else max(0.05, 0.0001 * com))
        itens[str(r[1])].append(dict(eid=f[1] if f else '', empresa=str(r[4] or '')[:60], op=str(r[1]), titulo=str(r[9] or ''),
            compFat=str(f[0]) if f else '', venc=d2(f[6]) if f else '', premio=premio,
            valorFat=None if not f else n(f[9]) if porto else round(n(f[9]) + n(f[30]), 2),  # Bradesco: boleto = fatura + IOF
            premioEsp=esp, taxa=taxa, comissao=com, parcela=r[10],
            tipo=tipo[:2] if tipo[:2] in ('AN', 'CO') else ('AN' if taxa > 50 else 'CO'),
            copart=round(cop.get(str(r[9]), 0), 2) if porto and f else 0, taxaFixa=n(r[27]),
            conf='Bônus' if porto and tipo.startswith('226') else 'Sem fatura na base' if not f else ('OK' if ok else 'Divergência')))
        if porto:
            itens[str(r[1])][-1]['comFat'] = round(n(f[9]) * taxa / 100, 2) if f else None
    com = defaultdict(lambda: dict(ordens=[], itens=[]))
    for r in wb['Pagamentos Comissão'].iter_rows(min_row=2, values_only=True):
        if not r[0] or str(r[0]) < '2026-01' or not seguradora(str(r[1] or '').upper()):
            continue
        op, c = str(r[2]), com[str(r[0])]
        it = itens.get(op, [])
        ok = abs(sum(i['comissao'] for i in it) - n(r[4])) <= 0.01
        if r[8] is not None:  # extrato com tributação (Porto): líquido + déb/créd = a pagar e bruto - tributos = líquido
            ok = ok and abs(n(r[8]) + n(r[9]) - n(r[10])) <= 0.01 and abs(n(r[4]) - n(r[5]) - n(r[6]) - n(r[7]) - n(r[8])) <= 0.01
        c['ordens'].append(dict(op=op, data=d2(r[3]), dataISO=r[3].strftime('%Y-%m-%d') if r[3] else '', empresa=rotulo(str(r[1])),
            bruto=n(r[4]), debcred=n(r[9]), apagar=n(r[10]) if r[10] is not None else n(r[4]), conf='OK' if ok else 'Divergência',
            lanc=len(it), debitos=r[14] or ''))
        c['itens'] += it
    return com


def gravar_comissoes(com, out, slug, operadora, faturas_mes, obs, agora):
    """Grava só os meses com ordens novas (juntar). Devolve os nomes dos arquivos gravados."""
    arqs = []
    for comp, c in sorted(com.items()):
        p = os.path.join(out, f'com_{comp}_{slug}.json')
        doc = juntar(dict(comp=comp, operadora=operadora, atualizadoEm=agora, ordens=c['ordens'], itens=c['itens'],
                          extras=[], pendentes=[], faturasMes=faturas_mes(comp), obs=obs), p)
        if doc:
            json.dump(doc, open(p, 'w', encoding='utf-8'), ensure_ascii=False); arqs.append(os.path.basename(p))
    return arqs


def ler(p):
    d = json.load(open(p, encoding='utf-8'))
    return d.get('data', d) if set(d) == {'data'} else d


def escolher(fat, out, comps=None):
    """Competências de faturas a gerar: as pedidas, ou as que ainda não têm documento e a mais recente da base.
    shortcut: um mês antigo corrigido na base só volta para a Central com --comp (os publicados têm ajustes à mão)."""
    if comps:
        return set(comps)
    ult = max(c for c, _ in fat)
    return {c for c, s in fat if c == ult or not os.path.exists(os.path.join(out, f'fat_{c}_{s}.json'))}


def preservar(itens, caminho):
    """NF que já está no documento: mantém o que foi ajustado à mão (nome, obs, verificado...; com obs, também as
    conferências). O detalhe das subs (subsDet) é refeito pelo painel_subs.py."""
    if not os.path.exists(caminho):
        return itens
    velho = {i['titulo']: i for i in ler(caminho)['itens']}
    for i in itens:
        for k, x in velho.get(i['titulo'], {}).items():
            if k != 'subsDet' and (k not in i or k == 'empresa' or (k in ('confValor', 'confVidas') and 'obs' in velho[i['titulo']])):
                i[k] = x
    return itens


def gravar_faturas(fat, comps, out, nome, agora):
    arqs = []
    for (comp, slug), it in sorted(fat.items()):
        if comp in comps:
            p = os.path.join(out, f'fat_{comp}_{slug}.json')
            json.dump(dict(comp=comp, operadora=nome(slug), atualizadoEm=agora, parcial=False, pendentes=[], itens=preservar(it, p)),
                      open(p, 'w', encoding='utf-8'), ensure_ascii=False)
            arqs.append(os.path.basename(p))
    return arqs


def gerar(base, out, comissao=False, comps=None):
    wb = openpyxl.load_workbook(base, read_only=True)
    agora = datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    fat, nfs, cop = faturas(wb, {'Porto Saúde': 'porto'})
    os.makedirs(out, exist_ok=True)
    comps = escolher(fat, out, comps)
    fats = gravar_faturas(fat, comps, out, lambda s: 'Porto Saúde', agora)
    coms = []
    if comissao:
        com = comissoes(wb, lambda s: s.startswith(SAUDE), nfs, cop,
                        lambda s: 'Porto · ' + ('Portomed' if 'PORTOMED' in s.upper() else 'Saúde'))
        coms = gravar_comissoes(com, out, 'porto', 'Porto Saúde', lambda c: len(fat.get((c, 'porto'), [])),
                                'Porto Saúde: extratos da Porto Seguro Saúde e da Portomed. Agenciamento = taxa acima de 50%; vitalício = até 10%.', agora)
    print(f'Porto: faturas de {", ".join(sorted(comps))}; {len(coms)} documento(s) de comissão com ordens novas.')
    return fats + coms


def juntar(doc, caminho):
    """Comissão: o documento já publicado (com extras, pendentes e obs feitos à mão) fica como está; só entram as ordens
    que ele ainda não tem, com os lançamentos delas. Devolve None quando não há nada novo (o arquivo não é regravado)."""
    if not os.path.exists(caminho):
        return doc
    atual = json.load(open(caminho, encoding='utf-8'))
    atual = atual.get('data', atual) if set(atual) == {'data'} else atual
    novas = {o['op'] for o in doc['ordens']} - {o['op'] for o in atual['ordens']}
    if not novas:
        return None
    atual['ordens'] += [o for o in doc['ordens'] if o['op'] in novas]
    atual['itens'] += [i for i in doc['itens'] if i['op'] in novas]
    atual['atualizadoEm'] = doc['atualizadoEm']
    return atual


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if a not in ('--publicar', '--comissao')]
    comps = set()
    while '--comp' in args:
        i = args.index('--comp'); comps.add(args[i + 1]); del args[i:i + 2]
    base = args[0] if args else os.path.join(os.path.dirname(AQUI), 'Base_Faturas_Vidas.xlsx')
    arqs = gerar(base, os.path.join(AQUI, '_painel'), '--comissao' in sys.argv, comps)
    for f in arqs:
        print('  _painel/' + f)
    if '--publicar' in sys.argv:
        import banco
        banco.publicar(arqs)

"""Agente Projeção — grava a rodada na Base_Faturas_Vidas.xlsx.
O cálculo é feito pelo motor (motor_projecao.js), o mesmo modelo da página Projeção da Central.
Uso:
  python3 projecao.py ler   <base.xlsx> <saida_implantacoes.json>      # lê a aba Implantações
  python3 projecao.py gravar <base.xlsx> <rodada.json> <implantacoes.json> <inicio ISO> <quem>
"""
import sys, json, shutil, os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

IMP_COLS = ['ID', 'Empresa', 'ID Empresa (se já estiver na base)', 'Seguradora', 'Competência da 1ª fatura (AAAA-MM)', 'Dia de vencimento',
            'Valor mensal estimado (R$)', 'Vidas', 'Agenciamento (%)', 'Vitalício (%)', 'Situação', 'Observações', 'Origem', 'Cadastrada em', 'Atualizada em']
IMP_KEYS = ['id', 'empresa', 'eid', 'operadora', 'inicio', 'diaVenc', 'valor', 'vidas', 'agPct', 'vitPct', 'status', 'obs', 'origem', 'criadaEm', 'atualizadoEm']
PRJ_COLS = ['Rodada', 'Data da rodada', 'Gatilho', 'Mês projetado (AAAA-MM)', 'Seguradora', 'Tipo', 'Já recebido (R$)', 'A receber programado (R$)',
            'Projetado (R$)', 'Total previsto (R$)', 'Realizado no mês (R$)', 'Diferença (R$)', 'Observações']
HEAD = dict(font=Font(bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='0F2A44'))
BLUE = Font(color='1F4E9E')

def header(ws, cols, widths):
    for i, c in enumerate(cols, 1):
        x = ws.cell(1, i, c); x.font = HEAD['font']; x.fill = HEAD['fill']; x.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 32

def aba_imp(wb):
    if 'Implantações' in wb.sheetnames: return wb['Implantações']
    ws = wb.create_sheet('Implantações', index=wb.sheetnames.index('Comissões') + 1)
    header(ws, IMP_COLS, [20, 38, 16, 16, 16, 11, 16, 8, 13, 11, 13, 40, 11, 18, 18])
    dv = DataValidation(type='list', formula1='"Prevista,Implantada,Cancelada"', allow_blank=True); ws.add_data_validation(dv); dv.add('K2:K2000')
    dv2 = DataValidation(type='list', formula1='"Porto Saúde,Bradesco Saúde,Bradesco Dental,Amil,Amil Dental,Hapvida,SulAmérica,Santa Helena"', allow_blank=True); ws.add_data_validation(dv2); dv2.add('D2:D2000')
    return ws

def aba_prj(wb):
    if 'Projeção' in wb.sheetnames: return wb['Projeção']
    ws = wb.create_sheet('Projeção', index=wb.sheetnames.index('Implantações') + 1)
    header(ws, PRJ_COLS, [24, 17, 22, 14, 13, 14, 15, 15, 15, 15, 15, 14, 50])
    return ws

def ler_imps(ws):
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[1]: continue
        d = dict(zip(IMP_KEYS, r))
        for k in ('diaVenc', 'valor', 'vidas', 'agPct', 'vitPct'):
            try: d[k] = float(d[k]) if d[k] not in (None, '') else 0
            except Exception: d[k] = 0
        d['inicio'] = d['inicio'].strftime('%Y-%m') if hasattr(d['inicio'], 'strftime') else str(d['inicio'] or '')[:7]
        for k in ('id', 'empresa', 'eid', 'operadora', 'status', 'obs', 'origem', 'criadaEm', 'atualizadoEm'):
            v = d.get(k); d[k] = v.isoformat() if hasattr(v, 'isoformat') else ('' if v is None else str(v))
        d['status'] = d['status'] or 'Prevista'
        out.append(d)
    return out

if sys.argv[1] == 'ler':
    wb = openpyxl.load_workbook(sys.argv[2], read_only=True, data_only=True)
    imps = ler_imps(wb['Implantações']) if 'Implantações' in wb.sheetnames else []
    json.dump(imps, open(sys.argv[3], 'w', encoding='utf-8'), ensure_ascii=False); print(len(imps), 'implantações na planilha')
    sys.exit()

base, rodada, impf, inicio, quem = sys.argv[2:7]
R = json.load(open(rodada, encoding='utf-8')); foto, realizado = R['foto'], R['realizado']
imps = json.load(open(impf, encoding='utf-8'))
hoje = datetime.now()
bk = os.path.join(os.path.dirname(base), 'Backups', f'Base_Faturas_Vidas antes Projecao {hoje:%Y-%m-%d}.xlsx')
if not os.path.exists(bk): shutil.copy2(base, bk)
wb = openpyxl.load_workbook(base)
wi, wp = aba_imp(wb), aba_prj(wb)
# implantações: a lista combinada (planilha + Central) reescreve a aba, uma linha por ID
ex = {}
for i, r in enumerate(wi.iter_rows(min_row=2, values_only=True), 2):
    if r[0]: ex[str(r[0])] = i
nxt = wi.max_row + 1 if wi.max_row > 1 else 2
for d in imps:
    row = ex.get(d['id'])
    if not row: row = nxt; nxt += 1
    for c, k in enumerate(IMP_KEYS, 1):
        v = d.get(k, '')
        if k == 'inicio' and v: v = str(v)[:7]
        wi.cell(row, c, v if v != '' else None)
    for c in (7,): wi.cell(row, c).number_format = '#,##0.00'
# rodada: uma linha por mês × seguradora × tipo
tipo = {'ag': 'Agenciamento', 'vit': 'Vitalício'}
dt = datetime.fromisoformat(foto['geradoEm'].replace('Z', '+00:00')).astimezone().replace(tzinfo=None)
r0 = wp.max_row + 1 if wp.max_row > 1 else 2
for l in foto['linhas']:
    wp.append([foto['id'], dt, foto['gatilho'], l['comp'], l['seg'], tipo[l['kind']], l['real'], l['prev'], l['proj'], l['total'], None, None,
               'mês atual: inclui o que já caiu' if l['comp'] == foto['mesAtual'] else ''])
# realizado: atualiza todas as linhas de meses que já fecharam (ou o mês atual, parcial)
real_kind = R.get('realizadoTipo', {})
for row in wp.iter_rows(min_row=2):
    m, seg, tp = row[3].value, row[4].value, row[5].value
    if not m or m > foto['mesAtual']: continue
    k = 'ag' if tp == 'Agenciamento' else 'vit'
    v = (real_kind.get(m, {}).get(seg, {}) or {}).get(k)
    if v is None: v = 0.0
    row[10].value = v; row[11].value = round(v - (row[9].value or 0), 2) if m < foto['mesAtual'] else None
    if m == foto['mesAtual'] and not row[12].value: row[12].value = 'realizado parcial (mês em andamento)'
for row in wp.iter_rows(min_row=r0):
    row[1].number_format = 'dd/mm/yyyy hh:mm'
    for c in range(6, 12): row[c].number_format = '#,##0.00'
# próxima semana detalhada: a aba é refeita a cada rodada
PS = foto.get('proximaSemana')
if PS:
    if 'Próxima semana' not in wb.sheetnames:
        ws = wb.create_sheet('Próxima semana', index=wb.sheetnames.index('Projeção') + 1)
    else:
        ws = wb['Próxima semana']; ws.delete_rows(1, ws.max_row)
    ws.append([f"Próxima semana: {PS['de'][8:10]}/{PS['de'][5:7]} a {PS['ate'][8:10]}/{PS['ate'][5:7]} · total previsto R$ {PS['total']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') + f" · rodada {dt:%d/%m/%Y %H:%M}"])
    ws['A1'].font = Font(bold=True, size=12)
    cols = ['Data', 'Empresa', 'ID Empresa', 'Seguradora', 'Tipo', 'Situação', 'Valor (R$)', 'Como foi calculado', 'Data prevista']
    for i, c in enumerate(cols, 1):
        x = ws.cell(2, i, c); x.font = HEAD['font']; x.fill = HEAD['fill']
    for i, w in enumerate([11, 42, 10, 12, 13, 12, 13, 70, 50], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for it in PS['itens']:
        ws.append([datetime.fromisoformat(it['data']), it['empresa'], it['eid'], it['seg'], tipo[it['kind']], it['situacao'], it['valor'], it['como'], it['quando']])
        ws.cell(ws.max_row, 1).number_format = 'dd/mm/yyyy'; ws.cell(ws.max_row, 7).number_format = '#,##0.00'
        if it['situacao'] == 'Em atraso': ws.cell(ws.max_row, 6).font = Font(bold=True, color='8A4100')
    ws.freeze_panes = 'A3'
# Agentes + Execuções
AG = wb['Agentes']
lin = next((i for i in range(2, AG.max_row + 1) if AG.cell(i, 1).value == 'Projeção'), None)
if not lin: lin = next(i for i in range(2, AG.max_row + 2) if not AG.cell(i, 1).value)
for c, v in enumerate(['Projeção', 'Projeta a comissão dos próximos 3 meses (agenciamento × vitalício, por seguradora) e guarda cada rodada para comparar com o realizado',
                       'Todas', 'Implantações, Projeção, Próxima semana', 'Depois de cada agente', None, 'Fichas/Projeção.md', 'Ativo'], 1):
    AG.cell(lin, c, v)
EX = wb['Execuções']
n = len(foto['linhas']); nimp = len([d for d in imps if d.get('status') != 'Cancelada'])
tot = {m['comp']: m['total'] for m in foto['meses']}
obs = 'Previsto: ' + '; '.join(f"{c}: R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') for c, v in tot.items()) + f". {nimp} implantações ativas."
EX.append(['Projeção', foto['mesAtual'], datetime.fromisoformat(inicio), datetime.now().replace(microsecond=0), '✅ OK', n, n, 0, 0,
           R.get('pendencias', ''), os.path.relpath(bk, os.path.dirname(base)).replace('\\', '/'), obs, quem])
wb.save(base)
print('gravado:', n, 'linhas de projeção,', len(imps), 'implantações; backup', os.path.basename(bk))

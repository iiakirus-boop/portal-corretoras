# -*- coding: utf-8 -*-
"""Agente Faturas Porto Saúde - leitura.
Lê a pasta Faturas/AAAA-MM/PortoSaude (uma subpasta "<sub> - <NOME CURTO>" por empresa, como o baixar_porto.js deixa)
e gera um JSON por fatura (título) para um gravar_porto.py: dados da fatura (PDF FATURADOS), boleto (PDF BOLETO) e
vidas (TXT do .zip de Exportação Arquivos). Só lê: não abre a base para escrita nem grava nada fora do JSON de saída.
Uso: python faturas_porto.py <pasta PortoSaude> <saida.json> [--base <Base_Faturas_Vidas.xlsx>]
  --base: só leitura, para achar o ID Empresa pela aba Subestipulantes/Empresas (padrão: a base da pasta FATURAS).
Regras (Agente Faturas — Porto Saúde.md): o TXT vem em centavos; linhas 02 = um beneficiário, 03 = total da família
(ignoradas); prêmio base = VALOR PREMIO - pró-rata - coparts - despesas; título sem zeros à esquerda; unidade 02/03
de uma subestipulante com título próprio leva " U02" no nome dos arquivos. TXT "PARCIAL" (exportação incompleta) é ignorado.
"""
import os, re, sys, json, zipfile, subprocess, tempfile
from collections import Counter

NUM = r'[\d.]+,\d\d'
AQUI = os.path.dirname(os.path.abspath(__file__))
ARQ = re.compile(r'^(\d+) - (.+?) (FATURADOS|BOLETO|VIDAS) PORTO (SAUDE|ODONTO) (\d\d)-(\d\d)( U\d\d)?( PARCIAL [\d-]+)?\.(pdf|txt|zip)$', re.I)
PAREN = {'Conjuge': 'Cônjuge', 'CONJUGE': 'Cônjuge'}


def n(s):
    return float(s.replace('.', '').replace(',', '.')) if s else None


def iso(d):
    return f'{d[6:10]}-{d[3:5]}-{d[0:2]}' if d and re.match(r'\d\d/\d\d/\d{4}$', d) else None


def pdf_txt(p, modo):
    # o Imprimir do portal às vezes vem com o HTML da página colado depois do %%EOF: lê só o PDF
    b = open(p, 'rb').read()
    h = b.find(b'<!DOCTYPE')
    e = b.rfind(b'%%EOF', 0, h) if h >= 0 else -1
    if e >= 0:
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as t:
            t.write(b[:e + 5])
        p = t.name
    try:
        return subprocess.run(['pdftotext', modo, '-enc', 'UTF-8', p, '-'], capture_output=True).stdout.decode('utf-8', 'replace').replace('\r\n', '\n')
    finally:
        if e >= 0:
            os.remove(p)


def ler_fatura(p):
    t = pdf_txt(p, '-raw')
    g = lambda r: (re.search(r, t) or [None, None])[1]
    if not g(r'T[íi]tulo:\s*(\d+)'):
        return dict(tipo='ignorado', arquivo=p, motivo='PDF sem "Título:" (não é o Resumo da Fatura da Porto)')
    d = dict(tipo='fatura', arquivo=p)
    m = re.search(r'Referencia:\s*(\d\d)/(\d{4})', t); d['ref'] = f'{m.group(2)}-{m.group(1)}' if m else None
    d['venc'] = iso(g(r'Vencimento:\s*(\d\d/\d\d/\d{4})'))
    m = re.search(r'Vig\. Inicial:\s*(\d\d/\d\d/\d{4})\s*Vig\. Final:\s*(\d\d/\d\d/\d{4})', t)
    d['cob'] = (iso(m.group(1)), iso(m.group(2))) if m else (None, None)
    m = re.search(r'Data emiss[ãa]o fatura:\s*(\d{1,2})/(\d{1,2})/(\d{4})', t)  # vem em M/D/AAAA
    d['emissao'] = f'{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}' if m else None
    d['estipulante'] = g(r'Estipulante:\s*\n?(\d+)')
    d['sub'] = g(r'Subestipulante:\s*\n?(\d+)')
    u = g(r'Unidade:\s*\n?(\d+)'); d['unidade'] = f'{int(u):02d}' if u else None
    d['apolice'], d['titulo'] = g(r'Ap[óo]lice:\s*(\d+)'), str(int(g(r'T[íi]tulo:\s*(\d+)')))
    m = re.search(r'Saldo Atual\s+(\d+)\s+R\$\s*(' + NUM + r')\s+(\d+)\s+R\$\s*(' + NUM + r')\s+(\d+)\s+R\$\s*(' + NUM + ')', t)
    if m:
        d.update(titulares=int(m.group(1)), dependentes=int(m.group(3)), vidas=int(m.group(5)), valor_saldo=n(m.group(6)))
    for k, r in (('premio_real', r'Pr[êe]mio Real\(=\)'), ('copart', r'Co-Participa[çc][ãa]o\(\+\)'),
                 ('desc_copart', r'Desc\.Co-Participa[çc][ãa]o\(-\)'), ('desp_nao_cobertas', r'Desp\.N[ãa]o Cobertas \(\+\)')):
        m = re.search(r + r'\s+R\$\s*' + NUM + r'\s+R\$\s*' + NUM + r'\s+R\$\s*(' + NUM + ')', t)
        d[k] = n(m.group(1)) if m else None
    for k, r in (('premio_liquido', 'Pr[êe]mio L[íi]quido'), ('prorata', 'Pr[êe]mio Pr[óo]-rata'), ('taxa_inscricao', 'Taxa de Inscri[çc][ãa]o'),
                 ('bonus', 'B[ôo]nus'), ('iof', r'I\.O\.F'), ('total', 'Total')):
        m = re.search(r'(?m)^' + r + r'\s+R\$\s*(' + NUM + ')', t)
        d[k] = n(m.group(1)) if m else None
    d['layout'] = 'empresarial' if 'RELAÇÃO ATUALIZADA DE SEGURADOS' in t else 'PME'
    return d


def ler_boleto(p):
    t = pdf_txt(p, '-layout')
    d = dict(tipo='boleto', arquivo=p)
    m = re.search(r'(\d{5}\.\d{5} \d{5}\.\d{6} \d{5}\.\d{6} \d \d{14})', t); d['linha'] = m.group(1) if m else None
    m = re.search(r'Valor do Documento[^\n]*\n(?:[ \t]*\n)*[^\n]*?(' + NUM + r')[ \t]*$', t, re.M); d['valor'] = n(m.group(1).replace('.', '')) if m else None
    m = re.search(r'Data de Vencimento\s*\n[^\n]*?(\d\d/\d\d/\d{4})', t); d['venc'] = iso(m.group(1)) if m else None
    c = Counter(re.findall(r'\b(\d{13})\b', t)); d['nosso'] = c.most_common(1)[0][0] if c else None
    m = re.search(r'Sacado\s*\n\s*\n?\s*0?(\d{14})\s+(.+)', t)
    if m:
        c = m.group(1); d['cnpj'] = f'{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}'; d['sacado'] = m.group(2).strip()
    if d['linha'] and d['valor'] is not None and int(d['linha'][-10:]) != round(d['valor'] * 100):
        d['aviso'] = f'valor do boleto {d["valor"]} não bate com a linha digitável'
    return d


def ler_txt(dados, nome):
    """Linhas 02 do TXT da Porto (campos separados por |, valores em centavos)."""
    linhas = dados.decode('latin-1').splitlines()
    cab = linhas[0].rstrip('|').split('|') if linhas else []
    out = []
    for l in linhas[1:]:
        c = l.split('|')
        if c[0] != '02':
            continue
        r = dict(zip(cab, c))
        cent = lambda k: int(r.get(k) or 0) / 100
        cop = dict(consultas=cent('VALOR CONSULTA'), exames=cent('VALOR EXAME'), psocorro=cent('VALOR PSOCORRO'), internacao=cent('VALOR INTERNACAO'))
        desp, pro, tot = cent('DESPESAS NAO COBERTAS'), cent('VALOR PRORATA'), cent('VALOR PREMIO')
        par = r.get('PARENTESCO', '').strip()
        out.append(dict(seguro=r.get('SEGURO'), dep=int(r['NR DEPENDENTE']) if (r.get('NR DEPENDENTE') or '').isdigit() else None,
                        nome=r.get('NOME SEGURADO', '').strip(), cpf=(r.get('CPF') or '').strip() or None,
                        idade=int(r['IDADE']) if (r.get('IDADE') or '').isdigit() else None,
                        parentesco=PAREN.get(par, par), titular=par.lower() == 'titular',
                        cod_plano=r.get('CODIGO PLANO'), plano=(r.get('PLANO') or '').strip(), matricula=(r.get('MATRICULA') or '').strip() or None,
                        situacao={'A': 'Ativo', 'E': 'Excluído'}.get(r.get('SITUACAO'), r.get('SITUACAO')),
                        adesao=iso(r.get('DATA ADESAO')), exclusao=iso(r.get('DATA EXCLUSAO')),
                        valor_premio=tot, prorata=pro, copart=cop, despesas=desp,
                        premio_base=round(tot - pro - sum(cop.values()) - desp, 2), txt=nome))
    return out


def ids_da_base(base):
    """sub -> ID Empresa e estipulante -> ID Empresa (abas Subestipulantes e Empresas), só leitura."""
    if not base or not os.path.exists(base):
        return {}, {}
    import openpyxl
    wb = openpyxl.load_workbook(base, read_only=True, data_only=True)
    sub, est = {}, {}
    for r in wb['Empresas'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[3] == 'Porto Saúde':
            if r[5]: sub.setdefault(str(r[5]), r[0])
            if r[4] and r[8] == 'Sim': est.setdefault(str(r[4]), r[0])
    for r in wb['Subestipulantes'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[2]: sub.setdefault(str(r[2]), r[0])
    wb.close()
    return sub, est


def main(root, out, base):
    subs, ests = ids_da_base(base)
    fats, docs = [], []
    for pasta in sorted(os.listdir(root)):
        dp = os.path.join(root, pasta)
        if not os.path.isdir(dp):
            continue
        arqs = {}
        for f in sorted(os.listdir(dp)):
            m = ARQ.match(f)
            if not m:
                docs.append(dict(tipo='ignorado', arquivo=os.path.join(dp, f), motivo='nome fora do padrão <sub> - <NOME> <TIPO> PORTO SAUDE MM-AA'))
            elif m.group(8):
                docs.append(dict(tipo='ignorado', arquivo=os.path.join(dp, f), motivo='exportação PARCIAL (substituída pela completa)'))
            else:
                arqs[(m.group(1), m.group(3).upper(), (m.group(7) or '').strip(), m.group(9).lower())] = (os.path.join(dp, f), m)
        for (sub, tp, un, ext), (p, m) in sorted(arqs.items()):
            if tp != 'FATURADOS':
                continue
            try:
                d = ler_fatura(p)
            except Exception as e:
                d = dict(tipo='erro', arquivo=p, erro=repr(e))
            if d['tipo'] != 'fatura':
                docs.append(d); continue
            d.update(pasta=pasta, nome_curto=m.group(2), produto='Porto Odonto' if m.group(4).upper() == 'ODONTO' else 'Porto Saúde',
                     comp=f'20{m.group(6)}-{m.group(5)}', sufixo=un, avisos=[])
            av = d['avisos']
            if d['ref'] != d['comp']: av.append(f'PDF é da competência {d["ref"]}, pasta é {d["comp"]}')
            if d['sub'] != sub: av.append(f'PDF é da sub {d["sub"]}, arquivo diz {sub}')
            b = arqs.get((sub, 'BOLETO', un, 'pdf'))
            d['boleto'] = ler_boleto(b[0]) if b else None
            if d['boleto']:
                bb = d['boleto']
                if bb.get('aviso'): av.append(bb['aviso'])
                if bb.get('valor') is not None and d.get('total') is not None and abs(bb['valor'] - d['total']) > 0.01:
                    av.append(f'boleto R$ {bb["valor"]:.2f} ≠ total da fatura R$ {d["total"]:.2f}')
                if bb.get('venc') and bb['venc'] != d['venc']: av.append(f'vencimento do boleto {bb["venc"]} ≠ fatura {d["venc"]}')
            # vidas: o .zip traz o título no nome de cada TXT (est_sub_uni_comp_data_título_seq.txt)
            z = arqs.get((sub, 'VIDAS', un, 'zip'))
            d['arquivo_vidas'], d['txts'], d['vidas_linhas'] = None, [], []
            if z:
                d['arquivo_vidas'] = z[0]
                with zipfile.ZipFile(z[0]) as zf:
                    for e in sorted(zf.namelist()):
                        k = os.path.basename(e).split('_')
                        if not e.lower().endswith('.txt') or len(k) < 7:
                            continue
                        if str(int(k[5])) != d['titulo']:
                            av.append(f'{e} é do título {int(k[5])}, não do {d["titulo"]}'); continue
                        lv = ler_txt(zf.read(e), e)
                        for v in lv: v.update(sub=str(int(k[1])), unidade=k[2])
                        d['txts'].append(dict(arquivo=e, sub=str(int(k[1])), unidade=k[2], linhas=len(lv)))
                        d['vidas_linhas'] += lv
            else:
                # sem .zip: usa os TXT extraídos da pasta com o mesmo sufixo
                for (s2, t2, u2, e2), (p2, m2) in sorted(arqs.items()):
                    if t2 == 'VIDAS' and e2 == 'txt' and u2 == un:
                        lv = ler_txt(open(p2, 'rb').read(), p2)
                        for v in lv: v.update(sub=s2, unidade=(u2[1:] if u2 else d['unidade']))
                        d['txts'].append(dict(arquivo=p2, sub=s2, unidade=u2[1:] if u2 else d['unidade'], linhas=len(lv)))
                        d['vidas_linhas'] += lv
                if len(d['txts']) == 1: d['arquivo_vidas'] = d['txts'][0]['arquivo']
            if not d['txts']: av.append('sem TXT de vidas')
            ativos = [v for v in d['vidas_linhas'] if v['situacao'] == 'Ativo']
            d['vidas_ativas'] = len(ativos)
            d['soma_vidas'] = round(sum(v['valor_premio'] for v in d['vidas_linhas']), 2)
            if d['txts'] and d.get('vidas') is not None and d['vidas_ativas'] != d['vidas']:
                av.append(f'vidas ativas no TXT {d["vidas_ativas"]} ≠ vidas na fatura {d["vidas"]}')
            # o VALOR PREMIO do TXT já vem com pró-rata, coparts e IOF: a soma tem de dar o total da fatura
            if d['txts'] and d.get('total') is not None and abs(d['soma_vidas'] - d['total']) > 0.01:
                av.append(f'soma do TXT R$ {d["soma_vidas"]:.2f} ≠ total R$ {d["total"]:.2f} (dif. {d["soma_vidas"] - d["total"]:.2f})')
            d['id_empresa'] = subs.get(sub) or ests.get(str(int(d['estipulante'] or 0)))
            if not d['id_empresa']: av.append(f'sub {sub} / estipulante {d["estipulante"]} sem cadastro na base')
            for s2 in {t['sub'] for t in d['txts']} - {sub}:
                if s2 not in subs: av.append(f'sub {s2} (no TXT) sem cadastro na aba Subestipulantes')
            fats.append(d)
        if not any(k[1] == 'FATURADOS' for k in arqs):
            docs.append(dict(tipo='erro', arquivo=dp, erro='pasta sem PDF FATURADOS: nada lido'))
    json.dump(fats + docs, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    rs = lambda v: f'R$ {v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.') if v is not None else '-'
    print(f'{len(fats)} faturas, {sum(f["vidas"] or 0 for f in fats)} vidas, {rs(sum(f["total"] or 0 for f in fats))}')
    for f in fats:
        print(f'  {f["pasta"]}{" " + f["sufixo"] if f["sufixo"] else ""}: título {f["titulo"]}, venc {f["venc"]}, {f["vidas"]} vidas '
              f'(TXT {f["vidas_ativas"]}), {rs(f["total"])}, boleto {rs(f["boleto"]["valor"]) if f["boleto"] else "-"}, {f["id_empresa"] or "SEM ID"}')
        for a in f['avisos']:
            print('    aviso:', a)
    for d in docs:
        print(f'  {d["tipo"]}: {os.path.relpath(d["arquivo"], root)} ({d.get("motivo") or d.get("erro")})')


if __name__ == '__main__':
    a = sys.argv[1:]
    base = os.path.join(os.path.dirname(AQUI), 'Base_Faturas_Vidas.xlsx')
    if '--base' in a:
        i = a.index('--base'); base = a[i + 1]; del a[i:i + 2]
    main(a[0], a[1], base)

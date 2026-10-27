# -*- coding: utf-8 -*-
"""Agente Faturas Hapvida - gravação.
Grava na Base_Faturas_Vidas.xlsx o JSON gerado por faturas_hapvida.py (pasta da equipe + coleta do Portal Empresa), copia os
arquivos para Faturas/AAAA-MM/Hapvida/ e gera os JSON do painel (Scripts/_painel/fat_AAAA-MM_hapvida.json).
Uso: python3 gravar_hapvida.py <hap.json> <planilha entrada> <planilha saida> <pasta FATURAS> <AAAA-MM-DD> [--sem-copia] [--simular]
     --simular = não salva a planilha, não copia arquivos e não gera o painel (só mostra o que faria).
Regras:
- Os planos são da NotreDame Intermédica (grupo Hapvida); operadora na base = "Hapvida".
- Uma fatura por obrigação (= "Nº fatura (Fat/Eds)" do extrato de comissão Hapvida). Competência = mês do vencimento.
- Uma empresa (aba Empresas) por contrato que recebe o boleto (ex.: 1V38A, 1Z1FK). Código estipulante = código do contrato,
  igual ao "Código estipulante" da aba Comissões, para a comissão casar sozinha. Os outros códigos que aparecem no analítico
  (filiais, odonto, ex.: 0S85D, 0XE6G) vão para Subestipulantes.
- Contrato da fatura: portal > boleto (CÓDIGO XXXXX.000000) > NF-e ("CONTRATO NXXXXX") > código que já está na aba Comissões >
  código do nome do arquivo > primeira empresa do analítico.
- Valor total = "Valor" do analítico (já com descontos/acréscimos); sem analítico, valor do portal ou do boleto.
  Vidas: Relatório.csv do portal > CSV da pasta > PDF analítico. Prêmio da vida = "Cobrado".
  Descontos e acréscimos (pró-rata de cancelamento, mínimo contratual, reajuste retroativo) não são por vida: ficam nas
  Observações e a Conferência valor mostra a diferença.
- Copart (fatura de co-participação e boleto COPART) não vira fatura de mensalidade: o total vai nas Observações da fatura
  do mesmo contrato e mês. Planos individuais (3010V, boleto "PAGAR ESTA FATURA") entram como fatura própria da pessoa.
- Fatura já gravada (mesma competência + empresa + nº fatura) não é gravada de novo: só compara valor e relata.
- Nunca apaga nada; arquivos são copiados (não movidos) e não sobrescreve cópia que já exista.
"""
import json, sys, os, re, shutil
from datetime import datetime
from collections import defaultdict, OrderedDict, Counter
import openpyxl
from shaddai_comum import d, br, dbr, cs, last_row, fim_mes, mmaa, estender, NEMP, SUBFORM  # módulo comum

OP = 'Hapvida'
AGENTE = 'Faturas Hapvida'
PASTA = 'Hapvida'
PROD = 'HAPVIDA'
PAREN = {'TITULAR': 'Titular', 'CONJUGE': 'Cônjuge', 'CÔNJUGE': 'Cônjuge', 'COMPANHEIRO(A)': 'Companheiro(a)', 'FILHO(A)': 'Filho(a)',
         'ENTEADO(A)': 'Enteado(a)', 'AGREGADO': 'Agregado', 'DEPENDENTE': 'Dependente'}


def nome_curto_arquivo(rel):
    b = os.path.basename(rel)
    m = re.match(r'^[0-9A-Z]{5}\s*-\s*(.*?)\s+(?:FATURADOS|FARUTADOS|BOLETO|NF-?E|HAPVIDA)', b, re.I)
    if m and m.group(1).strip(' -'):
        return re.sub(r'\s+', ' ', m.group(1).strip(' -')).upper()
    p = rel.split('/')
    return re.sub(r'\s+', ' ', p[1]).upper() if len(p) > 2 else None


def codigos_comissao(wb):
    return {str(r[8]) for r in wb['Comissões'].iter_rows(min_row=2, values_only=True) if r[3] and 'HAP' in str(r[3]).upper() and r[8]}


def preparar(docs, log, cods_com):
    for x in docs:
        if x['tipo'] in ('ignorado', 'erro'):
            log.append((x.get('mes_pasta'), None, 'Extração', f'Arquivo não lido: {x["rel"]} ({x.get("motivo") or x.get("erro")}).',
                        'Ignorado pelo Agente Faturas Hapvida.', 'Sim'))
    por = lambda t: [x for x in docs if x['tipo'] == t]
    nfs = {}
    for n in por('nf'):
        if n.get('nf'): nfs.setdefault(n['nf'], n)
    # boletos de mensalidade x copart
    bols = [b for b in por('boleto') if b['layout'] != 'individual']
    bol_cop = [b for b in bols if 'COPART' in b['rel'].upper() or b.get('copart')]
    bol_men = [b for b in bols if b not in bol_cop]
    bol_nosso = {b['nosso']: b for b in bol_men if b.get('nosso')}
    bol_obr = {b['obrigacao']: b for b in bol_men if b.get('obrigacao')}
    bol_cv = {(b.get('contrato'), (b.get('venc') or '')[:7]): b for b in bol_men if b.get('contrato') and b.get('venc')}
    # junta tudo por obrigação
    F = OrderedDict()
    def fat(o):
        return F.setdefault(o, dict(obrigacao=o, an=[], csv=[], portal=None, boleto=None, nf=None, copart=[]))
    for p in por('portal'): fat(p['obrigacao'])['portal'] = p
    for a in por('fatura'): fat(a['obrigacao'])['an'].append(a)
    for c in por('csv'):
        if c.get('obrigacao'): fat(c['obrigacao'])['csv'].append(c)
    usados = set()
    for o, f in F.items():
        a = f['an'][0] if f['an'] else None; p = f['portal']
        f['controle'] = (a or {}).get('controle') or (p or {}).get('nosso')
        f['venc'] = (a or {}).get('venc') or (p or {}).get('venc') or next((c['venc'] for c in f['csv'] if c.get('venc')), None)
        b = bol_obr.get(o) or bol_nosso.get(f['controle'])
        f['boleto'] = b
        if b: usados.add(id(b))
        nfn = (a or {}).get('nf') or (p or {}).get('nf') or (b or {}).get('nf')
        f['nf_num'] = nfn; f['nf'] = nfs.get(nfn) if nfn else None
        # contrato
        emp = [e[0] for e in (a or {}).get('empresas', [])]
        cand = [(p or {}).get('contrato'), (b or {}).get('contrato'), (f['nf'] or {}).get('contrato'),
                next((e for e in emp if e in cods_com), None),
                next((x.get('cod_arquivo') for x in f['an'] + f['csv'] if x.get('cod_arquivo') in emp), None),
                next((x.get('cod_arquivo') for x in f['an'] + f['csv'] if x.get('cod_arquivo')), None),
                emp[0] if emp else None]
        f['contrato'] = next(c for c in cand if c)
        if not b:
            b = bol_cv.get((f['contrato'], (f['venc'] or '')[:7]))
            if b and id(b) not in usados: f['boleto'] = b; usados.add(id(b))
        if len(f['an']) > 1:
            f['copias'] = [x['rel'] for x in f['an'][1:]]
    # boletos sem analítico/portal: fatura só com boleto (sem vidas)
    for b in bol_men:
        if id(b) in usados: continue
        if b['layout'] == 'bloqueto' and not b.get('obrigacao'): continue
        if not (b.get('contrato') or b.get('cod_arquivo')):
            log.append((b.get('mes_pasta'), None, 'Extração', f'Boleto sem contrato e sem fatura que o explique: {b["rel"]}.', 'Não gravado; conferir.', 'Não'))
            continue
        o = b.get('obrigacao') or f'BOL {b.get("contrato")} {(b.get("venc") or "")[:7]}'
        f = fat(o); f.update(boleto=b, contrato=b.get('contrato') or b.get('cod_arquivo'), venc=b.get('venc') or None, controle=b.get('nosso'), nf_num=b.get('nf'),
                             nf=nfs.get(b.get('nf')))
        if not f['venc']:
            f['venc_arquivo'] = b.get('mes_arquivo') or b.get('mes_pasta')
    # individuais (PF)
    for b in por('boleto'):
        if b['layout'] != 'individual': continue
        o = f'DOC {b["documento"]}'
        if o in F: F[o].setdefault('copias', []).append(b['rel']); continue
        f = fat(o); f.update(boleto=b, contrato=b.get('carteira'), venc=b.get('venc'), controle=b.get('nosso'), nf_num=None, individual=True)
    # copart: vai para a fatura de mensalidade do mesmo contrato e mês
    idx = defaultdict(list)
    for f in F.values():
        if f.get('venc'): idx[(f['contrato'], f['venc'][:7])].append(f)
    sub_de = {}
    for f in F.values():
        for a in f['an']:
            for e, _ in a.get('empresas', []):
                if e != f['contrato']: sub_de.setdefault(e, f['contrato'])
        for fl in (f['portal'] or {}).get('filhos', []):
            sub_de.setdefault(fl['cod'], f['contrato'])
    for c in por('copart') + [dict(b, tipo='copart_boleto') for b in bol_cop]:
        k = sub_de.get(c.get('contrato'), c.get('contrato'))
        alvo = idx.get((k, (c.get('venc') or c.get('mes_arquivo') or '')[:7]))
        if alvo: alvo[0]['copart'].append(c)
        else:
            log.append(((c.get('venc') or c.get('mes_arquivo') or '')[:7] or None, None, 'Conferência',
                        f'Copart sem fatura de mensalidade do mesmo contrato/mês: {c["rel"]}' + (f' (R$ {br(c["valor"])})' if c.get('valor') else '') + '.',
                        'Não gravada (copart não é fatura de mensalidade).', 'Sim'))
    # monta as faturas
    fats = []
    for o, f in F.items():
        a = f['an'][0] if f['an'] else None; p = f['portal']; b = f['boleto'] or {}
        venc = f.get('venc')
        comp = venc[:7] if venc else f.get('venc_arquivo')
        if not comp:
            log.append((None, None, 'Extração', f'Fatura {o} sem vencimento.', 'Não gravada.', 'Não')); continue
        if p and p.get('linhas'): linhas, fonte = p['linhas'], 'portal (Relatório.csv)'
        elif f['csv'] and f['csv'][0]['linhas']: linhas, fonte = f['csv'][0]['linhas'], 'CSV da pasta'
        elif a: linhas, fonte = a['linhas'], 'PDF analítico'
        elif f.get('individual'):
            linhas = [dict(empresa=b['contrato'], familia=b['carteira'], dep=x['carteira'][-4:-2], cod=x['carteira'], nome=b.get('razao') or x['nome'],
                           cpf=re.sub(r'\D', '', b.get('cpf') or ''), parentesco='TITULAR', plano_cod=None, plano=x['plano'], valor=x['valor'],
                           nasc=None, idade=None, inicio=None, mat=None) for x in b.get('beneficiarios', []) if x['tipo'].startswith('M')]
            fonte = 'boleto individual'
        else: linhas, fonte = [], None
        # regra: vidas do CSV (portal ou pasta); o PDF analítico só entra sem CSV. Se os dois trazem vidas e divergem, vale o CSV com aviso
        if a and a.get('linhas') and fonte in ('portal (Relatório.csv)', 'CSV da pasta') and len(a['linhas']) != len(linhas):
            log.append((comp, f['contrato'], 'Conferência', f'Fatura {o}: {fonte} tem {len(linhas)} vidas e o PDF analítico {len(a["linhas"])}.',
                        f'Gravadas as vidas do {fonte} (regra: CSV antes do PDF). Conferir o PDF.', 'Sim'))
        aj = (p or {}).get('ajustes') or (a or {}).get('ajustes') or (f['csv'][0].get('ajustes') if f['csv'] else []) or []
        valor = (a or {}).get('valor') if a else ((p or {}).get('valor') if p else b.get('valor'))
        if valor is None and f['csv']:
            valor = round(sum(x['valor'] for x in f['csv'][0]['linhas']) + sum(-x['valor'] if x['tipo'] == 'DESCONTO' else x['valor'] for x in aj), 2)
        erros = {'portal (Relatório.csv)': (p or {}).get('erros'), 'CSV da pasta': f['csv'][0].get('erros') if f['csv'] else None,
                 'PDF analítico': (a or {}).get('erros')}.get(fonte) or []
        if f.get('individual'):
            aj = [dict(tipo='ACRÉSCIMO', valor=x['valor'], obs=f'{x["tipo"]} {x["plano"]}') for x in b.get('beneficiarios', []) if not x['tipo'].startswith('M') and x['valor']]
        rel_an = (a or {}).get('rel'); rel_csv = f['csv'][0]['rel'] if f['csv'] else None
        empresas = OrderedDict((e, n) for e, n in (a or {}).get('empresas', []))
        for x in linhas: empresas.setdefault(x['empresa'], None)
        fats.append(dict(principal=f['contrato'], comp=comp, nf=o if not o.startswith(('BOL ', 'DOC ')) else o, obrigacao=o, venc=venc,
                         emissao=(a or {}).get('emissao') or (p or {}).get('emissao') or b.get('emissao'), periodo=b.get('periodo') or (None, None),
                         valor=valor, boleto_valor=b.get('valor') or ((p or {}).get('valor') if p else None), nosso=f.get('controle'), linha=b.get('linha'),
                         nf_num=f.get('nf_num'), nf_doc=f.get('nf'), linhas=linhas, fonte=fonte, ajustes=aj, erros=erros, copart=f['copart'],
                         subs=[e for e in empresas if e != f['contrato']], status=(p or {}).get('status'), portal=p, an=a, rel_an=rel_an, rel_csv=rel_csv, csv_arq=f['csv'][0]['arquivo'] if f['csv'] else None,
                         bol=f['boleto'], copias=f.get('copias', []), razao=(p or {}).get('razao') or b.get('razao') or (f['nf'] or {}).get('tomador')
                         or empresas.get(f['contrato']) or next(iter(empresas.values()), None),
                         cnpj=(p or {}).get('cnpj') or b.get('cnpj') or (f['nf'] or {}).get('cnpj'),
                         nc=next((nome_curto_arquivo(x['rel']) for x in f['an'] + f['csv'] + ([f['boleto']] if f['boleto'] else []) if x.get('rel')), None),
                         sem_venc=not venc, individual=f.get('individual', False)))
    # dados por contrato
    info = defaultdict(lambda: dict(comps=[], nomes=Counter(), subs=OrderedDict(), razao=None, cnpj=None))
    for f in sorted(fats, key=lambda x: x['comp']):
        i = info[f['principal']]; i['comps'].append(f['comp'])
        if f['venc']: i['venc'] = f['venc']
        if f['nc']: i['nomes'][f['nc']] += 1
        i['razao'] = f['razao'] or i['razao']; i['cnpj'] = f['cnpj'] or i['cnpj']
        an_emp = dict((f['an'] or {}).get('empresas', []))
        for s in f['subs']:
            i['subs'].setdefault(s, an_emp.get(s))
        if f['portal']:
            for fl in f['portal'].get('filhos', []):
                if fl['cod'] != f['principal']: i['subs'].setdefault(fl['cod'], None)
            i['status_portal'] = f['portal'].get('status_contrato')
        i['individual'] = f['individual']
    return fats, info


def cadastrar(wb, info, proc):
    E, S = wb['Empresas'], wb['Subestipulantes']
    rE = last_row(E); tplE = rE
    nid = max(int(E.cell(r, 1).value[1:]) for r in range(2, rE + 1) if E.cell(r, 1).value)
    ids = {}
    for r in range(2, rE + 1):
        if E.cell(r, 4).value == OP and E.cell(r, 5).value:
            ids[str(E.cell(r, 5).value)] = E.cell(r, 1).value; info[str(E.cell(r, 5).value)]['nc'] = E.cell(r, 11).value
    # código que já é subestipulante de uma empresa Hapvida (ex.: 0Y5VK da BROGOTA, empresa 0Y5WG) usa essa empresa, sem cadastrar outra
    nc_de = {v: info[k]['nc'] for k, v in ids.items()}
    sub_eid = {str(S.cell(r, 3).value): S.cell(r, 1).value for r in range(2, last_row(S) + 1) if S.cell(r, 1).value in nc_de and S.cell(r, 3).value}
    for princ in list(info):
        if princ not in ids and princ in sub_eid:
            ids[princ] = sub_eid[princ]; info[princ]['nc'] = nc_de[sub_eid[princ]]
    novas = []
    for princ in sorted(info):
        if princ in ids: continue
        i = info[princ]
        if rE >= NEMP + 1:
            raise SystemExit(f'Aba Empresas cheia ({NEMP} empresas cobertas pela aba Faturamento). Ampliar antes de cadastrar.')
        nid += 1; eid = f'E{nid:03d}'; ids[princ] = eid; rE += 1
        nc = i['nomes'].most_common(1)[0][0] if i['nomes'] else re.sub(r'\s+', ' ', (i.get('razao') or princ)).upper()[:25]
        i['nc'] = nc
        ult = max(i['comps'])
        obs = (f'Cadastrada pelo Agente Faturas Hapvida ({proc:%d/%m/%Y}). Plano NotreDame Intermédica (grupo Hapvida). '
               f'Código estipulante = código do contrato (igual ao extrato de comissão Hapvida).')
        if i.get('individual'): obs += ' Plano individual (PF).'
        if i['subs']: obs += f' {len(i["subs"])} subestipulantes/filiais/odonto no mesmo contrato (aba Subestipulantes).'
        if i.get('status_portal'): obs += f' Status no Portal Empresa: {i["status_portal"]}.'
        obs += f' Faturas lidas: {min(i["comps"])} a {ult}.'
        ativa = 'Não' if i.get('status_portal') == 'Canceled' else ('Sim' if ult >= '2026-09' else 'Não')
        row = [eid, i.get('razao'), i.get('cnpj'), OP, princ, princ, None,
               int(i['venc'][8:10]) if i.get('venc') else None, ativa, obs, nc,
               f'=IF(A{rE}="","",COUNTIFS(Subestipulantes!$A:$A,A{rE}))']
        for c, v in enumerate(row, 1):
            cs(E.cell(tplE, c), E.cell(rE, c)); E.cell(rE, c).value = v
        novas.append(eid)
    rS = last_row(S); tplS = rS
    exist = {(S.cell(r, 1).value, str(S.cell(r, 3).value)) for r in range(2, rS + 1)}
    for princ, i in info.items():
        subs = OrderedDict([(princ, i.get('razao'))]); subs.update(i['subs'])
        for k, raz in subs.items():
            if (ids[princ], k) in exist: continue
            exist.add((ids[princ], k)); rS += 1
            row = [ids[princ], princ, k, None, raz or (i.get('razao') if k == princ else None), i.get('cnpj') if k == princ else None,
                   i['nc'] if k == princ else (re.sub(r'\s+', ' ', raz).upper()[:25] if raz else k), 'Sim', SUBFORM.format(r=rS),
                   'Contrato principal (recebe o boleto).' if k == princ else f'Código do mesmo contrato {princ} (filial, unidade ou odonto) faturado junto.']
            for c, v in enumerate(row, 1):
                cs(S.cell(tplS, c), S.cell(rS, c)); S.cell(rS, c).value = v
    estender(E, NEMP + 1, 12, tplE)
    estender(S, max(rS, S.max_row), 10, tplS)
    return ids, novas


def gravar_faturas(wb, fats, info, ids, proc, base_dir, log):
    F, V = wb['Faturas'], wb['Vidas']
    rF = last_row(F); tplF = rF
    rV = last_row(V); tplV = rV
    ja = {}
    for r in range(2, rF + 1):
        ja[(F.cell(r, 1).value, F.cell(r, 2).value, str(F.cell(r, 5).value))] = r
    res = Counter(); copias = []; textos = []; gravadas = []
    for f in sorted(fats, key=lambda x: (x['comp'], ids[x['principal']], x['nf'])):
        eid = ids[f['principal']]; comp = f['comp']; nf = f['nf']; total = f['valor']
        chave = (comp, eid, nf)
        if chave in ja:
            r0 = ja[chave]; v0 = F.cell(r0, 10).value if r0 else None
            res['ja_na_planilha'] += 1
            if v0 is not None and total is not None and abs(float(v0) - total) > 0.01:
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf} já estava na planilha (linha {r0}) com valor R$ {br(float(v0))}; '
                            f'os arquivos trazem R$ {br(total)}.', 'Não gravada de novo (regra: só comparar). Verificar qual está certo.', 'Não'))
            continue
        ja[chave] = None
        nc = info[f['principal']]['nc']
        pasta = f'Faturas/{comp}/{PASTA}/{f["principal"]} - {nc}'
        base = f'{pasta}/{f["principal"]} - {nc}'
        arq_f = arq_v = arq_b = None
        if f['rel_an']:
            arq_f = f'{base} FATURADOS {PROD} {mmaa(comp)}.pdf'; copias.append((f['an']['arquivo'], arq_f))
        if f['portal'] and f['portal'].get('csv_txt'):
            arq_v = f'{base} FATURADOS {PROD} {mmaa(comp)}.csv'; textos.append((f['portal']['csv_txt'], arq_v))
        elif f['rel_csv']:
            arq_v = f'{base} FATURADOS {PROD} {mmaa(comp)}.csv'; copias.append((f['csv_arq'], arq_v))
        if f['bol']:
            arq_b = f'{base} BOLETO {PROD} {mmaa(comp)}.pdf'; copias.append((f['bol']['arquivo'], arq_b))
        if f['nf_doc']:
            copias.append((f['nf_doc']['arquivo'], f'{base} NF-e {PROD} {mmaa(comp)}.pdf'))
        for k, c in enumerate(f['copart'], 1):
            tipo = 'BOLETO COPART' if c['tipo'] == 'copart_boleto' else 'FATURADOS COPART'
            cod = c.get('contrato') or f['principal']
            copias.append((c['arquivo'], f'{pasta}/{cod} - {nc} {tipo} {PROD} {mmaa(comp)}.pdf'))
        obs = []
        if f['individual']: obs.append(f'Plano individual (PF), carteira {f["principal"]}.')
        if f['portal']:
            obs.append(f'Portal Empresa: {f["portal"].get("status")} ({f["portal"].get("tipo_fatura")}).')
        if not f['linhas']: obs.append('Sem relação de vidas (só boleto).')
        else: obs.append(f'Vidas lidas do {f["fonte"]}.')
        soma = round(sum(x['valor'] or 0 for x in f['linhas']), 2)
        if f['ajustes']:
            obs.append('Descontos/acréscimos da fatura (não detalhados por vida): ' + '; '.join(f'{a["tipo"]} R$ {br(a["valor"])} - {a["obs"][:90]}' for a in f['ajustes']) + '.')
        if f['linhas'] and total is not None and abs(soma - total) > 0.015:
            aj = round(sum(-a['valor'] if a['tipo'] == 'DESCONTO' else a['valor'] for a in f['ajustes']), 2)
            if abs(soma + aj - total) <= 0.015:
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf}: soma das vidas R$ {br(soma)} + descontos/acréscimos R$ {br(aj)} = valor da fatura R$ {br(total)}.',
                            'Gravada com o valor da fatura; a Conferência valor mostra os descontos/acréscimos como diferença (explicado nas Observações).', 'Sim'))
            else:
                obs.append(f'Fatura R$ {br(total)} - soma das vidas R$ {br(soma)} = R$ {br(round(total - soma, 2))}.')
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf}: valor R$ {br(total)} ≠ soma das vidas R$ {br(soma)} (descontos/acréscimos lidos R$ {br(aj)}).',
                            'Gravada com o valor da fatura. Verificar o analítico.', 'Não'))
        if f['copart']:
            obs.append('Copart separada no mesmo mês: ' + '; '.join(f'{os.path.basename(c["rel"])}' + (f' R$ {br(c["valor"])}' if c.get('valor') else '') for c in f['copart'][:4]) + '.')
        if f['nf_num']: obs.append(f'NF-e {f["nf_num"]}.')
        if f['copias']: obs.append(f'Cópias repetidas na pasta: {len(f["copias"])}.')
        for e in f['erros'][:3]:
            log.append((comp, f['principal'], 'Extração', f'Fatura {nf}: {e}', 'Vida não gravada; conferir no analítico.', 'Não'))
        origem = [x for x in [f['rel_an'], f['rel_csv'], (f['bol'] or {}).get('rel'), (f['portal'] or {}).get('rel')] if x]
        obs.append('Origem: ' + '; '.join(('HAPVIDA/' + x if not x.startswith('portal/') else 'Portal Empresa ' + x[7:]) for x in origem[:3]) + '.')
        per = f['periodo']; ini = d(per[0]) or datetime(int(comp[:4]), int(comp[5:7]), 1); fim = d(per[1]) or fim_mes(comp)
        rF += 1; r = rF
        row = {1: comp, 2: eid, 4: OP, 5: nf, 6: d(f.get('emissao')), 7: d(f.get('venc')), 8: ini, 9: fim,
               10: total, 11: len(f['linhas']) if f['linhas'] else None, 17: f['boleto_valor'], 19: f['nosso'], 20: f['linha'],
               21: arq_f, 22: arq_v, 23: arq_b, 24: proc, 25: ' '.join(obs), 26: f['principal'], 27: ', '.join(f['subs']) or None,
               }
        for col in range(1, 33):
            cs(F.cell(tplF, col), F.cell(r, col)); F.cell(r, col).value = row.get(col)
        F.cell(r, 3).value = f'=IF(B{r}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{r},Empresas!$A:$A,0)),"ID não cadastrado"))'
        F.cell(r, 12).value = f'=IF(A{r}="","",COUNTIFS(Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r},Vidas!$P:$P,"Ativo"))'
        F.cell(r, 13).value = f'=IF(A{r}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r}))'
        F.cell(r, 14).value = f'=IF(A{r}="","",J{r}-M{r})'
        F.cell(r, 15).value = f'=IF(A{r}="","",IF(L{r}=0,"Sem vidas",IF(ABS(N{r})<=0.01,"OK","Divergência")))'
        F.cell(r, 16).value = f'=IF(OR(A{r}="",K{r}=""),"",IF(L{r}=K{r},"OK","Divergência"))'
        F.cell(r, 18).value = f'=IF(OR(A{r}="",Q{r}=""),"",IF(ABS(Q{r}-J{r}-N(AE{r}))<=0.01,"OK","Divergência"))'
        F.cell(r, 29).value = f'=IF(A{r}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{r}))'
        F.cell(r, 30).value = (f'=IF(A{r}="","",IF(AC{r}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{r},'
                               f'Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')
        res['faturas'] += 1; res['sem_vidas'] += not f['linhas']; res['sem_boleto'] += not f['bol'] and not f['portal']
        for x in f['linhas']:
            rV += 1
            par = (x.get('parentesco') or '').upper()
            tit = par == 'TITULAR' or (not par and x.get('dep') == '00')
            vals = {1: comp, 2: eid, 3: OP, 4: nf, 5: x['familia'], 6: x.get('dep'), 7: x['nome'], 8: x.get('cpf'), 9: d(x.get('nasc')), 10: x.get('idade'),
                    11: PAREN.get(par, par.title() if par else ('Titular' if tit else None)), 12: 'Titular' if tit else 'Dependente',
                    13: x.get('plano_cod'), 14: x.get('plano') or x.get('plano_cod'), 15: x.get('mat'), 16: 'Ativo', 17: d(x.get('inicio')),
                    19: x['valor'], 20: 0, 21: 0, 22: 0, 23: 0, 24: 0, 25: 0,
                    29: f'Credencial {x["cod"]}' + (f'; desconto R$ {br(x["desconto"])}' if x.get('desconto') else ''), 30: x['empresa'],
                    31: x.get('unidade'), 32: 0}
            for col in range(1, 33):
                cs(V.cell(tplV, col if col < 32 else 21), V.cell(rV, col)); V.cell(rV, col).value = vals.get(col)
            V.cell(rV, 26).value = f'=IF(A{rV}="","",N(U{rV})+N(V{rV})+N(W{rV})+N(X{rV})+N(AF{rV}))'
            V.cell(rV, 27).value = f'=IF(A{rV}="","",N(S{rV})+N(T{rV})+N(Z{rV})+N(Y{rV}))'
            V.cell(rV, 28).value = f'=IF(A{rV}="","",IF(P{rV}="Excluído","Exclusão",IF(N(T{rV})>0,"Inclusão (pró-rata)","-")))'
            res['vidas'] += 1
        ok = bool(f['linhas']) and total is not None and abs(soma - total) < 0.015 and not f['erros']
        if ok: res['ok'] += 1
        else: res['pend'] += 1
        if not f['linhas']:
            log.append((comp, f['principal'], 'Download', f'Fatura {nf} sem relação de vidas (só boleto: {(f["bol"] or {}).get("rel")}).',
                        'Gravada com o valor do boleto; sem vidas. Baixar o analítico (Relatório.csv) no portal.', 'Não'))
        gravadas.append(dict(comp=comp, eid=eid, principal=f['principal'], nf=nf, venc=f['venc'], total=total, vinf=len(f['linhas']) or None,
                             vidas=len(f['linhas']), soma=soma, subs=len(f['subs']), boleto=bool(f['bol'] or f['portal']), status=f['status']))
    estender(F, max(rF, F.max_row), 32, tplF)
    estender(V, max(rV + 500, V.max_row), 32, tplV)
    for r in range(tplV + 1, V.max_row + 1):
        cs(V.cell(tplV, 21), V.cell(r, 32))
    return res, copias, textos, gravadas, rF, rV


def registrar(wb, res, log, info, ids, novas, comps, proc, backup, inicio, fontes):
    AG, EX, LG, EXH = wb['Agentes'], wb['Execuções'], wb['Log de Erros'], wb['Execução']
    ra = last_row(AG)
    if not any(AG.cell(r, 1).value == AGENTE for r in range(2, ra + 1)):
        ra += 1
        for c in range(1, 13):
            cs(AG.cell(2, c), AG.cell(ra, c))
            v = AG.cell(2, c).value
            AG.cell(ra, c).value = re.sub(r'(?<=[A-Z])2\b', str(ra), v) if isinstance(v, str) and v.startswith('=') else None
        for c, v in enumerate([AGENTE, 'Coleta as faturas no Portal Empresa Hapvida/NotreDame (Extrato financeiro + Relatório.csv) e lê os analíticos, boletos e NF-e '
                               'da pasta da equipe; grava Faturas + Vidas', OP, 'Empresas, Subestipulantes, Faturas, Vidas, Log de Erros', 'Mensal', 10,
                               'Scripts/faturas_hapvida.py + Scripts/gravar_hapvida.py', 'Ativo'], 1):
            AG.cell(ra, c).value = v
    pend = sum(1 for x in log if x[5] == 'Não')
    re_ = last_row(EX) + 1
    for c in range(1, 14): cs(EX.cell(re_ - 1, c), EX.cell(re_, c))
    resumo = f'{res["sem_vidas"]} faturas sem relação de vidas; {pend} pendências no Log de Erros' if pend else ''
    for c, v in enumerate([AGENTE, f'{comps[0]} a {comps[-1]}' if comps else '', inicio, datetime.now().replace(microsecond=0),
                           '⚠️ OK com pendência' if pend else '✅ OK', res['faturas'], res['ok'], res['pend'], 0, resumo,
                           backup, f'Primeira carga Hapvida ({fontes}): {res["faturas"]} faturas, {res["vidas"]} linhas de vidas, '
                           f'{len(novas)} empresas novas; {res["ja_na_planilha"]} já estavam na planilha.', 'Bruno'], 1):
        EX.cell(re_, c).value = v
    rl = last_row(LG); tplL = rl
    for comp, k, etapa, erro, acao, ok in log:
        rl += 1
        for c in range(1, 10): cs(LG.cell(tplL, c), LG.cell(rl, c))
        for c, v in enumerate([proc, comp, OP, ids.get(k) if k else None, etapa, erro, acao, ok, 'Agente Faturas'], 1):
            LG.cell(rl, c).value = v
    estender(LG, max(rl, LG.max_row), 9, tplL)
    rh = last_row(EXH, 2) + 1
    for c in range(2, 10): cs(EXH.cell(rh - 1, c), EXH.cell(rh, c))
    for c, v in enumerate([proc, f'{comps[0]} a {comps[-1]}' if comps else '', f'Hapvida: {len(info)} contratos ({fontes})', res['faturas'],
                           res['ok'], res['pend'], f'{res["sem_vidas"]} sem vidas', 'Primeira carga Hapvida / NotreDame.'], 2):
        EXH.cell(rh, c).value = v


LEIA = ['AGENTE FATURAS HAPVIDA (NOTREDAME INTERMÉDICA)',
        'Fontes: (1) Portal Empresa Hapvida/NotreDame (portal-empresa.hapvidagndi.com.br, login feito pela equipe): para cada contrato, Extrato financeiro → Em aberto e Histórico → Relatório.csv; '
        '(2) pasta "Área de Trabalho/boletos a enviar/HAPVIDA/MM-AA/<empresa>/" para os contratos que não estão nesse login (analítico PDF/CSV, boleto, NF-e). Scripts: Scripts/faturas_hapvida.py (leitura) e Scripts/gravar_hapvida.py (gravação).',
        'Uma fatura por obrigação (= Nº fatura do extrato de comissão Hapvida). Competência = mês do vencimento. Nosso número = "Controle" do analítico (= Nº documento no portal).',
        'Empresas: uma linha por contrato que recebe o boleto; Código estipulante = código do contrato (igual à aba Comissões). Filiais, unidades e odonto do mesmo contrato ficam em Subestipulantes.',
        'Valor total = valor da fatura (já com descontos e acréscimos). Pró-rata de cancelamento, mínimo contratual e reajuste retroativo não vêm por vida: a Conferência valor mostra a diferença, explicada nas Observações.',
        'Copart (fatura de co-participação) não vira fatura de mensalidade: o total vai nas Observações da fatura do mês. Planos individuais (3010V) entram como fatura própria.',
        'Arquivos copiados para Faturas/AAAA-MM/Hapvida/<contrato> - <NOME CURTO>/ (FATURADOS .pdf/.csv, BOLETO, NF-e). O original fica na pasta da equipe.']


def leia_me(wb):
    ws = wb['Leia-me']
    if any(ws.cell(r, 2).value == LEIA[0] for r in range(1, ws.max_row + 1)):
        return
    hdr = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or '').startswith('AGENTE FATURAS BRADESCO'))
    r = last_row(ws, 2) + 1
    for i, t in enumerate(LEIA):
        r += 1
        cs(ws.cell(hdr if i == 0 else hdr + 1, 2), ws.cell(r, 2)); ws.cell(r, 2).value = t


def painel(gravadas, wb, base_dir, proc):
    """JSON do painel (coleção faturas da Central), um por competência a partir de 2026-01."""
    emp = {r[0]: r[1] for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0]}
    docs = defaultdict(list)
    for g in gravadas:
        if g['comp'] < '2026-01': continue
        venc = d(g['venc'])
        if not g['boleto']: bol = 'Sem boleto'
        elif g['status']: bol = 'Pago' if g['status'] == 'Liquidado' else 'Em aberto'
        else: bol = 'Em aberto' if venc and venc >= proc else 'Pago'
        docs[g['comp']].append(dict(eid=g['eid'], empresa=(emp.get(g['eid']) or g['principal'])[:40], titulo=g['nf'], venc=venc.strftime('%d/%m/%Y') if venc else '',
                                    valor=round(g['total'] or 0, 2), vidas=g['vidas'], ativos=g['vidas'], subs=1 + g['subs'], boleto=bol,
                                    confValor=('Sem vidas' if not g['vidas'] else ('OK' if abs(g['soma'] - (g['total'] or 0)) < 0.015 else 'Divergência')),
                                    confVidas=('' if not g['vidas'] else 'OK')))
    out = os.path.join(base_dir, 'Scripts', '_painel'); os.makedirs(out, exist_ok=True)
    agora = datetime.now().astimezone().replace(microsecond=0).isoformat()
    for comp, it in sorted(docs.items()):
        json.dump(dict(comp=comp, operadora=OP, atualizadoEm=agora, parcial=False, pendentes=[], itens=sorted(it, key=lambda x: x['eid'])),
                  open(os.path.join(out, f'fat_{comp}_hapvida.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    return sorted(docs)


def main(js, xin, xout, base_dir, proc, *opts):
    inicio = datetime.now().replace(microsecond=0)
    proc = datetime.fromisoformat(proc)
    if os.path.exists(os.path.join(os.path.dirname(os.path.abspath(xin)), '~$' + os.path.basename(xin))):
        raise SystemExit('A planilha está aberta no Excel (~$ encontrado). Feche e rode de novo.')
    docs = json.load(open(js, encoding='utf-8'))
    log = []
    wb = openpyxl.load_workbook(xin)
    fats, info = preparar(docs, log, codigos_comissao(wb))
    ids, novas = cadastrar(wb, info, proc)
    res, copias, textos, gravadas, rF, rV = gravar_faturas(wb, fats, info, ids, proc, base_dir, log)
    comps = sorted({g['comp'] for g in gravadas})
    fontes = ' + '.join(x for x, ok in [('Portal Empresa', any(x['tipo'] == 'portal' for x in docs)), ('pasta da equipe', any(x['tipo'] != 'portal' for x in docs))] if ok)
    backup = f'Backups/Base_Faturas_Vidas antes Faturas Hapvida {proc:%Y-%m-%d}.xlsx'
    registrar(wb, res, log, info, ids, novas, comps, proc, backup, inicio, fontes)
    leia_me(wb)
    n = 0; pain = []
    if '--simular' not in opts:
        wb.save(xout)
        if '--sem-copia' not in opts:
            for src, dst in copias:
                p = os.path.join(base_dir, dst)
                if not os.path.exists(p):
                    os.makedirs(os.path.dirname(p), exist_ok=True); shutil.copy2(src, p); n += 1
            for txt, dst in textos:
                p = os.path.join(base_dir, dst)
                if not os.path.exists(p):
                    os.makedirs(os.path.dirname(p), exist_ok=True); open(p, 'w', encoding='cp1252', errors='replace', newline='').write(txt); n += 1
        pain = painel(gravadas, wb, base_dir, proc)
    json.dump(dict(gravadas=gravadas, log=log, ids=ids), open(os.path.splitext(js)[0] + '_resultado.json', 'w', encoding='utf-8'), ensure_ascii=False, default=str)
    print(json.dumps(dict(res, novas_empresas=len(novas), ids=f'{novas[0]}..{novas[-1]}' if novas else '', copiados=n, a_copiar=len(copias) + len(textos),
                          log=len(log), log_pend=sum(1 for x in log if x[5] == 'Não'), ult_fatura=rF, ult_vida=rV, comps=comps, painel=len(pain)), ensure_ascii=False))


if __name__ == '__main__':
    main(*sys.argv[1:])

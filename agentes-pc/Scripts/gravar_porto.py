# -*- coding: utf-8 -*-
"""Grava na Base_Faturas_Vidas.xlsx as faturas Porto Saúde lidas por faturas_porto.py.
Uso: python gravar_porto.py <json do leitor> <planilha entrada> <planilha saida> <pasta FATURAS> <AAAA-MM-DD processamento> [--sem-copia] [--simular]
- Uma linha na aba Faturas por título e uma linha na aba Vidas por registro 02 do TXT (colunas azuis; as fórmulas
  cinzas são as mesmas das outras operadoras).
- Título já gravado na aba Faturas não é gravado de novo (conta como "já na planilha").
- Os arquivos já estão em Faturas/AAAA-MM/PortoSaude (o baixar_porto.js grava lá): nada é copiado; --sem-copia é aceito
  só para o rodar.py chamar igual às outras operadoras. --simular faz tudo menos salvar.
- Fatura sem ID Empresa (sub/estipulante sem cadastro) não é gravada: vai para o resumo como pendência.
"""
import json, sys, os
from datetime import datetime
from collections import defaultdict
import openpyxl
from shaddai_comum import d, cs, last_row  # módulo comum


def rel(p):
    """caminho como a base guarda: a partir de 'Faturas/'."""
    p = (p or '').replace('\\', '/')
    i = p.rfind('/Faturas/')
    return p[i + 1:] if i >= 0 else (p or None)


def vida_formulas(V, r):
    V.cell(r, 26).value = f'=IF(A{r}="","",N(U{r})+N(V{r})+N(W{r})+N(X{r})+N(AF{r}))'
    V.cell(r, 27).value = f'=IF(A{r}="","",N(S{r})+N(T{r})+N(Z{r})+N(Y{r}))'
    V.cell(r, 28).value = f'=IF(A{r}="","",IF(P{r}="Excluído","Exclusão",IF(N(T{r})>0,"Inclusão (pró-rata)","-")))'


def fatura_formulas(F, r):
    F.cell(r, 3).value = f'=IF(B{r}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{r},Empresas!$A:$A,0)),"ID não cadastrado"))'
    F.cell(r, 12).value = f'=IF(A{r}="","",COUNTIFS(Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r},Vidas!$P:$P,"Ativo"))'
    F.cell(r, 13).value = f'=IF(A{r}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r}))'
    F.cell(r, 14).value = f'=IF(A{r}="","",J{r}-M{r})'
    F.cell(r, 15).value = f'=IF(A{r}="","",IF(L{r}=0,"Sem vidas",IF(ABS(N{r})<=0.01,"OK","Divergência")))'
    F.cell(r, 16).value = f'=IF(OR(A{r}="",K{r}=""),"",IF(L{r}=K{r},"OK","Divergência"))'
    F.cell(r, 18).value = f'=IF(OR(A{r}="",Q{r}=""),"",IF(ABS(Q{r}-J{r}-N(AE{r}))<=0.01,"OK","Divergência"))'
    F.cell(r, 29).value = f'=IF(A{r}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{r}))'
    F.cell(r, 30).value = (f'=IF(A{r}="","",IF(AC{r}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{r},'
                           f'Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')


def arquivo_vidas(f):
    """TXT extraído quando a fatura tem um TXT só; o .zip quando o título junta várias subestipulantes/unidades."""
    if len(f['txts']) == 1 and f.get('arquivo_vidas', '').lower().endswith('.zip'):
        t = f['arquivo_vidas'][:-4] + '.txt'
        if os.path.exists(t):
            return rel(t)
    return rel(f.get('arquivo_vidas'))


def main(js, xin, xout, base_dir, proc, *opts):
    proc = datetime.fromisoformat(proc)
    fats = [x for x in json.load(open(js, encoding='utf-8')) if x.get('tipo') == 'fatura']
    wb = openpyxl.load_workbook(xin)
    F, V = wb['Faturas'], wb['Vidas']
    ja = {str(F.cell(r, 5).value) for r in range(2, F.max_row + 1) if F.cell(r, 1).value and F.cell(r, 5).value}
    rF = last_row(F); tplF = rF
    rV = last_row(V); tplV = rV
    res = defaultdict(int); pend = []; novos = []
    for f in sorted(fats, key=lambda x: (x['comp'], x['pasta'], x['sufixo'])):
        if f['titulo'] in ja:
            res['ja_na_planilha'] += 1; continue
        if not f.get('id_empresa'):
            pend.append(f'{f["pasta"]} título {f["titulo"]}: sem ID Empresa (cadastre a sub {f["sub"]} na aba Subestipulantes)'); continue
        ja.add(f['titulo'])
        b = f.get('boleto') or {}
        subs = sorted({t['sub'] for t in f['txts']})
        obs = [f'Layout {f["layout"]}. Origem dos dados: TXT do .zip.']
        if not b:
            obs.append('Boleto já pago (botão Boleto desabilitado no portal).')
        if len(subs) > 1:
            obs.append(f'Fatura única do grupo: {len(subs)} subestipulantes no mesmo título.')
        obs += f.get('avisos', [])
        rF += 1
        row = {1: f['comp'], 2: f['id_empresa'], 4: f['produto'], 5: f['titulo'], 6: d(f.get('emissao')), 7: d(f.get('venc')),
               8: d(f['cob'][0]), 9: d(f['cob'][1]), 10: f.get('total'), 11: f.get('vidas'),
               17: b.get('valor'), 19: b.get('nosso'), 20: b.get('linha'),
               21: rel(f['arquivo']), 22: arquivo_vidas(f), 23: rel(b.get('arquivo')), 24: proc, 25: ' '.join(obs),
               26: str(int(f['estipulante'])) if f.get('estipulante') else None, 27: f.get('sub'), 28: f.get('unidade')}
        for col in range(1, 33):
            cs(F.cell(tplF, col), F.cell(rF, col))
            F.cell(rF, col).value = row.get(col)
        fatura_formulas(F, rF)
        res['faturas'] += 1; novos.append(f['titulo'])
        for v in f['vidas_linhas']:
            rV += 1
            c = v['copart']
            vals = {1: f['comp'], 2: f['id_empresa'], 3: f['produto'], 4: f['titulo'], 5: v['seguro'], 6: v['dep'], 7: v['nome'],
                    8: v['cpf'], 10: v['idade'], 11: v['parentesco'], 12: 'Titular' if v['titular'] else 'Dependente',
                    13: v['cod_plano'], 14: v['plano'], 15: v['matricula'], 16: v['situacao'], 17: d(v['adesao']), 18: d(v['exclusao']),
                    19: v['premio_base'], 20: v['prorata'], 21: c['consultas'], 22: c['exames'], 23: c['psocorro'], 24: c['internacao'],
                    25: v['despesas'], 30: v['sub'], 31: v['unidade']}
            for col in range(1, 33):
                cs(V.cell(tplV, col), V.cell(rV, col))
                V.cell(rV, col).value = vals.get(col)
            vida_formulas(V, rV)
            res['vidas'] += 1
            res['vidas_ativas'] += v['situacao'] == 'Ativo'
    if '--simular' not in opts:
        wb.save(xout)
    print(json.dumps(dict(res, titulos_novos=novos, pendencias=pend, ultima_fatura=rF, ultima_vida=rV,
                          simulado='--simular' in opts), ensure_ascii=False))


if __name__ == '__main__':
    main(*sys.argv[1:6], *sys.argv[6:])

import openpyxl, json, sys, os
from collections import defaultdict
from datetime import datetime
src, outdir = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
SLUG = {'Amil': 'amil', 'Amil Dental': 'amildental', 'Santa Helena': 'santahelena'}
emp = {r[0]: r[1] for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0]}
docs = defaultdict(list); hoje = datetime(2026, 10, 17)
for r in wb['Faturas'].iter_rows(min_row=2, values_only=True):
    if r[3] not in SLUG or not r[0] or r[0] < '2026-01': continue
    venc = r[6]
    docs[(r[0], r[3])].append(dict(eid=r[1], empresa=(emp.get(r[1]) or '')[:40], titulo=str(r[4]), venc=venc.strftime('%d/%m/%Y') if venc else '',
        valor=round(r[9] or 0, 2), vidas=r[10] if r[10] is not None else (r[11] or 0), ativos=r[11] or 0,
        subs=1 + (len(str(r[26]).split(',')) if r[26] else 0),
        boleto='Em aberto' if venc and venc >= hoje else 'Pago', confValor=r[14] or '', confVidas=r[15] or ('OK' if r[14] == 'OK' else '')))
os.makedirs(outdir, exist_ok=True)
for (comp, op), it in docs.items():
    json.dump(dict(comp=comp, operadora=op, atualizadoEm='2026-10-17T17:30:00Z', parcial=False, pendentes=[], itens=it),
              open(f'{outdir}/fat_{comp}_{SLUG[op]}.json', 'w', encoding='utf-8'), ensure_ascii=False)
print(len(docs), sum(len(v) for v in docs.values()))

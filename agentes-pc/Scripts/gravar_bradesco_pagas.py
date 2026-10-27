# -*- coding: utf-8 -*-
"""Grava a lista de faturas PAGAS do Portal de Negócios Bradesco (aba Histórico de pagamento).
Entrada: JSON {keys, rows} salvo do portal (uma linha por subfatura), sem PDF e sem vidas.
- Aba "Bradesco Pagas": histórico completo, uma linha por subfatura/boleto pago.
- Aba Faturas: só as competências >= COMP_MIN (padrão 2026-01), com valores, sem vidas.
- Empresas: cria a apólice que ainda não existe (com fatura em competência >= COMP_MIN).
- Faturamento: amplia para N_EMP empresas por bloco.
Uso: python3 gravar_bradesco_pagas.py <json> <planilha entrada> <planilha saida> <AAAA-MM-DD> [COMP_MIN]
"""
import json, sys, copy, re
from datetime import datetime
from collections import OrderedDict
import openpyxl
from shaddai_comum import cs, last_row  # módulo comum
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment

N_EMP = 200
STOP = {'LTDA', 'ME', 'EIRELI', 'EPP', 'SA', 'S/A', 'CIA', 'DE', 'DA', 'DO', 'DOS', 'DAS', 'E', '&', '-'}


def d(s):
    return datetime.strptime(s, '%d/%m/%Y') if s else None


def nome_curto(est):
    w = [x for x in re.split(r'\s+', est.strip()) if x and x.upper() not in STOP and not re.fullmatch(r'\d{6,}', x)]
    return ' '.join(w[:2]) or est.strip()


def ampliar_faturamento(FT, N):
    # acha o total do bloco 1
    t1_old = next(r for r in range(6, 400) if FT.cell(r, 2).value == 'Total da carteira')
    if t1_old - 6 >= N:
        return
    tt_old = next(r for r in range(t1_old + 1, t1_old + 6) if FT.cell(r, 1).value and 'Vidas' in str(FT.cell(r, 1).value))
    hh_old = tt_old + 1; d0_old = hh_old + 1
    t2_old = next(r for r in range(d0_old, d0_old + 400) if FT.cell(r, 2).value == 'Total da carteira')

    def snap(r):
        return [copy.copy(FT.cell(r, c)._style) for c in range(1, 18)], FT.row_dimensions[r].height
    sd1, st1, sgap, sti2, sh2, sd2, st2 = (snap(r) for r in (6, t1_old, t1_old + 1, tt_old, hh_old, d0_old, t2_old))
    head2 = [FT.cell(hh_old, c).value for c in range(1, 18)]
    title2 = FT.cell(tt_old, 1).value
    for r in range(6, max(t2_old, 6 + 2 * N + 10) + 1):
        for c in range(1, 18):
            FT.cell(r, c).value = None

    def apply(r, st):
        for c in range(1, 18):
            FT.cell(r, c)._style = copy.copy(st[0][c - 1])
        FT.row_dimensions[r].height = st[1]
    mcols = [L(c) for c in range(4, 16)]
    t1 = 6 + N
    for i in range(N):
        r = 6 + i; apply(r, sd1)
        FT.cell(r, 1).value = f'=IF(INDEX(Empresas!$A$2:$A${N+1},{i+1})="","",INDEX(Empresas!$A$2:$A${N+1},{i+1}))'
        FT.cell(r, 2).value = f'=IF($A{r}="","",INDEX(Empresas!$B$2:$B${N+1},{i+1}))'
        FT.cell(r, 3).value = f'=IF($A{r}="","",INDEX(Empresas!$D$2:$D${N+1},{i+1}))'
        for col in mcols:
            FT[f'{col}{r}'] = f'=IF($A{r}="","",SUMIFS(Faturas!$J:$J,Faturas!$A:$A,{col}$5,Faturas!$B:$B,$A{r}))'
        FT[f'P{r}'] = f'=IF($A{r}="","",SUM(D{r}:O{r}))'
        FT[f'Q{r}'] = f'=IF(OR($A{r}="",$P${t1}=0),"",P{r}/$P${t1})'
    apply(t1, st1); FT.cell(t1, 2).value = 'Total da carteira'
    for col in mcols + ['P']:
        FT[f'{col}{t1}'] = f'=SUM({col}6:{col}{t1-1})'
    apply(t1 + 1, sgap); apply(t1 + 2, sgap)
    tt = t1 + 3; hh = tt + 1; d0 = hh + 1; t2 = d0 + N
    apply(tt, sti2); FT.cell(tt, 1).value = title2
    apply(hh, sh2)
    for c, v in enumerate(head2, 1):
        FT.cell(hh, c).value = v
    for i in range(N):
        r = d0 + i; apply(r, sd2)
        FT.cell(r, 1).value = f'=IF(INDEX(Empresas!$A$2:$A${N+1},{i+1})="","",INDEX(Empresas!$A$2:$A${N+1},{i+1}))'
        FT.cell(r, 2).value = f'=IF($A{r}="","",INDEX(Empresas!$B$2:$B${N+1},{i+1}))'
        FT.cell(r, 3).value = f'=IF($A{r}="","",INDEX(Empresas!$D$2:$D${N+1},{i+1}))'
        for col in mcols:
            FT[f'{col}{r}'] = f'=IF($A{r}="","",SUMIFS(Faturas!$L:$L,Faturas!$A:$A,{col}${hh},Faturas!$B:$B,$A{r}))'
        FT[f'P{r}'] = f'=IF($A{r}="","",SUM(D{r}:O{r}))'
    apply(t2, st2); FT.cell(t2, 2).value = 'Total da carteira'
    for col in mcols + ['P']:
        FT[f'{col}{t2}'] = f'=SUM({col}{d0}:{col}{t2-1})'


def main(js, xin, xout, proc, comp_min='2026-01'):
    proc = datetime.fromisoformat(proc)
    raw = json.load(open(js, encoding='utf-8'))
    rows = [dict(zip(raw['keys'], r)) for r in raw['rows']]
    for r in rows:
        r['comp'] = f"{r['subAno']}-{r['subMes']:02d}"
        r['nf'] = f"{r['sucursal']}{r['numeroFatura']:06d}"
        r['op'] = 'Bradesco Dental' if r['subProduto'] == 'Dental' else 'Bradesco Saúde'
    rows.sort(key=lambda r: (r['comp'], r['estipulante'], r['apolice'], r['subfatura']))
    wb = openpyxl.load_workbook(xin)
    E, F, FT = wb['Empresas'], wb['Faturas'], wb['Faturamento']

    # ---------- Empresas
    rE = last_row(E); tplE = rE
    nid = max(int(E.cell(r, 1).value[1:]) for r in range(2, rE + 1) if E.cell(r, 1).value)
    ids = {str(E.cell(r, 5).value): E.cell(r, 1).value for r in range(2, rE + 1)
           if str(E.cell(r, 4).value or '').startswith('Bradesco')}
    novas = []
    ult = {}
    for r in rows:
        ult[str(r['apolice'])] = r
    for ap in sorted({str(r['apolice']) for r in rows if r['comp'] >= comp_min}, key=lambda a: ult[a]['estipulante']):
        if ap in ids:
            continue
        r0 = ult[ap]
        nid += 1; eid = f'E{nid:03d}'; ids[ap] = eid; rE += 1
        row = [eid, r0['estipulante'], r0['cnpj'], r0['op'], ap, None, ap, int(r0['vencimento'][:2]), 'Sim',
               f'Cadastrada pelo Agente Faturas a partir da lista de faturas pagas do Portal Bradesco ({proc:%d/%m/%Y}). '
               f'Razão social abreviada como no portal; completar quando baixar um PDF.', nome_curto(r0['estipulante']),
               f'=IF(A{rE}="","",COUNTIFS(Subestipulantes!$A:$A,A{rE}))']
        for c, v in enumerate(row, 1):
            cs(E.cell(tplE, c), E.cell(rE, c)); E.cell(rE, c).value = v
        novas.append(eid)
    if rE > 201:
        raise SystemExit(f'Empresas passou de 200 linhas formatadas ({rE})')

    # ---------- Faturas (competência >= comp_min)
    if F.cell(1, 32).value != 'Data de pagamento':
        F.cell(1, 32).value = 'Data de pagamento'; cs(F.cell(1, 7), F.cell(1, 32))
        F.column_dimensions['AF'].width = F.column_dimensions['G'].width or 12
    rF = last_row(F); tplF = rF
    ja = {str(F.cell(r, 5).value) for r in range(2, rF + 1) if F.cell(r, 1).value}
    n_fat = 0; pulados = 0
    for r in rows:
        if r['comp'] < comp_min:
            continue
        if r['nf'] in ja:
            pulados += 1
            continue
        ja.add(r['nf'])
        rF += 1; n_fat += 1
        ap = str(r['apolice'])
        vals = {1: r['comp'], 2: ids[ap], 4: r['op'], 5: r['nf'], 6: d(r['emissao']), 7: d(r['vencimento']),
                10: round(r['premio'] - r['iof'], 2), 17: r['premio'], 24: proc,
                25: f"Fatura paga em {r['baixa']}. Origem: lista Histórico de pagamento do portal (sem PDF e sem vidas).",
                26: ap, 27: f"{r['subfatura']:04d}", 31: r['iof'], 32: d(r['baixa'])}
        for col in range(1, 33):
            cs(F.cell(tplF, col if col < 32 else 7), F.cell(rF, col))
            F.cell(rF, col).value = vals.get(col)
        x = rF
        F.cell(x, 3).value = f'=IF(B{x}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{x},Empresas!$A:$A,0)),"ID não cadastrado"))'
        F.cell(x, 12).value = f'=IF(A{x}="","",COUNTIFS(Vidas!$A:$A,A{x},Vidas!$B:$B,B{x},Vidas!$D:$D,E{x},Vidas!$P:$P,"Ativo"))'
        F.cell(x, 13).value = f'=IF(A{x}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{x},Vidas!$B:$B,B{x},Vidas!$D:$D,E{x}))'
        F.cell(x, 14).value = f'=IF(A{x}="","",J{x}-M{x})'
        F.cell(x, 15).value = f'=IF(A{x}="","",IF(L{x}=0,"Sem vidas",IF(ABS(N{x})<=0.01,"OK","Divergência")))'
        F.cell(x, 16).value = f'=IF(OR(A{x}="",K{x}=""),"",IF(L{x}=K{x},"OK","Divergência"))'
        F.cell(x, 18).value = f'=IF(OR(A{x}="",Q{x}=""),"",IF(ABS(Q{x}-J{x}-N(AE{x}))<=0.01,"OK","Divergência"))'
        F.cell(x, 29).value = f'=IF(A{x}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{x}))'
        F.cell(x, 30).value = (f'=IF(A{x}="","",IF(AC{x}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{x},'
                               f'Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')
    # formatação das linhas vazias seguintes (até +300)
    for x in range(rF + 1, rF + 301):
        for col in range(1, 33):
            cs(F.cell(tplF, col if col < 32 else 7), F.cell(x, col))
        F.cell(x, 3).value = f'=IF(B{x}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{x},Empresas!$A:$A,0)),"ID não cadastrado"))'
        for col, fm in ((12, f'=IF(A{x}="","",COUNTIFS(Vidas!$A:$A,A{x},Vidas!$B:$B,B{x},Vidas!$D:$D,E{x},Vidas!$P:$P,"Ativo"))'),
                        (13, f'=IF(A{x}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{x},Vidas!$B:$B,B{x},Vidas!$D:$D,E{x}))'),
                        (14, f'=IF(A{x}="","",J{x}-M{x})'),
                        (15, f'=IF(A{x}="","",IF(L{x}=0,"Sem vidas",IF(ABS(N{x})<=0.01,"OK","Divergência")))'),
                        (16, f'=IF(OR(A{x}="",K{x}=""),"",IF(L{x}=K{x},"OK","Divergência"))'),
                        (18, f'=IF(OR(A{x}="",Q{x}=""),"",IF(ABS(Q{x}-J{x}-N(AE{x}))<=0.01,"OK","Divergência"))'),
                        (29, f'=IF(A{x}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{x}))'),
                        (30, f'=IF(A{x}="","",IF(AC{x}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{x},Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')):
            F.cell(x, col).value = fm
    for dv in F.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(re.sub(r'(\D)501\b', rf'\g<1>{rF + 300}', str(dv.sqref)))

    # ---------- aba Bradesco Pagas (histórico completo)
    nome = 'Bradesco Pagas'
    if nome in wb.sheetnames:
        del wb[nome]
    H = wb.create_sheet(nome, index=wb.sheetnames.index('Faturas') + 1)
    head = ['Competência (AAAA-MM)', 'Apólice', 'ID Empresa', 'Estipulante (como no portal)', 'CNPJ', 'Operadora', 'Subfatura',
            'Nº fatura (End./Fatura)', 'Emissão', 'Vencimento', 'Data de pagamento', 'Valor pago (R$)', 'IOF (R$)',
            'Valor da fatura sem IOF (R$)', 'Na aba Faturas?', 'Salvo em']
    hdr = F.cell(1, 1)
    for c, v in enumerate(head, 1):
        cell = H.cell(1, c); cs(hdr, cell); cell.value = v
    blue = Font(color='FF0000FF', size=10); gray = PatternFill('solid', fgColor='FFEEF1F4')
    money = '"R$ "#,##0.00;"-R$ "#,##0.00;\\-'
    for i, r in enumerate(rows, 2):
        vals = [r['comp'], str(r['apolice']), None, r['estipulante'], r['cnpj'], r['op'], f"{r['subfatura']:04d}", r['nf'],
                d(r['emissao']), d(r['vencimento']), d(r['baixa']), r['premio'], r['iof'], None, None, proc]
        for c, v in enumerate(vals, 1):
            cell = H.cell(i, c); cell.value = v; cell.font = blue
            if c in (9, 10, 11, 16):
                cell.number_format = 'dd/mm/yyyy'
            if c in (12, 13, 14):
                cell.number_format = money
        H.cell(i, 3).value = (f'=IFERROR(INDEX(Empresas!$A:$A,MATCH(B{i},Empresas!$E:$E,0)),"")')
        H.cell(i, 14).value = f'=L{i}-M{i}'
        H.cell(i, 15).value = f'=IF(COUNTIF(Faturas!$E:$E,H{i})>0,"Sim","Não")'
        for c in (3, 14, 15):
            H.cell(i, c).fill = gray; H.cell(i, c).font = Font(size=10)
    widths = [12, 10, 9, 28, 19, 15, 9, 14, 11, 11, 11, 14, 11, 14, 10, 11]
    for c, w in enumerate(widths, 1):
        H.column_dimensions[L(c)].width = w
    H.freeze_panes = 'A2'
    H.auto_filter.ref = f'A1:{L(len(head))}{len(rows) + 1}'
    H.sheet_properties.tabColor = 'CC092F'

    ampliar_faturamento(FT, N_EMP)
    wb.save(xout)
    print(json.dumps(dict(historico=len(rows), faturas=n_fat, ja_existiam=pulados, novas_empresas=len(novas),
                          primeira=novas[:1], ultima=novas[-1:], ultima_linha_faturas=rF), ensure_ascii=False))


if __name__ == '__main__':
    main(*sys.argv[1:])

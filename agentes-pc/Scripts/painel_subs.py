# -*- coding: utf-8 -*-
"""Painel (Central) - faturas com subestipulantes.
Uso: python painel_subs.py <pasta com os docs atuais da coleção faturas (AAAA-MM_<op>.json ou fat_AAAA-MM_<op>.json)> <pasta de saída> [AAAA-MM-DD de hoje]
(o arquivo de saída tem o mesmo nome do de entrada)
- Amil, Amil Dental, Santa Helena e SulAmérica: refaz os itens a partir da planilha, um item por empresa principal no mês
  (soma das NFs das subs), com `subs` (quantidade) e `subsDet` (uma linha por sub, a principal primeiro).
- Porto e Bradesco: mantém os itens como estão e só acrescenta `subsDet` nas empresas que têm mais de uma sub.
subsDet: {cod, nome, cnpj, titulo, venc, valor, vidas, boleto, confValor, vidasSub, valorVidas}; titulo/venc/valor/boleto/confValor = null
quando a sub não tem NF própria no mês (vem na NF da principal). vidasSub/valorVidas = vidas ativas e prêmio da sub na aba Vidas.
Só lê a planilha (calcula as conferências em Python, sem depender de valores de fórmula salvos pelo Excel).
"""
import sys, os, json, glob, collections, datetime as dt
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(os.path.dirname(AQUI), 'Base_Faturas_Vidas.xlsx')
REFAZ = {'amil': 'Amil', 'amildental': 'Amil Dental', 'santahelena': 'Santa Helena', 'sulamerica': 'SulAmérica'}
SLUG = {'Porto Saúde': 'porto', 'Bradesco Saúde': 'bradesco', 'Bradesco Dental': 'bradescodental', **{v: k for k, v in REFAZ.items()}}


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def dbr(v):
    return v.strftime('%d/%m/%Y') if isinstance(v, (dt.date, dt.datetime)) else ''


def main(src, out, hoje=None):
    hoje = dt.datetime.fromisoformat(hoje) if hoje else dt.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    wb = openpyxl.load_workbook(BASE, read_only=True)
    emp = {r[0]: r for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0]}
    porcod = {(r[3], str(r[5] or r[4]).replace(' 0', '')): r for r in emp.values()}
    subs = collections.defaultdict(collections.OrderedDict)
    for r in wb['Subestipulantes'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[2] is not None:
            subs[r[0]].setdefault(str(r[2]), dict(nome=r[4] or r[6], cnpj=r[5], curto=r[6]))
    vnf = collections.defaultdict(lambda: [0, 0.0, 0]); vsub = collections.defaultdict(lambda: [0, 0.0])
    for r in wb['Vidas'].iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        val = num(r[18]) + num(r[19]) + sum(num(r[c]) for c in (20, 21, 22, 23, 31)) + num(r[24])
        at = r[15] in ('Ativo', None)
        k = (r[0], r[1], str(r[3])); vnf[k][0] += at; vnf[k][1] += val; vnf[k][2] += 1
        if r[29] is not None:
            s = (r[0], r[1], str(r[29])); vsub[s][0] += at; vsub[s][1] += val
    fats = collections.defaultdict(list)
    for r in wb['Faturas'].iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[0]) >= '2026-01' and r[3] in SLUG:
            fats[(r[0], SLUG[r[3]])].append(r)

    def nf_info(r):
        comp, eid, nf = r[0], r[1], str(r[4])
        at, soma, nl = vnf.get((comp, eid, nf), [0, 0.0, 0])
        tot = num(r[9])
        conf = 'Sem vidas' if not nl else ('OK' if abs(soma - tot) < 0.015 else 'Divergência')
        confv = '' if r[10] is None else ('OK' if at == r[10] else 'Divergência')
        bol = ('Em aberto' if r[6] and r[6] >= hoje else 'Pago') if r[16] is not None or r[18] is not None else 'Sem boleto'
        return dict(titulo=nf, venc=dbr(r[6]), valor=round(tot, 2), vidas=r[10] if r[10] is not None else at, ativos=at,
                    boleto=bol, confValor=conf, confVidas=confv, soma=soma)

    def nf_cod(r, princ_cod):
        """Contrato/sub que recebe a NF."""
        aa = str(r[26] or '')
        if r[3] == 'SulAmérica':
            return aa if (aa and ',' not in aa and 'Fatura própria' in str(r[24] or '')) else princ_cod
        return aa.split(',')[0].strip() if aa else princ_cod

    def pcod(eid):
        """Código da sub principal: o contrato/empresa da principal, ou a 1ª sub cadastrada (Bradesco: 0001)."""
        e = emp[eid]; pc = str(e[5] or e[4]).replace(' 0', '') if e[3] == 'SulAmérica' else str(e[5] or e[4])
        return pc if (pc in subs[eid] or not subs[eid]) else next(iter(subs[eid]))

    def subs_det(comp, eid, rows):
        e = emp[eid]; pc = pcod(eid)
        por = collections.defaultdict(list)
        for r in rows: por[nf_cod(r, pc)].append(r)
        cods = list(subs[eid]) or [pc]
        if pc in cods: cods.remove(pc)
        cods = [pc] + cods + [c for c in por if c not in cods and c != pc]
        det = []
        for c in cods:
            s = subs[eid].get(c, {})
            vs = vsub.get((comp, eid, c), [0, 0.0])
            rs = por.get(c, [])
            ex = porcod.get((e[3], c))  # empresa antiga da sub (contrato que virou sub): razão social e CNPJ completos
            d = dict(cod=c, nome=(ex[1] if ex else None) or s.get('nome') or (e[1] if c == pc else None),
                     cnpj=(ex[2] if ex else None) or s.get('cnpj') or (e[2] if c == pc else None),
                     titulo=None, venc=None, valor=None, vidas=vs[0], boleto=None, confValor=None, vidasSub=vs[0], valorVidas=round(vs[1], 2))
            if rs:
                ii = [nf_info(r) for r in rs]
                d.update(titulo=', '.join(i['titulo'] for i in ii), venc=ii[0]['venc'], valor=round(sum(i['valor'] for i in ii), 2),
                         vidas=sum(i['vidas'] or 0 for i in ii) if len(subs[eid]) > 1 or e[3] != 'Porto Saúde' else vs[0],
                         boleto='Em aberto' if any(i['boleto'] == 'Em aberto' for i in ii) else ii[0]['boleto'],
                         confValor='OK' if all(i['confValor'] == 'OK' for i in ii) else ('Sem vidas' if all(i['confValor'] == 'Sem vidas' for i in ii) else 'Divergência'))
            det.append(d)
        return det

    os.makedirs(out, exist_ok=True); escritos = []
    for p in sorted(glob.glob(os.path.join(src, '*.json'))):
        nome = os.path.basename(p); doc_id = nome[:-5].removeprefix('fat_'); comp, slug = doc_id.split('_', 1)
        doc = json.load(open(p, encoding='utf-8')); doc = doc.get('data', doc)
        rows = fats.get((comp, slug), [])
        if slug in REFAZ:
            por = collections.OrderedDict()
            for r in sorted(rows, key=lambda r: r[1]): por.setdefault(r[1], []).append(r)
            itens = []
            for eid, rs in por.items():
                e = emp[eid]; pc = pcod(eid)
                ii = [nf_info(r) for r in rs]
                prin = next((i for r, i in zip(rs, ii) if nf_cod(r, pc) == pc), ii[0])
                det = subs_det(comp, eid, rs)
                it = dict(eid=eid, empresa=(e[1] or '')[:40], titulo=prin['titulo'], venc=prin['venc'], valor=round(sum(i['valor'] for i in ii), 2),
                          vidas=sum(i['vidas'] or 0 for i in ii), ativos=sum(i['ativos'] for i in ii), subs=max(len(det), 1), nfs=len(ii),
                          boleto='Em aberto' if any(i['boleto'] == 'Em aberto' for i in ii) else prin['boleto'],
                          confValor='OK' if all(i['confValor'] == 'OK' for i in ii) else ('Sem vidas' if all(i['confValor'] == 'Sem vidas' for i in ii) else 'Divergência'),
                          confVidas='OK' if all(i['confVidas'] in ('OK', '') for i in ii) else 'Divergência')
                if len(det) > 1: it['subsDet'] = det
                if 'confirmar' in str(e[9] or '') and 'Agrupamento a confirmar' in str(e[9]): it['grupoPendente'] = 'a confirmar'
                itens.append(it)
            doc['itens'] = itens
        else:
            por = collections.defaultdict(list)
            for r in rows: por[r[1]].append(r)
            feito = set()  # Bradesco tem um item por NF: o detalhe vai só no 1º item da empresa
            for it in doc.get('itens', []):
                eid = it.get('eid')
                if eid in subs and len(subs[eid]) > 1 and por.get(eid) and eid not in feito:
                    feito.add(eid)
                    it['subsDet'] = subs_det(comp, eid, por[eid]); it['subs'] = len(it['subsDet'])
        doc['atualizadoEm'] = dt.datetime.now().astimezone().replace(microsecond=0).isoformat()
        json.dump(doc, open(os.path.join(out, nome), 'w', encoding='utf-8'), ensure_ascii=False)
        escritos.append((doc_id, len(doc.get('itens', [])), sum(1 for i in doc.get('itens', []) if i.get('subsDet'))))
    for x in escritos: print(*x)


if __name__ == '__main__':
    main(*sys.argv[1:])

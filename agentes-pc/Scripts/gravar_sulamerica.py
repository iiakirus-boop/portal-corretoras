# -*- coding: utf-8 -*-
"""Agente Faturas SulAmérica - gravação.
Grava na Base_Faturas_Vidas.xlsx o JSON gerado por faturas_sulamerica.py, copia os PDFs para Faturas/AAAA-MM/SulAmerica/
e gera os JSON do painel (Scripts/_painel/fat_AAAA-MM_sulamerica.json).
Uso: python3 gravar_sulamerica.py <sul.json> <planilha entrada> <planilha saida> <pasta FATURAS> <AAAA-MM-DD> [--sem-copia] [--simular]
     --simular = não salva a planilha, não copia arquivos e não gera o painel (só mostra o que faria).
Regras:
- Uma empresa (aba Empresas) por código de empresa SulAmérica que recebe o boleto (ex.: 8J3U6). Código estipulante = "<empresa> 0",
  igual ao da aba Comissões (extrato de comissão); Subestipulante/contrato = código da empresa; Apólice = nº da apólice do relatório.
- Faturamento agrupado (um boleto para a apólice inteira, ex.: GABB 90809, nº documento terminando em 000000000): as outras empresas
  da mesma apólice e vencimento vão para Subestipulantes e as vidas delas entram na fatura do boleto.
- Competência = mês do vencimento do boleto "Fatura Mensal" (bate com o MM-AA que a equipe põe no nome dos arquivos).
  Boleto avulso (2ª via/reemissão, vencimento novo): competência = MM-AA do nome do arquivo.
- Nº fatura = endosso (últimos 9 dígitos do nº do documento = "Nº fatura (Fat/Eds)" da aba Comissões).
  Boleto agrupado sem endosso: "LF <nº LF>". Relatório sem boleto: "REL <apólice> <AAAAMMDD do vencimento>".
- Valor total da fatura = valor líquido do boleto (sem IOF); IOF na coluna IOF; Valor do boleto = valor total cobrado.
  Sem boleto: Total Geral do relatório.
- Vidas: uma linha por segurado do "Relatório de Faturamento" / "Relação de Segurados Ativos". Prêmio = prêmio do segurado.
  Nº seguro (família) = código de identificação sem os 4 últimos dígitos; nº dependente = 3 dígitos seguintes.
- Fatura já gravada (mesma competência + empresa + nº fatura) não é gravada de novo: só compara valor e relata.
- Nunca apaga nada; arquivos são copiados (não movidos) e não sobrescreve cópia que já exista.
"""
import json, sys, os, re, copy, shutil, calendar
from datetime import datetime
from collections import defaultdict, OrderedDict, Counter
import openpyxl
from agrupar_subestipulantes import DE  # subs vinculadas manualmente à empresa principal (se houver)
from shaddai_comum import d, br, dbr, cs, last_row, fim_mes, mmaa, estender, NEMP, SUBFORM  # módulo comum
from openpyxl.worksheet.cell_range import MultiCellRange

OP = 'SulAmérica'
AGENTE = 'Faturas SulAmérica'
PASTA = 'SulAmerica'
PROD = 'SULAMERICA SAUDE'
PAREN = {'TITULAR': 'Titular', 'CÔNJUGE': 'Cônjuge', 'CONJUGE': 'Cônjuge', 'CONJ/COMP.': 'Cônjuge', 'COMPANHEIRO (A)': 'Companheiro(a)',
         'FILHO(A)': 'Filho(a)', 'FILHOS': 'Filho(a)', 'OUTROS': 'Outros'}


def grupo_de(rel):
    p = rel.split('/')
    if len(p) > 2 and re.match(r'\d\d-\d\d$', p[0]): return p[1]
    if len(p) > 2 and re.match(r'GABB', p[0]): return 'GABB'
    return None


def nome_curto_arquivo(rel):
    b = os.path.basename(rel)
    m = re.match(r'^[0-9A-Z]{5}\s*-\s*(.*?)\s+(?:FATURADOS|BOLETO|COPART|SUL\s+AMERICA)', b, re.I)
    if m and m.group(1).strip(' -'):
        return re.sub(r'\s+', ' ', m.group(1).strip(' -')).upper()
    return None


def preparar(docs, log):
    """Junta boletos e relatórios em faturas. Devolve (faturas, info por empresa principal)."""
    for x in docs:
        if x['tipo'] in ('ignorado', 'erro'):
            log.append((None, None, 'Extração', f'Arquivo não lido: {x["rel"]} ({x.get("motivo") or x.get("erro")}).',
                        'Ignorado pelo Agente Faturas SulAmérica.', 'Sim'))
    # 1) boletos sem repetição (mesma empresa + documento + vencimento = cópia do mesmo arquivo)
    bols = OrderedDict()
    for b in sorted([x for x in docs if x['tipo'] == 'boleto'], key=lambda x: (x['rel'].startswith('GABB 25-26'), 'JA ENVIADO' in x['rel'], x['rel'])):
        k = (b['empresa'], b['documento'], b['venc'], b['valor'])
        if k in bols:
            bols[k].setdefault('copias', []).append(b['rel']); continue
        bols[k] = b
    mensal = [b for b in bols.values() if b['layout'] == 'fatura_mensal']
    end_mensal = {(b['empresa'], b['endosso']): b for b in mensal if b['endosso']}
    for b in mensal:
        b['comp'] = b['venc'][:7]
        if b.get('mes_arquivo') and b['mes_arquivo'] != b['comp']:
            log.append((b['comp'], b['empresa'], 'Conferência',
                        f'Boleto {b["rel"]} tem {mmaa(b["mes_arquivo"])} no nome, mas vence em {dbr(b["venc"])}.',
                        f'Gravado na competência do vencimento ({b["comp"]}).', 'Sim'))
    avulsos = []
    for b in [b for b in bols.values() if b['layout'] == 'avulso']:
        tw = end_mensal.get((b['empresa'], b['endosso']))
        if tw:
            tw.setdefault('segunda_via', []).append(b)
            continue
        b['comp'] = b.get('mes_arquivo') or b['venc'][:7]
        avulsos.append(b)
    # 2) relatórios sem repetição
    rels = OrderedDict()
    for f in sorted([x for x in docs if x['tipo'] == 'fatura'], key=lambda x: (x['rel'].startswith('GABB 25-26'), 'JA ENVIADO' in x['rel'], x['rel'])):
        k = (f['empresa'], f['apolice'], f['venc'], f['total'])
        if k in rels:
            rels[k].setdefault('copias', []).append(f['rel']); continue
        rels[k] = f
    # 3) fatura = boleto (+ relatórios)
    fats = []
    por_venc = {(b['empresa'], b['venc']): b for b in mensal}
    por_comp_av = {(b['empresa'], b['comp']): b for b in avulsos}
    agrup = defaultdict(list)
    for b in mensal:
        if not b['endosso']: agrup[(b['apl5'], b['venc'])].append(b)
    for b in mensal + avulsos:
        b['rels'] = []; b['subs'] = OrderedDict()
    sem_boleto = []
    for r in rels.values():
        comp = r['venc'][:7]
        b = por_venc.get((r['empresa'], r['venc'])) or por_comp_av.get((r['empresa'], comp))
        if b:
            b['rels'].append(r); continue
        ag = agrup.get((r['apl5'], r['venc']))
        if ag:
            ag[0]['rels'].append(r); ag[0]['subs'][r['empresa']] = r; continue
        sem_boleto.append(r)
    for b in mensal + avulsos:
        prin = next((r for r in b['rels'] if r['empresa'] == b['empresa']), None)
        if b['endosso']: nf = b['endosso']
        elif b.get('lf'): nf = f'LF {int(b["lf"])}'
        elif b.get('lancamento'): nf = f'LANC {b["lancamento"]}'
        else: nf = 'DOC ' + (b['documento'] or '')
        fats.append(dict(principal=b['empresa'], comp=b['comp'], nf=nf, boleto=b, rels=b['rels'], subs=list(b['subs']), rel_prin=prin,
                         venc=b['venc'] if b['layout'] == 'fatura_mensal' else ((prin or {}).get('venc') or b['venc']),
                         emissao=(prin or {}).get('emissao') or b.get('emissao'),
                         periodo=b['periodo'] if b['periodo'][0] else ((prin or {}).get('periodo') or (None, None))))
    for r in sem_boleto:
        fats.append(dict(principal=r['empresa'], comp=r['venc'][:7], nf=f'REL {r["apolice"]} {r["venc"].replace("-", "")}', boleto=None, rels=[r], subs=[],
                         rel_prin=r, venc=r['venc'], emissao=r['emissao'], periodo=r['periodo']))
    # 4) empresas que aparecem como subestipulante de faturamento agrupado ficam sempre sob a principal do grupo
    sub_de = {}
    for f in fats:
        for s in f['subs']: sub_de[s] = f['principal']
    for f in fats:
        if f['principal'] in sub_de:
            f['sub_de_grupo'] = f['principal']; f['principal'] = sub_de[f['principal']]
    # 5) dados por empresa principal
    info = defaultdict(lambda: dict(comps=[], nomes=Counter(), grupo=None, subs=OrderedDict(), apolices=Counter()))
    for f in sorted(fats, key=lambda x: x['comp']):
        i = info[f['principal']]
        i['comps'].append(f['comp']); i['venc'] = f['venc']
        b = f['boleto'] or {}
        for x in ([b] if b else []) + f['rels']:
            cod = x.get('empresa')
            nc = nome_curto_arquivo(x['rel'])
            if nc and cod == f['principal']: i['nomes'][nc] += 1
            g = grupo_de(x['rel'])
            if g and not i['grupo']: i['grupo'] = re.sub(r'\s+\d\d-\d\d$', '', g)
            if cod and cod != f['principal']:
                s = i['subs'].setdefault(cod, dict(nomes=Counter(), razao=None, apolice=None))
                if nc: s['nomes'][nc] += 1
                s['razao'] = x.get('razao') or s['razao']; s['apolice'] = x.get('apolice') or s['apolice']
        if f['rel_prin'] and f['rel_prin']['empresa'] == f['principal']:
            i['razao'] = f['rel_prin']['razao'] or i.get('razao'); i['apolices'][f['rel_prin']['apolice']] += 1
        if b.get('razao') and not i.get('razao'): i['razao'] = b['razao']
        if b.get('cnpj'): i['cnpj'] = b['cnpj']
        if not i['apolices'] and b.get('apl5'): i['apl5'] = b['apl5']
    return fats, info


def cadastrar(wb, info, proc):
    E, S = wb['Empresas'], wb['Subestipulantes']
    rE = last_row(E); tplE = rE
    nid = max(int(E.cell(r, 1).value[1:]) for r in range(2, rE + 1) if E.cell(r, 1).value)
    ids = {}
    for r in range(2, rE + 1):
        if E.cell(r, 4).value == OP and E.cell(r, 6).value:
            ids[str(E.cell(r, 6).value)] = E.cell(r, 1).value; info[str(E.cell(r, 6).value)]['nc'] = E.cell(r, 11).value
    novas = []
    for princ in sorted(info, key=lambda k: (info[k]['grupo'] or 'ZZZ', k)):
        if princ in ids: continue
        i = info[princ]
        if rE >= NEMP + 1:
            raise SystemExit(f'Aba Empresas cheia ({NEMP} empresas cobertas pela aba Faturamento). Ampliar antes de cadastrar.')
        nid += 1; eid = f'E{nid:03d}'; ids[princ] = eid; rE += 1
        nc = i['nomes'].most_common(1)[0][0] if i['nomes'] else re.sub(r'\s+', ' ', (i.get('razao') or princ)).upper()[:25]
        i['nc'] = nc
        ult = max(i['comps'])
        apol = i['apolices'].most_common(1)[0][0] if i['apolices'] else i.get('apl5')
        obs = (f'Cadastrada pelo Agente Faturas SulAmérica a partir dos boletos/relatórios da pasta ({proc:%d/%m/%Y}). '
               f'Código estipulante = código da empresa SulAmérica + " 0" (igual ao extrato de comissão).')
        if i['grupo']: obs += f' Grupo/pasta: {i["grupo"]}.'
        if i['subs']: obs += f' Faturamento agrupado da apólice: {len(i["subs"])} subestipulantes (aba Subestipulantes).'
        obs += f' Faturas na pasta: {min(i["comps"])} a {ult}.'
        row = [eid, i.get('razao'), i.get('cnpj'), OP, f'{princ} 0', princ, apol,
               int(i['venc'][8:10]) if i.get('venc') else None, 'Sim' if ult >= '2026-09' else 'Não', obs, nc,
               f'=IF(A{rE}="","",COUNTIFS(Subestipulantes!$A:$A,A{rE}))']
        for c, v in enumerate(row, 1):
            cs(E.cell(tplE, c), E.cell(rE, c)); E.cell(rE, c).value = v
        novas.append(eid)
    rS = last_row(S); tplS = rS
    exist = {(S.cell(r, 1).value, str(S.cell(r, 3).value)) for r in range(2, rS + 1)}
    for princ, i in info.items():
        subs = OrderedDict([(princ, dict(nomes=Counter({i['nc']: 1}), razao=i.get('razao'), apolice=None))]); subs.update(i['subs'])
        for k, s in subs.items():
            if (ids[princ], k) in exist: continue
            exist.add((ids[princ], k)); rS += 1
            nc = i['nc'] if k == princ else (s['nomes'].most_common(1)[0][0] if s['nomes'] else re.sub(r'\s+', ' ', s['razao'] or k).upper()[:25])
            row = [ids[princ], f'{princ} 0', k, None, s['razao'], i.get('cnpj') if k == princ else None, nc, 'Sim', SUBFORM.format(r=rS),
                   'Empresa principal (recebe o boleto).' if k == princ else
                   f'Empresa SulAmérica faturada junto, no boleto agrupado da empresa {princ} (apólice {s["apolice"]}).']
            for c, v in enumerate(row, 1):
                cs(S.cell(tplS, c), S.cell(rS, c)); S.cell(rS, c).value = v
    estender(E, NEMP + 1, 12, tplE)
    estender(S, max(rS, S.max_row), 10, tplS)
    return ids, novas


def nome_sub(info, princ, cod):
    if cod == princ: return info[princ]['nc']
    s = info[princ]['subs'].get(cod)
    return (s['nomes'].most_common(1)[0][0] if s and s['nomes'] else cod)


def gravar_faturas(wb, fats, info, ids, proc, log):
    F, V = wb['Faturas'], wb['Vidas']
    rF = last_row(F); tplF = rF
    rV = last_row(V); tplV = rV
    ja = {}
    for r in range(2, rF + 1):
        ja[(F.cell(r, 1).value, F.cell(r, 2).value, str(F.cell(r, 5).value))] = r
    res = Counter(); copias = []; gravadas = []
    for f in sorted(fats, key=lambda x: (x['comp'], ids[x['principal']], x['nf'])):
        eid = DE.get(ids[f['principal']], ids[f['principal']]); comp = f['comp']; nf = f['nf']; b = f['boleto'] or {}
        rp = f['rel_prin'] or {}
        soma_rel = round(sum(r['total'] or 0 for r in f['rels']), 2)
        vinf = sum(r['vidas'] or 0 for r in f['rels']) if f['rels'] else None
        if b.get('layout') == 'fatura_mensal':
            total, iof = b.get('liquido') or b.get('valor'), b.get('iof')
        elif b:
            total, iof = (soma_rel or b.get('valor')), (rp.get('iof') if f['rels'] else None)
        else:
            total, iof = rp.get('total'), rp.get('iof')
        chave = (comp, eid, nf)
        if chave in ja:
            r0 = ja[chave]; v0 = F.cell(r0, 10).value
            res['ja_na_planilha'] += 1
            if v0 is not None and total is not None and abs(float(v0) - total) > 0.01:
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf} já estava na planilha (linha {r0}) com valor R$ {br(float(v0))}; '
                            f'o arquivo traz R$ {br(total)}.', 'Não gravada de novo (regra: só comparar). Verificar qual está certo.', 'Não'))
                res['ja_divergente'] += 1
            continue
        ja[chave] = None
        nc = info[f['principal']]['nc']
        pasta = f'Faturas/{comp}/{PASTA}/{f["principal"]} - {nc}'
        arq_f = arq_b = None
        for r in f['rels']:
            cod = r['empresa']
            dst = f'{pasta}/{cod} - {nome_sub(info, f["principal"], cod)} FATURADOS {PROD} {mmaa(comp)}.pdf'
            copias.append((r['arquivo'], dst))
            if r is f['rel_prin'] or (arq_f is None and cod == f['principal']): arq_f = dst
        if arq_f is None and f['rels']: arq_f = copias[-len(f['rels'])][1]
        if b:
            cb = f.get('sub_de_grupo') or f['principal']; ncb = nome_sub(info, f['principal'], cb)  # boleto próprio de subestipulante: nome dela
            arq_b = f'{pasta}/{cb} - {ncb} BOLETO {PROD} {mmaa(comp)}{" 2A VIA" if b["layout"] == "avulso" else ""}.pdf'
            copias.append((b['arquivo'], arq_b))
            for k, sv in enumerate(b.get('segunda_via', []), 1):
                copias.append((sv['arquivo'], f'{pasta}/{cb} - {ncb} BOLETO {PROD} {mmaa(comp)} 2A VIA{"" if k == 1 else " " + str(k)}.pdf'))
        obs = []
        if not b: obs.append('Só o relatório de faturamento (sem boleto na pasta).')
        elif b['layout'] == 'avulso':
            obs.append(f'Boleto avulso (2ª via/reemissão) da apólice {b.get("apolice")} endosso {b.get("endosso")} parcela {b.get("parcela")}, '
                       f'vencimento {dbr(b["venc"])}, R$ {br(b["valor"])} (pode incluir multa/juros).')
        else: obs.append('Boleto Fatura Mensal' + (' (faturamento agrupado da apólice)' if b.get('agrupado') or not b.get('endosso') else '') + '.')
        if b and not f['rels']: obs.append('Sem relatório de faturamento na pasta (sem vidas).')
        if f['subs']: obs.append(f'{len(f["rels"])} relatórios (empresa principal + {len(f["subs"])} subestipulantes) somam R$ {br(soma_rel)}.')
        if f.get('sub_de_grupo'): obs.append(f'Fatura própria da empresa {f["sub_de_grupo"]} (antes do faturamento agrupado).')
        if b.get('segunda_via'):
            obs.append('2ª via na pasta: ' + '; '.join(f'venc. {dbr(s["venc"])} R$ {br(s["valor"])}' for s in b['segunda_via']) + '.')
        if b.get('layout') == 'fatura_mensal' and f['rels'] and abs(soma_rel - total) > 0.015:
            dif = round(total - soma_rel, 2); cop = b.get('copart') or 0
            pre = round(sum(v for k, v in (b.get('itens') or {}).items() if k.lower().startswith('prêmio saúde')), 2)
            outros = '; '.join(f'{k} R$ {br(v)}' for k, v in (b.get('itens') or {}).items() if not k.lower().startswith('prêmio saúde'))
            if pre and -0.015 <= pre - soma_rel <= (iof or 0) + 0.015 and dif > 0:
                obs.append(f'Prêmio Saúde do boleto (R$ {br(pre)}, com IOF) confere com os relatórios (R$ {br(soma_rel)}); '
                           f'a diferença de R$ {br(dif)} para o boleto é {outros} (sem detalhe por vida).')
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf}: prêmio dos relatórios R$ {br(soma_rel)} confere com o Prêmio Saúde do boleto; '
                            f'boleto líquido R$ {br(total)} inclui também {outros}.',
                            'Gravada com o valor do boleto; a Conferência valor da planilha mostra essa parte como diferença (sem detalhe por vida).', 'Sim'))
            elif cop and abs(dif - cop) < 0.015:
                obs.append(f'Boleto inclui coparticipação de R$ {br(cop)} (não detalhada por vida): boleto líquido R$ {br(total)} = relatórios R$ {br(soma_rel)} + copart.')
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf}: boleto líquido R$ {br(total)} = prêmio dos relatórios R$ {br(soma_rel)} + coparticipação R$ {br(cop)}.',
                            'Gravada com o valor do boleto; a Conferência valor da planilha mostra a copart como diferença (sem detalhe por vida).', 'Sim'))
            else:
                obs.append(f'Boleto (líquido) R$ {br(total)} - relatórios R$ {br(soma_rel)} = R$ {br(dif)}' + (f' (copart no boleto R$ {br(cop)})' if cop else '') + '.')
                itens = '; '.join(f'{k} R$ {br(v)}' for k, v in (b.get('itens') or {}).items()) or 'composição não lida'
                log.append((comp, f['principal'], 'Conferência', f'Fatura {nf}: boleto líquido R$ {br(total)} ≠ soma dos relatórios R$ {br(soma_rel)} '
                            f'(diferença R$ {br(dif)}; boleto: {itens}).', 'Gravada com o valor do boleto. Verificar se falta relatório de alguma '
                            'subestipulante ou se há acerto/pró-rata no mês.', 'Não'))
        rel_txt = '; '.join(['SUL AMERICA/' + x['rel'] for x in ([b] if b else []) + f['rels']][:3])
        obs.append(f'Arquivo original: {rel_txt}{" ..." if len(f["rels"]) + bool(b) > 3 else ""}.')
        per = f['periodo'] or (None, None)
        ini = d(per[0]) or datetime(int(comp[:4]), int(comp[5:7]), 1)
        fim = d(per[1]) or fim_mes(comp)
        rF += 1; r = rF
        subs_lst = list(f['subs']) + ([f['sub_de_grupo']] if f.get('sub_de_grupo') else [])
        row = {1: comp, 2: eid, 4: OP, 5: nf, 6: d(f.get('emissao')), 7: d(f.get('venc')), 8: ini, 9: fim,
               10: total, 11: vinf, 17: b.get('valor'), 19: b.get('nosso'), 20: b.get('linha'),
               21: arq_f, 22: None, 23: arq_b, 24: proc, 25: ' '.join(obs), 26: f'{f["principal"]} 0', 27: ', '.join(subs_lst) or None,
               31: iof}
        for col in range(1, 33):
            cs(F.cell(tplF, col), F.cell(r, col)); F.cell(r, col).value = row.get(col)
        F.cell(r, 3).value = f'=IF(B{r}="","",IFERROR(INDEX(Empresas!$B:$B,MATCH(B{r},Empresas!$A:$A,0)),"ID não cadastrado"))'
        F.cell(r, 12).value = f'=IF(A{r}="","",COUNTIFS(Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r},Vidas!$P:$P,"Ativo"))'
        F.cell(r, 13).value = f'=IF(A{r}="","",SUMIFS(Vidas!$AA:$AA,Vidas!$A:$A,A{r},Vidas!$B:$B,B{r},Vidas!$D:$D,E{r}))'
        F.cell(r, 14).value = f'=IF(A{r}="","",J{r}-M{r})'
        F.cell(r, 15).value = f'=IF(A{r}="","",IF(L{r}=0,"Sem vidas",IF(ABS(N{r})<=0.01,"OK","Divergência")))'
        F.cell(r, 16).value = f'=IF(OR(A{r}="",K{r}=""),"",IF(L{r}=K{r},"OK","Divergência"))'
        F.cell(r, 18).value = f'=IF(OR(A{r}="",Q{r}=""),"",IF(ABS(Q{r}-J{r}-N(AE{r}))<=0.01,"OK","Divergência"))'
        F.cell(r, 29).value = f'=IF(A{r}="","",SUMIFS(Comissões!$O:$O,Comissões!$J:$J,E{r}))'
        F.cell(r, 30).value = (f'=IF(A{r}="","",IF(AC{r}=0,"Pendente",IF(COUNTIFS(Comissões!$J:$J,E{r},'
                               f'Comissões!$X:$X,"Divergência")>0,"Divergência","Recebida")))')
        res['faturas'] += 1; res['sem_boleto'] += not b; res['sem_relatorio'] += bool(b) and not f['rels']
        soma = 0.0; n_v = 0
        for rr in f['rels']:
            for x in rr['linhas']:
                rV += 1; n_v += 1; soma += x['valor']
                tit = x['parentesco'] == 'TITULAR'
                vals = {1: comp, 2: eid, 3: OP, 4: nf, 5: x['familia'], 6: x['dep'], 7: x['nome'], 8: x['cpf'], 9: d(x['nasc']), 10: x['idade'],
                        11: PAREN.get(x['parentesco'], x['parentesco'].title()), 12: 'Titular' if tit else 'Dependente', 13: x['plano_cod'], 14: x['plano'],
                        15: x['mat'], 16: 'Ativo', 17: d(x['inicio']), 19: x['valor'], 20: 0, 21: 0, 22: 0, 23: 0, 24: 0, 25: 0,
                        29: f'Cód. identificação {x["cod"]}', 30: rr['empresa'], 32: 0}
                for col in range(1, 33):
                    cs(V.cell(tplV, col if col < 32 else 21), V.cell(rV, col)); V.cell(rV, col).value = vals.get(col)
                V.cell(rV, 26).value = f'=IF(A{rV}="","",N(U{rV})+N(V{rV})+N(W{rV})+N(X{rV})+N(AF{rV}))'
                V.cell(rV, 27).value = f'=IF(A{rV}="","",N(S{rV})+N(T{rV})+N(Z{rV})+N(Y{rV}))'
                V.cell(rV, 28).value = f'=IF(A{rV}="","",IF(P{rV}="Excluído","Exclusão",IF(N(T{rV})>0,"Inclusão (pró-rata)","-")))'
                res['vidas'] += 1
        soma = round(soma, 2)
        ok_v = bool(f['rels']) and abs(soma - (total or 0)) < 0.015
        ok_n = vinf is not None and n_v == vinf
        if ok_v and ok_n and b: res['ok'] += 1; st = 'OK'
        else: res['pend'] += 1; st = 'Pendência'
        if not b:
            log.append((comp, f['principal'], 'Download', f'Sem boleto na pasta para a fatura de vencimento {dbr(f["venc"])} ({rp.get("rel")}).',
                        'Fatura gravada com o Total Geral do relatório (sem nosso número/linha digitável).', 'Não'))
        elif not f['rels']:
            log.append((comp, f['principal'], 'Download', f'Só boleto, sem relatório de faturamento (fatura {nf}, {b["rel"]}).',
                        'Fatura gravada com o valor do boleto; sem vidas.', 'Não'))
        if f['rels'] and not ok_n:
            log.append((comp, f['principal'], 'Conferência', f'Vidas lidas ({n_v}) ≠ "Total de Segurados" ({vinf}) na fatura {nf}.', 'Verificar o PDF.', 'Não'))
        gravadas.append(dict(comp=comp, eid=eid, principal=f['principal'], nf=nf, venc=f['venc'], total=total, vinf=vinf, vidas=n_v, soma=soma,
                             subs=len(f['subs']), st=st, boleto=bool(b)))
    estender(F, max(rF, F.max_row), 32, tplF)
    estender(V, max(rV + 500, V.max_row), 32, tplV)
    for r in range(tplV + 1, V.max_row + 1):
        cs(V.cell(tplV, 21), V.cell(r, 32))
    return res, copias, gravadas, rF, rV


def registrar(wb, res, log, info, ids, novas, comps, proc, backup, inicio):
    AG, EX, LG, EXH = wb['Agentes'], wb['Execuções'], wb['Log de Erros'], wb['Execução']
    ra = last_row(AG)
    if not any(AG.cell(r, 1).value == AGENTE for r in range(2, ra + 1)):
        ra += 1
        for c in range(1, 13):
            cs(AG.cell(2, c), AG.cell(ra, c))
            v = AG.cell(2, c).value
            AG.cell(ra, c).value = re.sub(r'(?<=[A-Z])2\b', str(ra), v) if isinstance(v, str) and v.startswith('=') else None
        for c, v in enumerate([AGENTE, 'Lê os boletos e relatórios de faturamento (PDF) da SulAmérica Saúde que a equipe coloca na pasta e grava Faturas + Vidas',
                               OP, 'Empresas, Subestipulantes, Faturas, Vidas, Log de Erros', 'Mensal', 10,
                               'Scripts/faturas_sulamerica.py + Scripts/gravar_sulamerica.py', 'Ativo'], 1):
            AG.cell(ra, c).value = v
    pend = sum(1 for x in log if x[5] == 'Não')
    re_ = last_row(EX) + 1
    for c in range(1, 14): cs(EX.cell(re_ - 1, c), EX.cell(re_, c))
    resumo = (f'{res["sem_boleto"]} faturas só com relatório (sem boleto); {res["sem_relatorio"]} só com boleto (sem relatório); '
              f'{pend} pendências no Log de Erros') if pend else ''
    for c, v in enumerate([AGENTE, f'{comps[0]} a {comps[-1]}' if comps else '', inicio, datetime.now().replace(microsecond=0),
                           '⚠️ OK com pendência' if pend else '✅ OK', res['faturas'], res['ok'], res['pend'], 0, resumo,
                           backup, f'Primeira carga SulAmérica a partir da pasta "boletos a enviar/SUL AMERICA": {res["faturas"]} faturas, '
                           f'{res["vidas"]} linhas de vidas, {len(novas)} empresas novas; {res["ja_na_planilha"]} já estavam na planilha.', 'Bruno'], 1):
        EX.cell(re_, c).value = v
    rl = last_row(LG); tplL = rl
    for comp, k, etapa, erro, acao, ok in log:
        rl += 1
        for c in range(1, 10): cs(LG.cell(tplL, c), LG.cell(rl, c))
        for c, v in enumerate([proc, comp, OP, ids.get(k) if k else None, etapa, erro, acao, ok, 'Agente Faturas'], 1):
            LG.cell(rl, c).value = v
    estender(LG, max(rl, LG.max_row), 9, tplL)
    rh = last_row(EXH, 2) + 1
    for c in range(2, 10): cs(EXH.cell(rh - 1, c), EXH.cell(rh, c))
    for c, v in enumerate([proc, f'{comps[0]} a {comps[-1]}' if comps else '', f'SulAmérica: {len(info)} empresas principais (arquivos da equipe)', res['faturas'],
                           res['ok'], res['pend'], f'{res["sem_boleto"]} sem boleto', 'Primeira carga SulAmérica Saúde a partir dos arquivos da equipe.'], 2):
        EXH.cell(rh, c).value = v


LEIA = ['AGENTE FATURAS SULAMÉRICA (SULAMÉRICA SAÚDE)',
        'Sem portal: a equipe salva os arquivos em "Área de Trabalho/boletos a enviar/SUL AMERICA/MM-AA/<empresa>/" e o agente lê a pasta inteira. Scripts: Scripts/faturas_sulamerica.py (leitura) e Scripts/gravar_sulamerica.py (gravação).',
        'Arquivos lidos (PDF): boleto "Fatura Mensal", boleto avulso (2ª via) e "Relatório de Faturamento de Segurados Ativos" / "Relação de Segurados Ativos". Copart (TXT/XLS), CSV/XLSX que repetem o PDF, IR, sinistralidade e seguro de vida são ignorados e vão para o Log de Erros.',
        'Empresas: uma linha por código de empresa SulAmérica que recebe o boleto. Código estipulante = "<empresa> 0" (igual ao extrato de comissão). Faturamento agrupado (ex.: GABB, apólice 90809): as outras empresas da apólice ficam em Subestipulantes.',
        'Competência = mês do vencimento do boleto (= MM-AA do nome dos arquivos). Nº fatura = endosso (o mesmo "Fat/Eds" do extrato de comissão); boleto agrupado sem endosso = "LF <nº>"; relatório sem boleto = "REL <apólice> <vencimento>".',
        'Valor total da fatura = valor líquido do boleto (sem IOF); IOF em coluna própria. Quando o boleto traz coparticipação, a diferença para a soma das vidas aparece como Divergência (copart não detalhada por vida).',
        'Vidas: uma linha por segurado do relatório (Nº seguro = código de identificação sem os 4 últimos dígitos). Prêmio = prêmio do segurado.',
        'Arquivos copiados para Faturas/AAAA-MM/SulAmerica/<empresa> - <NOME CURTO>/<empresa> - <NOME CURTO> FATURADOS | BOLETO SULAMERICA SAUDE <MM-AA>.pdf. O original fica na pasta da equipe.']


def leia_me(wb):
    ws = wb['Leia-me']
    if any(ws.cell(r, 2).value == LEIA[0] for r in range(1, ws.max_row + 1)):
        return
    hdr = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or '').startswith('AGENTE FATURAS BRADESCO'))
    r = last_row(ws, 2) + 1
    for i, t in enumerate(LEIA):
        r += 1
        cs(ws.cell(hdr if i == 0 else hdr + 1, 2), ws.cell(r, 2)); ws.cell(r, 2).value = t


def painel(gravadas, wb, ids, base_dir, proc):
    """JSON do painel (coleção faturas da Central), um por competência a partir de 2026-01."""
    emp = {r[0]: r[1] for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0]}
    docs = defaultdict(list)
    for g in gravadas:
        if g['comp'] < '2026-01': continue
        venc = d(g['venc'])
        docs[g['comp']].append(dict(eid=g['eid'], empresa=(emp.get(g['eid']) or '')[:40], titulo=g['nf'], venc=venc.strftime('%d/%m/%Y') if venc else '',
                                    valor=round(g['total'] or 0, 2), vidas=g['vinf'] if g['vinf'] is not None else g['vidas'], ativos=g['vidas'], subs=1 + g['subs'],
                                    boleto='Sem boleto' if not g['boleto'] else ('Em aberto' if venc and venc >= proc else 'Pago'),
                                    confValor=('Sem vidas' if not g['vidas'] else ('OK' if abs(g['soma'] - (g['total'] or 0)) < 0.015 else 'Divergência')),
                                    confVidas=('' if g['vinf'] is None else ('OK' if g['vidas'] == g['vinf'] else 'Divergência'))))
    out = os.path.join(base_dir, 'Scripts', '_painel'); os.makedirs(out, exist_ok=True)
    agora = datetime.now().astimezone().replace(microsecond=0).isoformat()
    for comp, it in sorted(docs.items()):
        json.dump(dict(comp=comp, operadora=OP, atualizadoEm=agora, parcial=False, pendentes=[], itens=sorted(it, key=lambda x: x['eid'])),
                  open(os.path.join(out, f'fat_{comp}_sulamerica.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    return sorted(docs)


def main(js, xin, xout, base_dir, proc, *opts):
    inicio = datetime.now().replace(microsecond=0)
    proc = datetime.fromisoformat(proc)
    if os.path.exists(os.path.join(os.path.dirname(os.path.abspath(xin)), '~$' + os.path.basename(xin))):
        raise SystemExit('A planilha está aberta no Excel (~$ encontrado). Feche e rode de novo.')
    docs = json.load(open(js, encoding='utf-8'))
    log = []
    fats, info = preparar(docs, log)
    wb = openpyxl.load_workbook(xin)
    ids, novas = cadastrar(wb, info, proc)
    res, copias, gravadas, rF, rV = gravar_faturas(wb, fats, info, ids, proc, log)
    comps = sorted({g['comp'] for g in gravadas})
    backup = f'Backups/Base_Faturas_Vidas antes Faturas SulAmerica {proc:%Y-%m-%d}.xlsx'
    registrar(wb, res, log, info, ids, novas, comps, proc, backup, inicio)
    leia_me(wb)
    n = 0; pain = []
    if '--simular' not in opts:
        wb.save(xout)
        if '--sem-copia' not in opts:
            for src, dst in copias:
                p = os.path.join(base_dir, dst)
                if not os.path.exists(p):
                    os.makedirs(os.path.dirname(p), exist_ok=True); shutil.copy2(src, p); n += 1
        pain = painel(gravadas, wb, ids, base_dir, proc)
    json.dump(dict(gravadas=gravadas, log=log, ids=ids), open(os.path.splitext(js)[0] + '_resultado.json', 'w', encoding='utf-8'), ensure_ascii=False, default=str)
    print(json.dumps(dict(res, novas_empresas=len(novas), ids=f'{novas[0]}..{novas[-1]}' if novas else '', copiados=n, a_copiar=len(copias),
                          log=len(log), log_pend=sum(1 for x in log if x[5] == 'Não'), ult_fatura=rF, ult_vida=rV, comps=comps, painel=len(pain)), ensure_ascii=False))


if __name__ == '__main__':
    main(*sys.argv[1:])

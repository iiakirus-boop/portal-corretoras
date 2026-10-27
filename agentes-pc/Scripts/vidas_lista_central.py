# -*- coding: utf-8 -*-
"""Agente Vidas -> Central: lista de beneficiários por empresa (coleção 'vidas' da Central) para uma operadora.
Uso: python vidas_lista_central.py [Operadora] [competência mínima AAAA-MM]
     padrão: SulAmérica, todas as competências. Gera _painel/vidas_<operadora>.json ({eid: documento}).
Documento: eid, comp (última), cnpj, diaVenc, hist{comp:{ativos, incl, excl, valor}}, vidas[beneficiários da última competência],
atualizadoEm. incl/excl = beneficiários (CPF, ou nome+nascimento) que entraram/saíram em relação à competência anterior.
A distribuição (coleção 'vidasdist') sai do vidas_central.py, que já cobre todas as operadoras.
Só lê a planilha; não grava nada nela.
"""
import sys, os, json, re, unicodedata, collections, datetime as dt
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(os.path.dirname(AQUI), 'Base_Faturas_Vidas.xlsx')


def slug(s):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower())


def iso(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (dt.date, dt.datetime)) else (v or None)


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else 0


def main(op='SulAmérica', cmin=''):
    wb = openpyxl.load_workbook(BASE, read_only=True)
    emp = {r[0]: r for r in wb['Empresas'].iter_rows(min_row=2, values_only=True) if r[0] and r[3] == op}
    por = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in wb['Vidas'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[2] == op and r[1] in emp and str(r[0]) >= cmin:
            por[r[1]][r[0]].append(r)
    agora = dt.datetime.now().astimezone().replace(microsecond=0).isoformat()
    docs = {}
    for eid, comps in sorted(por.items()):
        hist, ant = {}, None
        for comp in sorted(comps):
            at = [x for x in comps[comp] if x[15] in ('Ativo', None)]
            ks = {(x[7] or f'{x[6]}|{iso(x[8])}') for x in at}
            valor = sum(num(x[18]) + num(x[19]) + sum(num(x[c]) for c in (20, 21, 22, 23, 24, 31)) for x in comps[comp])
            hist[comp] = dict(ativos=len(at), incl=len(ks - ant) if ant is not None else 0, excl=len(ant - ks) if ant is not None else 0, valor=round(valor, 2))
            ant = ks
        ult = max(comps)
        vidas = []
        for x in comps[ult]:
            cop = round(sum(num(x[c]) for c in (20, 21, 22, 23, 31)), 2)
            vidas.append(dict(n=x[6], nasc=iso(x[8]), id=x[9], par=x[10], td='T' if x[11] == 'Titular' else 'D', fam=x[4], plano=x[13],
                              sit=x[15] or 'Ativo', ades=iso(x[16]), excl=iso(x[17]), premio=num(x[18]), pr=num(x[19]), cop=cop,
                              tot=round(num(x[18]) + num(x[19]) + cop + num(x[24]), 2), fat=str(x[3]), op=op, sub=x[29]))
        e = emp[eid]
        docs[eid] = dict(eid=eid, comp=ult, cnpj=e[2], diaVenc=e[7], hist=hist, vidas=vidas, atualizadoEm=agora)
    out = os.path.join(AQUI, '_painel', f'vidas_{slug(op)}.json')
    json.dump(docs, open(out, 'w', encoding='utf-8'), ensure_ascii=False)
    for k, d in docs.items():
        print(k, d['comp'], len(d['vidas']), 'vidas;', len(d['hist']), 'competências')
    print(out)


if __name__ == '__main__':
    main(*sys.argv[1:])

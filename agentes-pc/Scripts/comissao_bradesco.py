"""Agente Comissão Bradesco — importa o extrato de comissões da Bradesco Saúde (.xls do Portal de Negócios)
para a Base_Faturas_Vidas.xlsx, nas mesmas abas da Porto ("Pagamentos Comissão" e "Comissões").

Caminho no portal: Saúde → Comissões → Extrato de comissões → período → Buscar → Ver detalhes → aba SAÚDE S.A. → ".XLS"

Uso:  python comissao_bradesco.py <Base_Faturas_Vidas.xlsx> <arquivo.xls> [arquivo2.xls ...]

- O .xls traz uma linha por beneficiário/parcela. O agente agrupa por
  ordem de pagamento (coluna "NÚMERO DA FATURA" do portal) + fatura (sucursal + endosso) + taxa + tipo.
- Nº da fatura: o extrato traz só o endosso (18728); o agente acha na base a fatura 962318728 (sucursal + ... + endosso).
- Ordem de pagamento já gravada não é gravada de novo: o script compara o total e relata.
- Conferência: prêmio do extrato = valor da fatura sem IOF (aba Faturas ou "Bradesco Pagas") e comissão = prêmio × taxa.
- Tipo AN = agenciamento; CO = comissão (vitalício). Taxa 0,01% = lançamento por apólice (parcela 0).
- O .xls não traz tributações (INSS/IRRF/ISS) nem crédito líquido: essas colunas ficam em branco.
"""
import sys, os, re, datetime, collections
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from shaddai_comum import preparar, _prox, _w, MOEDA, DATA, CINZA, N, PAG, COM

SEGURADORA = "BRADESCO SAUDE S.A."
TIPOS = {"AN": "AN - Agenciamento", "CO": "CO - Comissão (vitalício)"}


def ler_xls(path):
    """O '.xls' do portal é um xlsx com outra extensão: abre pelo conteúdo."""
    with open(path, "rb") as f:
        wb = openpyxl.load_workbook(f, data_only=True)
    if "DET_FATURAS" not in wb.sheetnames:
        raise ValueError(f"{os.path.basename(path)} não é o extrato de comissões da Bradesco (falta a aba DET_FATURAS)")
    ws = wb["DET_FATURAS"]
    cab = [str(c or "").strip() for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    ix = {k: cab.index(k) for k in ("NÚMERO DA FATURA", "SUCURSAL", "DATA LANÇ", "LANÇ", "APÓLICE / SUBFATURA",
                                     "ENDOSSO", "RAMO", "SEGURADO", "PARCELA", "VL_PRÊMIO", "%", "VALOR",
                                     "TIPO DE REMUNERAÇÃO", "PROPOSTA")}
    linhas = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r[ix["NÚMERO DA FATURA"]]:
            continue
        g = lambda k: r[ix[k]]
        apl, sub = str(g("APÓLICE / SUBFATURA")).split("/")
        linhas.append(dict(
            op=str(g("NÚMERO DA FATURA")).strip(), suc=str(g("SUCURSAL")).strip(), data=g("DATA LANÇ"),
            lanc=str(g("LANÇ")).strip(), apolice=str(int(apl)), sub=f"{int(sub):04d}", endosso=str(g("ENDOSSO")).strip(),
            ramo=str(g("RAMO")).strip(), segurado=re.sub(r"\s+\d{2}$", "", str(g("SEGURADO") or "")).strip(),
            parcela=str(g("PARCELA")).strip(), premio=round(float(g("VL_PRÊMIO") or 0), 2),
            taxa=round(float(g("%") or 0), 4), valor=round(float(g("VALOR") or 0), 2),
            tipo=str(g("TIPO DE REMUNERAÇÃO") or "").strip()))
    return linhas


def _indice_faturas(wb):
    """nº fatura -> (competência, apólice) de tudo que a base conhece (Faturas + Bradesco Pagas)."""
    idx = {}
    for r in wb["Bradesco Pagas"].iter_rows(min_row=2, values_only=True):
        if r[7]:
            idx[str(r[7])] = (str(r[0] or ""), str(r[1] or ""))
    for r in wb["Faturas"].iter_rows(min_row=2, values_only=True):
        if r[4] and r[3] and "Bradesco" in str(r[3]):
            idx.setdefault(str(r[4]), (str(r[0] or ""), str(r[25] or "")))
    return idx


def fatura_base(l, idx):
    """O extrato traz só o endosso (5 dígitos: 18728); a fatura é 962318728. Procura na base a fatura da mesma
    sucursal que termina com o endosso, de preferência da mesma apólice e a mais recente.
    Sem endosso (prêmio de campanha, ex.: TOP RETENCAO, apólice 0): não há fatura, devolve ''."""
    if not l["endosso"].isdigit():
        return "", False
    end = f"{int(l['endosso']):05d}"
    cand = [f for f in idx if f.startswith(l["suc"]) and f.endswith(end) and len(f) == len(l["suc"]) + 6]
    cand.sort(key=lambda f: (idx[f][1] == l["apolice"], idx[f][0]), reverse=True)
    if cand:
        return cand[0], True
    return f"{l['suc']}3{end}", False


def _estipulantes(wb):
    nomes = {}
    for r in wb["Empresas"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[3] and "Bradesco" in str(r[3]) and r[4]:
            nomes[str(r[4])] = r[1]
    for r in wb["Bradesco Pagas"].iter_rows(min_row=2, values_only=True):
        if r[1] and str(r[1]) not in nomes:
            nomes[str(r[1])] = r[3]
    return nomes


def _formulas_bradesco(wc, k):
    """Linha da aba Comissões com conferência própria da Bradesco (fatura e IOF vêm de Faturas ou Bradesco Pagas)."""
    BP = "'Bradesco Pagas'"
    # Q e V também aqui, para linhas além das pré-formatadas pelo preparar()
    wc[f"Q{k}"] = (f"=IF(B{k}=\"\",\"\",IFERROR(INDEX(Faturas!$B:$B,MATCH(J{k},Faturas!$E:$E,0)),"
                   f"IFERROR(INDEX(Empresas!$A:$A,MATCH(I{k},Empresas!$E:$E,0)),\"\")))")
    wc[f"V{k}"] = f"=IF(OR(B{k}=\"\",U{k}=\"\"),\"\",M{k}-U{k})"
    for col in "QRSTUVWX":
        wc[f"{col}{k}"].fill = CINZA
    for col in "STUVW":
        wc[f"{col}{k}"].number_format = MOEDA
    wc[f"R{k}"] = (f"=IFERROR(INDEX(Faturas!$A:$A,MATCH(J{k},Faturas!$E:$E,0)),"
                   f"IFERROR(INDEX({BP}!$A:$A,MATCH(J{k},{BP}!$H:$H,0)),\"\"))")
    # S = valor da fatura SEM IOF (na aba Faturas a coluna J já é sem IOF; na Bradesco Pagas é pago − IOF)
    wc[f"S{k}"] = (f"=IFERROR(INDEX(Faturas!$J:$J,MATCH(J{k},Faturas!$E:$E,0)),"
                   f"IFERROR(INDEX({BP}!$L:$L,MATCH(J{k},{BP}!$H:$H,0))-INDEX({BP}!$M:$M,MATCH(J{k},{BP}!$H:$H,0)),\"\"))")
    wc[f"T{k}"] = (f"=IF(S{k}=\"\",\"\",IFERROR(INDEX(Faturas!$AE:$AE,MATCH(J{k},Faturas!$E:$E,0)),"
                   f"IFERROR(INDEX({BP}!$M:$M,MATCH(J{k},{BP}!$H:$H,0)),0)))")
    wc[f"U{k}"] = f"=IF(S{k}=\"\",\"\",ROUND(S{k}-N(AB{k}),2))"
    wc[f"W{k}"] = f"=ROUND(M{k}*N{k}/100,2)"
    wc[f"X{k}"] = (f"=IF(S{k}=\"\",\"Sem fatura na base\",IF(AND(ABS(V{k})<=MAX(1,0.0005*U{k}),"
                   f"ABS(W{k}-O{k})<=MAX(0.05,0.0001*O{k})),\"OK\",\"Divergência\"))")


def importar(wb, arquivos, pasta_rel):
    preparar(wb)
    wp, wc = wb[PAG], wb[COM]
    # cabeçalhos passam a valer para as duas seguradoras
    if wp["B1"].value == "Empresa Porto":
        wp["B1"] = "Seguradora"
    if wc["D1"].value == "Empresa Porto":
        wc["D1"] = "Seguradora"
        wc["T1"].comment = Comment("Porto: copart + despesas não cobertas (descontados do prêmio). Bradesco: IOF da fatura, só informativo — a coluna S já vem sem IOF.", "Agente Comissão")
    # Faturas: colunas AC/AD (comissão recebida / status) até a última fatura da base
    fat = wb["Faturas"]
    ultima = max(r for r in range(1, fat.max_row + 1) if fat[f"A{r}"].value not in (None, ""))
    for r in range(2, ultima + 1):
        if fat[f"AC{r}"].value is None:
            fat[f"AC{r}"] = f"=IF(A{r}=\"\",\"\",SUMIFS('{COM}'!$O:$O,'{COM}'!$J:$J,E{r}))"
            fat[f"AD{r}"] = (f"=IF(A{r}=\"\",\"\",IF(AC{r}=0,\"Pendente\",IF(COUNTIFS('{COM}'!$J:$J,E{r},"
                             f"'{COM}'!$X:$X,\"Divergência\")>0,\"Divergência\",\"Recebida\")))")
            fat[f"AC{r}"].number_format = MOEDA
            fat[f"AC{r}"].fill = fat[f"AD{r}"].fill = CINZA
    nomes = _estipulantes(wb)
    idx = _indice_faturas(wb)
    ja = {str(wp[f"C{r}"].value): r for r in range(2, N + 2) if wp[f"C{r}"].value}
    hoje = datetime.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    rel = []
    for path in arquivos:
        arq = f"{pasta_rel}/{os.path.basename(path)}"
        linhas = ler_xls(path)
        por_op = collections.OrderedDict()
        for l in sorted(linhas, key=lambda x: (x["data"], x["op"])):
            por_op.setdefault(l["op"], []).append(l)
        for op, ls in por_op.items():
            bruta = round(sum(l["valor"] for l in ls), 2)
            data = ls[0]["data"]
            comp = data.strftime("%Y-%m")
            if op in ja:
                r = ja[op]
                ant = wp[f"E{r}"].value or 0
                rel.append((op, "já gravada", "sem diferenças" if abs(ant - bruta) <= 0.01
                            else f"total no arquivo {bruta:.2f} ≠ gravado {ant:.2f}"))
                continue
            r = _prox(wp, "C")
            _w(wp, f"A{r}", comp); _w(wp, f"B{r}", SEGURADORA); _w(wp, f"C{r}", op)
            _w(wp, f"D{r}", data, DATA); _w(wp, f"E{r}", bruta, MOEDA)
            _w(wp, f"P{r}", arq); _w(wp, f"Q{r}", hoje, DATA)
            _w(wp, f"R{r}", "Extrato Bradesco (.xls) não traz tributações nem crédito líquido; conferência só da soma dos lançamentos")
            wp[f"L{r}"] = f"=IF(C{r}=\"\",\"\",COUNTIFS('{COM}'!$B:$B,C{r}))"
            wp[f"M{r}"] = f"=IF(C{r}=\"\",\"\",SUMIFS('{COM}'!$O:$O,'{COM}'!$B:$B,C{r}))"
            wp[f"N{r}"] = f"=IF(C{r}=\"\",\"\",IF(ABS(M{r}-E{r})<=0.01,\"OK\",\"Divergência\"))"
            for col in "LMN":
                wp[f"{col}{r}"].fill = CINZA
            wp[f"M{r}"].number_format = MOEDA
            ja[op] = r
            grupos = collections.OrderedDict()
            achou = {}
            for l in ls:
                fat, ok = fatura_base(l, idx)
                achou[fat] = ok
                grupos.setdefault((fat, l["taxa"], l["tipo"]), []).append(l)
            for (fat, taxa, tipo), gs in grupos.items():
                l0 = gs[0]
                k = _prox(wc, "B")
                premio = round(sum(x["premio"] for x in gs), 2)
                parcelas = sorted({x["parcela"] for x in gs}, key=lambda p: int(p) if p.isdigit() else 0)
                _w(wc, f"A{k}", comp); _w(wc, f"B{k}", op); _w(wc, f"C{k}", data, DATA); _w(wc, f"D{k}", SEGURADORA)
                _w(wc, f"E{k}", nomes.get(l0["apolice"]) or l0["segurado"]); _w(wc, f"F{k}", l0["suc"])
                _w(wc, f"G{k}", l0["ramo"]); _w(wc, f"H{k}", f"{l0['apolice']}/{l0['sub']}")
                _w(wc, f"I{k}", l0["apolice"]); _w(wc, f"J{k}", fat)
                _w(wc, f"K{k}", int(parcelas[0]) if len(parcelas) == 1 and parcelas[0].isdigit() else ", ".join(parcelas))
                _w(wc, f"L{k}", data, DATA); _w(wc, f"M{k}", premio, MOEDA)
                _w(wc, f"N{k}", round(taxa, 2), "0.00"); _w(wc, f"O{k}", round(sum(x["valor"] for x in gs), 2), MOEDA)
                _w(wc, f"P{k}", TIPOS.get(tipo, tipo)); _w(wc, f"Y{k}", arq); _w(wc, f"Z{k}", hoje, DATA)
                _formulas_bradesco(wc, k)
                o = [f"{len(gs)} lançamento(s) no extrato, endosso {l0['endosso']}"]
                if not fat: o[0] = f"{len(gs)} lançamento(s) no extrato, sem endosso: pagamento sem fatura (tipo {tipo}, {l0['segurado']})"
                elif not achou[fat]: o.append("Fatura não encontrada na base (nº deduzido do endosso)")
                if tipo == "AN": o.append(f"Agenciamento {taxa:g}% (parcela {', '.join(parcelas)})".replace(".", ","))
                if taxa < 0.1: o.append("Lançamento 0,01% por apólice (parcela 0)")
                if len(parcelas) > 1: o.append("Mais de uma parcela no mesmo pagamento (retroativo/inclusão)")
                _w(wc, f"AA{k}", "; ".join(o))
            rel.append((op, "gravada", f"{data:%d/%m/%Y}, {len(ls)} lançamentos em {len(grupos)} linhas, bruta {bruta:.2f}"))
    return rel


if __name__ == "__main__":
    base, *arqs = sys.argv[1:]
    wb = openpyxl.load_workbook(base)
    pasta = os.path.relpath(os.path.dirname(os.path.abspath(arqs[0])),
                            os.path.dirname(os.path.abspath(base))).replace("\\", "/")
    for linha in importar(wb, arqs, pasta):
        print(*linha, sep=" | ")
    wb.save(base)
